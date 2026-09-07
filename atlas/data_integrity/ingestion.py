"""Workbook ingestion: openpyxl workbook -> typed ingestion records.

Schema-agnostic. Every sheet is resolved through the :class:`MappingConfig`;
every header is resolved to a canonical field or preserved as an *unknown*
column with its value (never dropped). Malformed inputs are handled
defensively:

* unreadable / corrupt workbook -> :class:`IngestionResult` with ``load_error``
* sheet that maps to no known entity -> diagnostic, no records
* header-only sheet -> zero records, ``data_row_count == 0``
* fully-blank rows -> skipped and counted (``blank_row_count``)
* duplicate headers -> first mapping wins, duplicates recorded
* blank / merged-cell header gaps -> captured as unnamed unknown columns

Nothing here raises on bad data: bad data becomes findings later.
"""

from __future__ import annotations

import datetime
from pathlib import Path
from typing import Any, Callable, Optional

from openpyxl import load_workbook

from atlas.data_integrity.mapping import MappingConfig, default_mapping
from atlas.data_integrity.normalizers import apply_normalizer
from atlas.data_integrity.records import (
    FieldValue,
    IngestionRecord,
    IngestionResult,
    Provenance,
    SheetDiagnostic,
)


def _default_clock() -> str:
    return datetime.datetime.now(datetime.timezone.utc).isoformat()


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _trim_trailing(cells: list[Any]) -> list[Any]:
    out = list(cells)
    while out and _is_blank(out[-1]):
        out.pop()
    return out


class WorkbookIngestor:
    def __init__(
        self,
        mapping: Optional[MappingConfig] = None,
        clock: Callable[[], str] = _default_clock,
    ):
        self.mapping = mapping or default_mapping()
        self.clock = clock

    # ------------------------------------------------------------------
    def ingest_path(self, path: Path) -> IngestionResult:
        path = Path(path)
        result = IngestionResult(source_file=str(path))
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - defensive: any openpyxl/zip error
            result.load_error = f"{type(exc).__name__}: {exc}"
            return result

        try:
            for ws in wb.worksheets:
                diag, records = self._ingest_sheet(str(path), ws)
                result.sheets.append(diag)
                result.records.extend(records)
        finally:
            wb.close()
        return result

    # ------------------------------------------------------------------
    def _ingest_sheet(self, source_file: str, ws) -> tuple[SheetDiagnostic, list[IngestionRecord]]:
        sheet_name = ws.title
        schema = self.mapping.resolve_sheet(sheet_name)
        diag = SheetDiagnostic(
            sheet_name=sheet_name,
            entity_type=schema.entity_type if schema else None,
            resolved=schema is not None,
        )
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            diag.notes.append("empty sheet (no rows at all)")
            return diag, []

        header_cells = _trim_trailing(list(rows[0]))
        diag.header_columns = [("" if _is_blank(c) else str(c)) for c in header_cells]

        if schema is None:
            diag.notes.append("sheet did not resolve to any known entity; skipped")
            return diag, []

        # Resolve each header column position -> canonical (or unknown).
        col_canonical: list[Optional[str]] = []
        seen_canonical: dict[str, int] = {}
        for idx, header in enumerate(header_cells):
            if _is_blank(header):
                col_canonical.append(None)  # unnamed; may still carry data
                continue
            canonical = self.mapping.resolve_column(schema, str(header))
            if canonical is None:
                col_canonical.append(None)
                diag.unknown_columns.append(str(header))
            elif canonical in seen_canonical:
                # duplicate header mapping -> keep first, record duplicate
                diag.duplicate_headers.append(str(header))
                col_canonical.append(None)
            else:
                seen_canonical[canonical] = idx
                col_canonical.append(canonical)
                diag.mapped_columns.append(canonical)

        records: list[IngestionRecord] = []
        source_stem = Path(source_file).stem
        for row_offset, raw_row in enumerate(rows[1:], start=2):
            cells = list(raw_row)
            if all(_is_blank(c) for c in cells):
                diag.blank_row_count += 1
                continue
            diag.data_row_count += 1
            record = self._build_record(
                source_file, source_stem, sheet_name, schema.entity_type,
                row_offset, header_cells, col_canonical, cells,
            )
            records.append(record)

        diag.record_count = len(records)
        return diag, records

    # ------------------------------------------------------------------
    def _build_record(
        self,
        source_file: str,
        source_stem: str,
        sheet_name: str,
        entity_type: str,
        row_index: int,
        header_cells: list[Any],
        col_canonical: list[Optional[str]],
        cells: list[Any],
    ) -> IngestionRecord:
        provenance = Provenance(
            source_file=source_file,
            sheet_name=sheet_name,
            entity_type=entity_type,
            row_index=row_index,
            ingested_at=self.clock(),
            mapping_version=self.mapping.version,
        )
        record_id = f"{source_stem}::{sheet_name}::r{row_index}"
        record = IngestionRecord(
            record_id=record_id,
            entity_type=entity_type,
            provenance=provenance,
        )
        schema = self.mapping.schema_for(entity_type)
        for col_idx, canonical in enumerate(col_canonical):
            raw = cells[col_idx] if col_idx < len(cells) else None
            if canonical is None:
                if not _is_blank(raw):
                    header = header_cells[col_idx] if col_idx < len(header_cells) else None
                    key = str(header) if not _is_blank(header) else f"__unnamed_col{col_idx + 1}"
                    # avoid clobbering; keep deterministic distinct keys
                    if key in record.unknown:
                        key = f"{key}__dup{col_idx + 1}"
                    record.unknown[key] = raw
                continue
            spec = schema.spec_for(canonical) if schema else None
            normalizer = spec.normalizer if spec else "text"
            normalized = apply_normalizer(normalizer, raw)
            changed = normalized != raw and not (_is_blank(raw) and normalized in (None, ""))
            record.fields[canonical] = FieldValue(
                canonical=canonical,
                raw=raw,
                normalized=normalized,
                normalizer=normalizer,
                changed=changed,
            )
        return record


def ingest_workbook(
    path: Path,
    mapping: Optional[MappingConfig] = None,
    clock: Callable[[], str] = _default_clock,
) -> IngestionResult:
    """Convenience wrapper."""
    return WorkbookIngestor(mapping=mapping, clock=clock).ingest_path(Path(path))
