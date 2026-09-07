"""Phase 0.9 callable pipeline.

``run_phase09`` executes the whole data-integrity flow end to end, fully
offline (no browser, no network, ``controller="none"``):

    audit fixture -> ingest -> normalize -> validate -> identity -> reconcile
    -> adversarial scenarios (A..AU) + expansion timings
    -> atomic XLSX + JSON quality report + audit JSON

It returns a :class:`Phase09Result` carrying every output path, per-stage
timing, headline counts, and the fixture's pre/post SHA-256 (which MUST be
identical — the pipeline never mutates the immutable fixture).

The same building blocks (:func:`run_scenario`, :func:`run_expansion`) are
reused by the Phase 0.9 test-suite so tests assert on exactly what the
deliverable pipeline produces.
"""

from __future__ import annotations

import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Optional

from atlas.controllers.base import Controller, NullController
from atlas.data_integrity import adversarial as adv
from atlas.data_integrity.identity import IdentityResolver
from atlas.data_integrity.ingestion import WorkbookIngestor
from atlas.data_integrity.mapping import MappingConfig, default_mapping
from atlas.data_integrity.reconciliation import Reconciler
from atlas.data_integrity.report import build_quality_report
from atlas.data_integrity.report_writer import (
    WriteResult,
    write_json_atomic,
    write_workbook_atomic,
)
from atlas.data_integrity.validation import Action, Validator
from atlas.persistence.sqlite import StateStore

FIXED_CLOCK = "2026-09-06T00:00:00+00:00"


# ---------------------------------------------------------------------------
# Observed signals for a single workbook run
# ---------------------------------------------------------------------------
@dataclass
class Observed:
    record_count: int = 0
    load_ok: bool = True
    load_error: Optional[str] = None
    finding_codes: set[str] = field(default_factory=set)
    relationships: set[str] = field(default_factory=set)
    quarantine: int = 0
    unknown_col: bool = False
    dup_header: bool = False
    blank_rows: int = 0
    rec_closed: int = 0
    rec_created: int = 0
    ingest_seconds: float = 0.0
    total_seconds: float = 0.0

    def satisfies(self, signal: str) -> bool:
        if signal == "load_ok":
            return self.load_ok
        if signal == "unknown_col":
            return self.unknown_col
        if signal == "dup_header":
            return self.dup_header
        if signal.startswith("code:"):
            return signal[5:] in self.finding_codes
        if signal.startswith("rel:"):
            return signal[4:] in self.relationships
        if signal.startswith("quar>="):
            return self.quarantine >= int(signal[6:])
        if signal.startswith("rec_closed>="):
            return self.rec_closed >= int(signal[12:])
        if signal.startswith("blank>="):
            return self.blank_rows >= int(signal[7:])
        if signal.startswith("records=="):
            return self.record_count == int(signal[9:])
        if signal.startswith("records>="):
            return self.record_count >= int(signal[9:])
        return False


def _run_workbook(
    path: Path,
    mapping: MappingConfig,
    store: Optional[StateStore],
    di_run_id: str,
    *,
    reconcile: bool = True,
) -> Observed:
    obs = Observed()
    ingestor = WorkbookIngestor(mapping=mapping, clock=lambda: FIXED_CLOCK)
    t0 = time.perf_counter()
    ingestion = ingestor.ingest_path(path)
    obs.ingest_seconds = time.perf_counter() - t0

    if ingestion.load_error is not None:
        obs.load_ok = False
        obs.load_error = ingestion.load_error
        obs.total_seconds = time.perf_counter() - t0
        return obs

    obs.record_count = len(ingestion.records)
    obs.blank_rows = sum(s.blank_row_count for s in ingestion.sheets)
    obs.unknown_col = any(s.unknown_columns for s in ingestion.sheets)
    obs.dup_header = any(s.duplicate_headers for s in ingestion.sheets)

    validator = Validator(mapping)
    resolver = IdentityResolver(mapping)
    resolution = resolver.resolve(ingestion.records)
    report = validator.validate(ingestion.records)

    obs.finding_codes = {f.code for f in report.findings}
    obs.quarantine = len(report.quarantined) + len(report.rejected)
    for c in resolution.clusters:
        obs.relationships.update(r.value for r in c.relationships)

    if reconcile and store is not None:
        rc = Reconciler(store).reconcile(
            di_run_id, str(path), ingestion.records, report, resolution,
            source_hash="scenario", mapping_version=mapping.version,
        )
        obs.rec_closed = rc.closed
        obs.rec_created = rc.created

    obs.total_seconds = time.perf_counter() - t0
    return obs


# ---------------------------------------------------------------------------
# Scenario + expansion runners
# ---------------------------------------------------------------------------
def run_scenario(
    scenario: adv.Scenario,
    base_rows: list[dict[str, Any]],
    out_dir: Path,
    *,
    mapping: Optional[MappingConfig] = None,
    seed: int = adv.DEFAULT_SEED,
) -> dict[str, Any]:
    """Generate one scenario copy, run the full pipeline, verify signals."""
    mapping = mapping or default_mapping()
    out_dir = Path(out_dir)
    path = adv.generate_scenario_copy(base_rows, scenario, out_dir, seed=seed)

    store_path = out_dir / f"_scen_{scenario.code}.sqlite"
    if store_path.exists():
        store_path.unlink()
    store = StateStore(store_path)
    try:
        reconcile = scenario.category != "performance"
        obs = _run_workbook(path, mapping, store, f"scen-{scenario.code}", reconcile=reconcile)
    finally:
        store.close()
    store_path.unlink(missing_ok=True)

    missing = [sig for sig in scenario.expected if not obs.satisfies(sig)]
    passed = not missing
    return {
        "code": scenario.code,
        "name": scenario.name,
        "category": scenario.category,
        "description": scenario.description,
        "represented": scenario.represented,
        "copy_file": path.name,
        "records": obs.record_count,
        "load_ok": obs.load_ok,
        "load_error": obs.load_error or "",
        "finding_codes": ", ".join(sorted(obs.finding_codes)),
        "relationships": ", ".join(sorted(obs.relationships)),
        "quarantine": obs.quarantine,
        "rec_closed": obs.rec_closed,
        "expected": ", ".join(scenario.expected),
        "missing_signals": ", ".join(missing),
        "passed": passed,
        "ingest_seconds": round(obs.ingest_seconds, 4),
    }


def run_expansion(
    n: int,
    base_rows: list[dict[str, Any]],
    out_dir: Path,
    *,
    mapping: Optional[MappingConfig] = None,
    seed: int = adv.DEFAULT_SEED,
    reconcile: bool = False,
) -> dict[str, Any]:
    """Generate an N-row workbook and measure ingest (and optional reconcile)."""
    mapping = mapping or default_mapping()
    out_dir = Path(out_dir)
    path = adv.generate_expansion(base_rows, n, out_dir, seed=seed)

    store = None
    store_path = out_dir / f"_exp_{n}.sqlite"
    if reconcile:
        if store_path.exists():
            store_path.unlink()
        store = StateStore(store_path)
    try:
        obs = _run_workbook(path, mapping, store, f"exp-{n}", reconcile=reconcile)
    finally:
        if store is not None:
            store.close()
            store_path.unlink(missing_ok=True)

    throughput = round(obs.record_count / obs.ingest_seconds, 1) if obs.ingest_seconds > 0 else 0.0
    return {
        "rows": n,
        "file": path.name,
        "file_bytes": path.stat().st_size,
        "records_ingested": obs.record_count,
        "ingest_seconds": round(obs.ingest_seconds, 4),
        "total_seconds": round(obs.total_seconds, 4),
        "rows_per_second": throughput,
        "reconciled": reconcile,
        "passed": obs.record_count == n and obs.load_ok,
    }


# ---------------------------------------------------------------------------
# Full pipeline
# ---------------------------------------------------------------------------
@dataclass
class Phase09Result:
    source_file: str
    source_sha256_before: str
    source_sha256_after: str
    report_xlsx: str
    report_json: str
    audit_json: str
    timings: dict[str, float] = field(default_factory=dict)
    counts: dict[str, Any] = field(default_factory=dict)
    scenarios: list[dict[str, Any]] = field(default_factory=list)
    performance: list[dict[str, Any]] = field(default_factory=list)
    write_results: dict[str, Any] = field(default_factory=dict)

    @property
    def fixture_unchanged(self) -> bool:
        return self.source_sha256_before == self.source_sha256_after

    @property
    def all_scenarios_passed(self) -> bool:
        return all(s["passed"] for s in self.scenarios)

    def to_dict(self) -> dict[str, Any]:
        return {
            "source_file": self.source_file,
            "source_sha256_before": self.source_sha256_before,
            "source_sha256_after": self.source_sha256_after,
            "fixture_unchanged": self.fixture_unchanged,
            "report_xlsx": self.report_xlsx,
            "report_json": self.report_json,
            "audit_json": self.audit_json,
            "timings": self.timings,
            "counts": self.counts,
            "scenario_pass_count": sum(1 for s in self.scenarios if s["passed"]),
            "scenario_total": len(self.scenarios),
            "performance": self.performance,
            "write_results": self.write_results,
        }


def audit_fixture(path: Path, mapping: MappingConfig) -> dict[str, Any]:
    path = Path(path)
    sheets: list[dict[str, Any]] = []
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            header = [("" if c is None else str(c)) for c in (rows[0] if rows else ())]
            while header and header[-1] == "":
                header.pop()
            data_rows = sum(
                1 for r in rows[1:]
                if any(c is not None and not (isinstance(c, str) and c.strip() == "") for c in r)
            )
            schema = mapping.resolve_sheet(ws.title)
            sheets.append(
                {
                    "sheet": ws.title,
                    "entity_type": schema.entity_type if schema else None,
                    "header_columns": len(header),
                    "headers": header,
                    "data_rows": data_rows,
                }
            )
    finally:
        wb.close()
    return {
        "source_file": str(path),
        "sha256": adv.sha256_file(path),
        "bytes": path.stat().st_size,
        "mapping_version": mapping.version,
        "sheets": sheets,
    }


def run_phase09(
    fixture_path: Path,
    output_dir: Path,
    *,
    mapping: Optional[MappingConfig] = None,
    controller: Optional[Controller] = None,
    seed: int = adv.DEFAULT_SEED,
    generated_at: str = FIXED_CLOCK,
    expansion_sizes: tuple[int, ...] = (500, 5000, 20000),
    include_scenarios: bool = True,
    generated_dir: Optional[Path] = None,
) -> Phase09Result:
    """Execute Phase 0.9 end to end and write the deliverable artifacts."""
    mapping = mapping or default_mapping()
    controller = controller or NullController()  # offline, no-LLM by contract
    fixture_path = Path(fixture_path)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    # Mutated workbooks are fixtures, rather than report output.  Keeping
    # them beside (but never inside) the immutable real fixture makes their
    # disposable provenance explicit and prevents a report rerun from
    # treating an adversarial workbook as production-shaped output.
    adv_dir = (
        Path(generated_dir)
        if generated_dir is not None
        else fixture_path.parent.parent / "generated"
    )
    adv_dir.mkdir(parents=True, exist_ok=True)

    timings: dict[str, float] = {}
    before_hash = adv.sha256_file(fixture_path)

    # 1) Audit
    t = time.perf_counter()
    audit = audit_fixture(fixture_path, mapping)
    timings["audit_seconds"] = round(time.perf_counter() - t, 4)

    # 2) Main pipeline on the real fixture (into a fresh canonical store).
    state_path = output_dir / "phase09_state.sqlite"
    for suffix in ("", "-wal", "-shm"):
        p = Path(str(state_path) + suffix)
        if p.exists():
            p.unlink()
    store = StateStore(state_path)
    try:
        ingestor = WorkbookIngestor(mapping=mapping, clock=lambda: generated_at)
        t = time.perf_counter()
        ingestion = ingestor.ingest_path(fixture_path)
        timings["ingest_seconds"] = round(time.perf_counter() - t, 4)

        t = time.perf_counter()
        resolution = IdentityResolver(mapping).resolve(ingestion.records)
        timings["identity_seconds"] = round(time.perf_counter() - t, 4)

        t = time.perf_counter()
        validation = Validator(mapping).validate(ingestion.records)
        timings["validation_seconds"] = round(time.perf_counter() - t, 4)

        t = time.perf_counter()
        reconciliation = Reconciler(store).reconcile(
            "phase09-main", str(fixture_path), ingestion.records, validation, resolution,
            source_hash=before_hash, mapping_version=mapping.version,
        )
        timings["reconcile_seconds"] = round(time.perf_counter() - t, 4)

        canonical_count = store.count_canonical_jobs()
        observation_count = store.count_observations()
        quarantine_rows = store.count_quarantine()
    finally:
        store.close()

    # 3) Adversarial scenarios + expansion performance.
    scenario_rows: list[dict[str, Any]] = []
    performance_rows: list[dict[str, Any]] = []
    if include_scenarios:
        base_rows = adv.extract_sheet_rows(fixture_path, "All_Jobs")
        t = time.perf_counter()
        for scenario in adv.SCENARIOS:
            if scenario.category == "performance":
                # measured in the performance section (single ingest per size)
                continue
            scenario_rows.append(
                run_scenario(scenario, base_rows, adv_dir, mapping=mapping, seed=seed)
            )
        timings["scenarios_seconds"] = round(time.perf_counter() - t, 4)

        t = time.perf_counter()
        for n in expansion_sizes:
            perf = run_expansion(n, base_rows, adv_dir, mapping=mapping, seed=seed)
            performance_rows.append(perf)
            # represent the matching expansion scenario code in the scenario table
            code = {500: "AR", 5000: "AS", 20000: "AT"}.get(n)
            sc = adv.scenario_by_code(code) if code else None
            if sc is not None:
                scenario_rows.append(
                    {
                        "code": sc.code, "name": sc.name, "category": sc.category,
                        "description": sc.description, "represented": True,
                        "copy_file": perf["file"], "records": perf["records_ingested"],
                        "load_ok": perf["passed"], "load_error": "",
                        "finding_codes": "", "relationships": "", "quarantine": 0,
                        "rec_closed": 0, "expected": f"records=={n}",
                        "missing_signals": "" if perf["passed"] else f"records=={n}",
                        "passed": perf["passed"], "ingest_seconds": perf["ingest_seconds"],
                    }
                )
        timings["expansion_seconds"] = round(time.perf_counter() - t, 4)
        scenario_rows.sort(key=lambda r: (len(r["code"]), r["code"]))

    # 4) Verify the immutable fixture was never mutated.
    after_hash = adv.sha256_file(fixture_path)

    # 5) Build + atomically write the report artifacts.
    report = build_quality_report(
        generated_at=generated_at,
        audit=audit,
        ingestion=ingestion,
        validation=validation,
        identity=resolution,
        reconciliation=reconciliation.summary(),
        adversarial=scenario_rows,
        performance=performance_rows,
    )

    xlsx_path = output_dir / "Atlas_PHASE09_Data_Quality_Report.xlsx"
    json_path = output_dir / "Atlas_PHASE09_Data_Quality_Report.json"
    audit_path = output_dir / "PHASE09_AUDIT.json"

    t = time.perf_counter()
    wb = report.to_workbook()
    xlsx_res = write_workbook_atomic(xlsx_path, wb, expected_sheets=("Summary", "Findings"))
    json_res = write_json_atomic(json_path, report.to_dict())
    audit_doc = {
        "audit": audit,
        "source_sha256_before": before_hash,
        "source_sha256_after": after_hash,
        "fixture_unchanged": before_hash == after_hash,
        "timings": timings,
        "counts": {
            "records": len(ingestion.records),
            "canonical_jobs": canonical_count,
            "observations": observation_count,
            "quarantine_rows": quarantine_rows,
            "findings": len(validation.findings),
        },
    }
    audit_res = write_json_atomic(audit_path, audit_doc)
    timings["report_write_seconds"] = round(time.perf_counter() - t, 4)

    counts = {
        "records": len(ingestion.records),
        "sheets": len(ingestion.sheets),
        "findings": len(validation.findings),
        "findings_by_severity": validation.counts_by_severity(),
        "quarantined": len(validation.quarantined),
        "identity_clusters": len(resolution.clusters),
        "relationship_counts": resolution.relationship_counts(),
        "canonical_jobs": canonical_count,
        "observations": observation_count,
        "reconciliation": reconciliation.summary(),
    }

    return Phase09Result(
        source_file=str(fixture_path),
        source_sha256_before=before_hash,
        source_sha256_after=after_hash,
        report_xlsx=xlsx_res.written_path,
        report_json=json_res.written_path,
        audit_json=audit_res.written_path,
        timings=timings,
        counts=counts,
        scenarios=scenario_rows,
        performance=performance_rows,
        write_results={
            "xlsx": xlsx_res.to_dict(),
            "json": json_res.to_dict(),
            "audit": audit_res.to_dict(),
        },
    )
