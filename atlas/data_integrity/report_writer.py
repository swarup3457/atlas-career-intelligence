"""Atomic, crash-safe report writers.

The write protocol for every artifact is the same and is what makes the
report *reliable* (see docs/REPORT_RELIABILITY.md):

1. **temp** — serialize to a sibling ``<name>.part`` temp file (same
   directory, so :func:`os.replace` is atomic on the same filesystem).
2. **reopen / validate** — reopen the temp file and confirm it is a
   structurally valid artifact (a real xlsx zip with the expected sheets,
   or parseable JSON). A corrupt temp never becomes the final file.
3. **os.replace** — atomically swap the validated temp into place. Readers
   only ever observe the old complete file or the new complete file.

Two failure modes are handled explicitly:

* **locked final** (Windows: the workbook is open in Excel) — ``os.replace``
  raises ``PermissionError``. We fall back to a *deterministic* alternate
  file (``<stem>.locked<suffix>``, then ``.locked-2`` ...) so the run still
  produces output at a predictable path instead of losing it.
* **incomplete temp** — a crash between step 1 and 3 leaves a ``.part``
  file. :func:`recover_incomplete_temps` finds and removes those stale
  temps so they can never be mistaken for real output.
"""

from __future__ import annotations

import json
import os
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Optional

from openpyxl import Workbook, load_workbook

TEMP_MARKER = ".part"
_MAX_ALTERNATES = 50


class ReportWriteError(RuntimeError):
    """Raised when an artifact could not be written safely (e.g. the temp
    file failed validation and was discarded)."""


@dataclass
class WriteResult:
    requested_path: str
    written_path: str
    locked: bool = False
    validated: bool = True
    recovered: list[str] = field(default_factory=list)

    @property
    def used_alternate(self) -> bool:
        return Path(self.requested_path) != Path(self.written_path)

    def to_dict(self) -> dict[str, Any]:
        return {
            "requested_path": self.requested_path,
            "written_path": self.written_path,
            "locked": self.locked,
            "validated": self.validated,
            "used_alternate": self.used_alternate,
            "recovered": list(self.recovered),
        }


def _temp_path(final: Path) -> Path:
    # Keep the real suffix (e.g. .xlsx) so validators that check extension
    # still accept the temp; insert a ``.part`` marker before it so the temp
    # is unmistakably incomplete and discoverable for crash recovery.
    return final.with_name(f"{final.stem}{TEMP_MARKER}{final.suffix}")


def _alternate_path(final: Path, index: int) -> Path:
    # Deterministic: <stem>.locked<suffix>, then <stem>.locked-2<suffix>, ...
    if index == 1:
        return final.with_name(f"{final.stem}.locked{final.suffix}")
    return final.with_name(f"{final.stem}.locked-{index}{final.suffix}")


def recover_incomplete_temps(directory: Path, final_name: Optional[str] = None) -> list[str]:
    """Remove leftover ``*.part.*`` temp files from a previous crashed write.

    If ``final_name`` is given, only that artifact's temp is considered;
    otherwise every ``*.part.*`` file in ``directory`` is swept. Returns the
    list of removed paths (as strings), sorted for determinism.
    """
    directory = Path(directory)
    if not directory.exists():
        return []
    removed: list[str] = []
    if final_name is not None:
        final = directory / final_name
        candidates = [_temp_path(final)]
    else:
        candidates = sorted(
            set(directory.glob(f"*{TEMP_MARKER}.*")) | set(directory.glob(f"*{TEMP_MARKER}"))
        )
    for temp in candidates:
        if temp.exists() and temp.is_file():
            try:
                temp.unlink()
                removed.append(str(temp))
            except OSError:
                # best-effort recovery; a still-locked temp is left in place
                continue
    return sorted(removed)


def _atomic_promote(temp: Path, final: Path) -> WriteResult:
    """os.replace temp -> final, falling back to a deterministic alternate
    when the final path is locked."""
    try:
        os.replace(temp, final)
        return WriteResult(str(final), str(final), locked=False)
    except PermissionError:
        pass
    except OSError as exc:
        # Windows raises a plain OSError (winerror 32) when the target is
        # open in another process; treat that like a lock too.
        if getattr(exc, "winerror", None) not in (32, 33):
            # Not a sharing violation — clean up and re-raise as a write error.
            temp.unlink(missing_ok=True)
            raise ReportWriteError(f"Failed to promote temp to {final}: {exc}") from exc

    for index in range(1, _MAX_ALTERNATES + 1):
        alt = _alternate_path(final, index)
        try:
            os.replace(temp, alt)
            return WriteResult(str(final), str(alt), locked=True)
        except (PermissionError, OSError) as exc:
            if getattr(exc, "winerror", None) in (32, 33) or isinstance(exc, PermissionError):
                continue
            temp.unlink(missing_ok=True)
            raise ReportWriteError(f"Failed to promote temp to alternate {alt}: {exc}") from exc
    temp.unlink(missing_ok=True)
    raise ReportWriteError(
        f"All {_MAX_ALTERNATES} alternate paths for {final} were locked."
    )


def write_workbook_atomic(
    final_path: Path,
    workbook: Workbook,
    *,
    expected_sheets: Optional[Iterable[str]] = None,
    recover_first: bool = True,
) -> WriteResult:
    """Atomically write an openpyxl workbook with reopen/validate + replace."""
    final = Path(final_path)
    final.parent.mkdir(parents=True, exist_ok=True)
    recovered: list[str] = []
    if recover_first:
        recovered = recover_incomplete_temps(final.parent, final.name)

    temp = _temp_path(final)
    temp.unlink(missing_ok=True)
    try:
        workbook.save(str(temp))
    except Exception as exc:  # noqa: BLE001
        temp.unlink(missing_ok=True)
        raise ReportWriteError(f"Failed to serialize workbook to {temp}: {exc}") from exc

    # Reopen + validate the temp before it is allowed to become final.
    try:
        check = load_workbook(temp, read_only=True)
        sheetnames = set(check.sheetnames)
        check.close()
        if expected_sheets is not None:
            missing = set(expected_sheets) - sheetnames
            if missing:
                raise ReportWriteError(
                    f"Validation failed: temp workbook missing sheets {sorted(missing)}."
                )
    except ReportWriteError:
        temp.unlink(missing_ok=True)
        raise
    except Exception as exc:  # noqa: BLE001
        temp.unlink(missing_ok=True)
        raise ReportWriteError(f"Temp workbook failed reopen/validation: {exc}") from exc

    result = _atomic_promote(temp, final)
    result.recovered = recovered
    return result


def write_json_atomic(
    final_path: Path,
    obj: Any,
    *,
    recover_first: bool = True,
    indent: int = 2,
) -> WriteResult:
    """Atomically write JSON with reopen/validate + replace."""
    final = Path(final_path)
    final.parent.mkdir(parents=True, exist_ok=True)
    recovered: list[str] = []
    if recover_first:
        recovered = recover_incomplete_temps(final.parent, final.name)

    temp = _temp_path(final)
    temp.unlink(missing_ok=True)
    try:
        text = json.dumps(obj, indent=indent, sort_keys=True)
    except (TypeError, ValueError) as exc:
        raise ReportWriteError(f"Object is not JSON-serializable: {exc}") from exc
    with open(temp, "w", encoding="utf-8") as fh:
        fh.write(text)
        fh.flush()
        os.fsync(fh.fileno())

    # Reopen + validate.
    try:
        with open(temp, "r", encoding="utf-8") as fh:
            json.load(fh)
    except Exception as exc:  # noqa: BLE001
        temp.unlink(missing_ok=True)
        raise ReportWriteError(f"Temp JSON failed reopen/validation: {exc}") from exc

    result = _atomic_promote(temp, final)
    result.recovered = recovered
    return result
