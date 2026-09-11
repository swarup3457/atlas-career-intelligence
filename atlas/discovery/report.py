from __future__ import annotations

import datetime
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from atlas.reporting.trust_boundary import partition_jobs

SHEETS = (
    "Validated_Jobs", "Rejected_Jobs", "Foreign_Leads", "Raw_Job_Leads",
    "Discovery_Providers", "Company_Candidates", "Official_Verification",
    "Company_Coverage", "Lane_Coverage", "Source_Health", "Evidence_Audit",
    "Session_Audit", "Run_Summary",
)

V3_SHEETS = SHEETS + ("Selection_Audit", "Discovery_Funnel")


def build_discovery_v3_workbook(*, evidence_root: Path, live_root: Path, run_id: str, output_root: Path) -> Path:
    """Build the V3 workbook from the existing Discovery V2 evidence only."""
    output = build_discovery_workbook(evidence_root=evidence_root, live_root=live_root, run_id=run_id, output_root=output_root)
    from openpyxl import load_workbook
    timestamp = datetime.datetime.now(datetime.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    v3 = Path(output_root) / run_id / f"Atlas_VSCode_DiscoveryV3_{timestamp}_{run_id}.xlsx"
    wb = load_workbook(output)
    for sheet in ("Selection_Audit", "Discovery_Funnel"):
        wb.create_sheet(sheet)
    wb["Selection_Audit"].append(["company", "selected", "reason"])
    verification = _load(evidence_root / "official_verification.json", [])
    for candidate in verification:
        wb["Selection_Audit"].append([candidate.get("company", ""), candidate.get("selected", False), candidate.get("selection_reason", candidate.get("deferral_reason", ""))])
    wb["Discovery_Funnel"].append(["stage", "count"])
    raw = _load(evidence_root / "freehire_raw.json", [])
    queued = _load(evidence_root / "freehire_queued.json", [])
    wb["Discovery_Funnel"].append(["raw_leads", len(raw)])
    wb["Discovery_Funnel"].append(["unique_leads", len({_lead_identity(lead) for lead in raw})])
    wb["Discovery_Funnel"].append(["prefiltered_queued", len(queued)])
    wb["Discovery_Funnel"].append(["companies_verified", len(verification)])
    wb["Discovery_Funnel"].append(["companies_selected", sum(1 for item in verification if item.get("selected"))])
    wb.save(v3)
    load_workbook(v3).close()
    return v3


def reconcile_metrics(*, raw: int, deduped: int, queued: int, verified_companies: int, selected_companies: int, primary_attempts: int, correction_attempts: int, total_attempts: int, validated: int, rejected: int, foreign: int) -> list[str]:
    problems: list[str] = []
    if not raw >= deduped >= queued:
        problems.append("raw >= deduped >= queued invariant failed")
    if verified_companies < selected_companies:
        problems.append("verified companies fewer than selected companies")
    if primary_attempts + correction_attempts != total_attempts:
        problems.append("primary+correction attempts do not equal total attempts")
    if min(validated, rejected, foreign) < 0:
        problems.append("job counts cannot be negative")
    return problems


def _load(path: Path, default: Any) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return default


def _lead_identity(lead: dict[str, Any]) -> str:
    if lead.get("identity"):
        return str(lead["identity"])
    canonical = str(lead.get("canonical_url") or "").strip().lower().rstrip("/")
    if canonical:
        return "url:" + canonical
    provider = str(lead.get("provider") or "").strip().lower()
    stable = str(lead.get("provider_id") or lead.get("raw_reference") or "").strip()
    if provider and stable:
        return f"provider:{provider}:{stable}"
    key = "|".join(str(lead.get(k) or "").lower().strip() for k in ("company", "title", "location", "posted_date"))
    import hashlib
    return "lead:" + hashlib.sha256(key.encode("utf-8")).hexdigest()[:24]


def recommendation_tier(job: dict[str, Any]) -> str:
    experience = str(job.get("experience_text", "")).lower()
    stack = {str(value).lower() for value in (job.get("stack") or [])}
    lane = str(job.get("lane", "")).upper()
    if any(token in experience for token in ("4+", "5+", "6+", "7+", "8+", "10+", "12+")):
        return "STRETCH"
    if lane in {"JAVA_BACKEND", "JAVA_FULLSTACK", "REACT_FRONTEND", "DOTNET"} and stack:
        if any(token in stack for token in ("java", "spring", "react", "reactjs", ".net", "c#")):
            return "STRONG_MATCH" if any(token in experience for token in ("0-1", "1-2", "2+", "3+", "3-4")) else "GOOD_MATCH"
    return "REVIEWABLE"


def build_discovery_workbook(*, evidence_root: Path, live_root: Path, run_id: str, output_root: Path) -> Path:
    timestamp = datetime.datetime.now(datetime.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    output = Path(output_root) / run_id / f"Atlas_VSCode_DiscoveryV2_{timestamp}_{run_id}.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in SHEETS:
        wb.create_sheet(sheet)

    raw = _load(evidence_root / "freehire_raw.json", [])
    queued = _load(evidence_root / "freehire_queued.json", [])
    health = _load(evidence_root / "freehire_health.json", [])
    manifest = _load(evidence_root / "selection_manifest.json", {})
    result_paths = sorted(live_root.rglob("result.json"))
    result_paths += sorted(live_root.rglob("correction-result.json"))
    result_paths += sorted(live_root.rglob("final-result.json"))
    all_result_paths = list(result_paths)
    result_by_task: dict[str, dict[str, Any]] = {}
    for path in result_paths:
        result = _load(path, {})
        task_key = str(result.get("task_id", path.parent.name))
        result_by_task[task_key] = result
    results = list(result_by_task.values())
    attempt_by_id: dict[str, dict[str, Any]] = {}
    for path in all_result_paths:
        attempt = _load(path, {})
        attempt_id = str(attempt.get("attempt_id", path.name))
        # Later artifact forms are normalized replacements for the same attempt,
        # not additional worker invocations.
        attempt_by_id[attempt_id] = attempt
    all_attempt_results = list(attempt_by_id.values())
    verification = _load(evidence_root / "official_verification.json", [])

    ws = wb["Validated_Jobs"]
    ws.append(["company", "title", "India location", "official_url", "requisition", "product_category", "discovery_source", "role_lane", "backend_stack", "frontend_stack", "experience", "posted_date", "recommendation", "strengths", "gaps", "verification_status"])
    validated_total = 0
    proposal_rejections_all: list[tuple[Any, Any]] = []
    for result in results:
        company_label = result.get("company_id", result.get("company", ""))
        validated_jobs, proposal_rejections = partition_jobs(result.get("jobs", []), official_domain=str(result.get("official_domain", "")), company=str(result.get("company", result.get("company_id", ""))))
        validated_total += len(validated_jobs)
        proposal_rejections_all.extend((company_label, rej) for rej in proposal_rejections)
        for vj in validated_jobs:
            ev = vj.evidence
            stack = vj.source.get("stack", []) or []
            frontend = [x for x in stack if str(x).lower() in {"react", "reactjs", "angular", "vue", "typescript", "javascript", "html", "css"}]
            backend = [x for x in stack if x not in frontend]
            ws.append([company_label, ev.title, ev.location, ev.official_url, ev.requisition_id, "", "; ".join(result.get("discovery_provenance", {}).values()), vj.lane, ", ".join(backend), ", ".join(frontend), ev.experience_text, ev.posted_date, vj.recommendation, "", "", vj.verification_status])

    for name, headers in {
        "Rejected_Jobs": ["company", "title", "location", "url", "lane", "reason_code", "detail"],
        "Foreign_Leads": ["company", "title", "location", "url", "reason"],
        "Raw_Job_Leads": ["identity", "provider", "mechanism", "query", "title", "company", "location", "posted_date", "source_url", "canonical_url", "category", "seniority", "description_available", "prefilter_status", "prefilter_reason"],
        "Discovery_Providers": ["provider", "status", "query_count", "raw_yield", "queued_yield", "errors"],
        "Company_Candidates": ["company", "official_domain", "category", "discovery_query", "clue_title", "clue_location", "confidence", "freshness", "status"],
        "Official_Verification": ["company", "company_id", "official_domain", "careers_url", "verified", "verified_job_count", "strongest_job", "india_evidence", "role_stack_evidence", "experience_evidence", "source_health", "selected", "selection_reason", "deferral_reason"],
        "Company_Coverage": ["company", "task_id", "status", "attempt_id", "worker_invocation_id", "attempt_kind", "attempt_number", "model", "browser_used", "result_path", "details_opened", "browser_errors"],
        "Lane_Coverage": ["company", "lane", "result_state", "terminal"],
        "Source_Health": ["provider_or_source", "status", "details"],
        "Evidence_Audit": ["company", "attempt_id", "result_hash", "detail_urls", "quotes", "source_health"],
        "Session_Audit": ["attempt_id", "company", "backend", "attempt_number", "action", "model", "browser_used", "errors"],
        "Run_Summary": ["run_id", "raw_leads", "deduped_leads", "queued_leads", "companies", "validated_jobs", "rejected_jobs", "foreign_leads", "terminal_tasks"],
    }.items():
        wb[name].append(headers)

    seen_rejections: set[str] = set()
    seen_foreign: set[str] = set()
    for result in all_attempt_results:
        company = result.get("company", result.get("company_id", ""))
        for rejection in result.get("rejections", []):
            key = json.dumps([company, rejection], sort_keys=True, default=str)
            if key not in seen_rejections:
                seen_rejections.add(key)
                wb["Rejected_Jobs"].append([company, rejection.get("title", ""), rejection.get("location", ""), rejection.get("url", rejection.get("canonical_url", "")), rejection.get("lane", ""), rejection.get("reason_code", ""), rejection.get("detail", rejection.get("reason", ""))])
        for foreign in result.get("foreign_leads", []):
            key = json.dumps([company, foreign], sort_keys=True, default=str)
            if key not in seen_foreign:
                seen_foreign.add(key)
                wb["Foreign_Leads"].append([company, foreign.get("title", ""), foreign.get("location", ""), foreign.get("url", foreign.get("official_url", "")), foreign.get("reason", "NON_INDIA_LOCATION")])

    for company_label, rej in proposal_rejections_all:
        wb["Rejected_Jobs"].append([company_label, rej.title, rej.location, rej.url, rej.lane, rej.reason_code, rej.detail])

    for result in results:
        company = result.get("company", result.get("company_id", ""))
        states = result.get("result_states", {})
        if isinstance(states, dict):
            for lane, state in states.items():
                wb["Lane_Coverage"].append([company, lane, state.get("state", state.get("status", "")) if isinstance(state, dict) else str(state), result.get("completion_claim") is True])
        attempts_for_company = [r for r in all_attempt_results if r.get("task_id") == result.get("task_id")]
        for attempt_result in attempts_for_company:
            kind = "CORRECTION" if "correction" in str(attempt_result.get("attempt_id", "")).lower() else "PRIMARY"
            wb["Company_Coverage"].append([company, result.get("task_id", ""), "COMPLETE" if result.get("completion_claim") else "PARTIAL", attempt_result.get("attempt_id", ""), attempt_result.get("worker_invocation_id", ""), kind, 2 if kind == "CORRECTION" else 1, attempt_result.get("model", "GitHub Copilot"), True, next((str(path) for path in all_result_paths if _load(path, {}).get("attempt_id") == attempt_result.get("attempt_id")), ""), len(attempt_result.get("detail_urls", [])), len(attempt_result.get("browser_errors", []))])
            attempt_errors = attempt_result.get("browser_errors", [])
            wb["Session_Audit"].append([attempt_result.get("attempt_id", ""), company, "VSCODE_SUBAGENT", 2 if kind == "CORRECTION" else 1, kind, attempt_result.get("model", "GitHub Copilot"), True, "; ".join(str(x) if isinstance(x, str) else json.dumps(x) for x in attempt_errors)])
        wb["Evidence_Audit"].append([company, result.get("attempt_id", ""), "", len(result.get("detail_urls", [])), len(result.get("evidence_quotes", [])), json.dumps(result.get("source_health", {}), sort_keys=True)])

    for lead in raw:
        wb["Raw_Job_Leads"].append([_lead_identity(lead), lead.get("provider", "freehire"), lead.get("mechanism", "public_api"), lead.get("query", ""), lead.get("title", ""), lead.get("company", ""), lead.get("location", ""), lead.get("posted_date", ""), lead.get("source_url", ""), lead.get("canonical_url", ""), lead.get("category", ""), lead.get("seniority", ""), lead.get("description_available", False), lead.get("prefilter_status", ""), lead.get("prefilter_reason", "")])
    wb["Discovery_Providers"].append(["freehire", "OK", len(_load(evidence_root / "discovery_queries.json", {}).get("queries", [])), len(raw), len(queued), json.dumps(health, default=str)])
    wb["Source_Health"].append(["freehire", "OK", json.dumps(health, sort_keys=True)])
    for result in results:
        source_health = result.get("source_health", {})
        wb["Source_Health"].append([result.get("company", result.get("company_id", "")), source_health.get("status", "usable") if isinstance(source_health, dict) else "usable", json.dumps(source_health, sort_keys=True)])
    candidate_rows = verification or []
    for candidate in candidate_rows:
        wb["Company_Candidates"].append([candidate.get("company", ""), candidate.get("official_domain", ""), candidate.get("category", ""), candidate.get("discovery_query", "job-first"), candidate.get("strongest_job", ""), candidate.get("india_evidence", ""), candidate.get("confidence", ""), candidate.get("freshness", ""), candidate.get("status", "")])
        wb["Official_Verification"].append([candidate.get("company", ""), candidate.get("company_id", ""), candidate.get("official_domain", ""), candidate.get("careers_url", ""), candidate.get("verified", True), candidate.get("verified_job_count", 0), candidate.get("strongest_job", ""), candidate.get("india_evidence", ""), candidate.get("role_stack_evidence", ""), candidate.get("experience_evidence", ""), json.dumps(candidate.get("source_health", {}), sort_keys=True), candidate.get("selected", False), candidate.get("selection_reason", ""), candidate.get("deferral_reason", "")])
    deduped_count = len({_lead_identity(lead) for lead in raw})
    validated_count = validated_total
    rejected_count = sum(len(r.get("rejections", [])) for r in results)
    foreign_count = sum(len(r.get("foreign_leads", [])) for r in results)
    wb["Run_Summary"].append([run_id, len(raw), deduped_count, len(queued), len(results), validated_count, rejected_count, foreign_count, len(results)])
    problems = reconcile_metrics(raw=len(raw), deduped=len({_lead_identity(lead) for lead in raw}), queued=len(queued), verified_companies=len(candidate_rows), selected_companies=sum(1 for candidate in candidate_rows if candidate.get("selected")), primary_attempts=sum(1 for result in all_attempt_results if "correction" not in str(result.get("attempt_id", "")).lower()), correction_attempts=sum(1 for result in all_attempt_results if "correction" in str(result.get("attempt_id", "")).lower()), total_attempts=len(all_attempt_results), validated=validated_count, rejected=rejected_count, foreign=foreign_count)
    if problems:
        raise ValueError("discovery report reconciliation failed: " + "; ".join(problems))
    wb.save(output)
    load_workbook(output).close()
    return output
