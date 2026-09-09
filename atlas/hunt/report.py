"""Unique per-run Hunt report writer (build spec 18, 25; architecture s.12).

Writes an immutable run directory:
``output/production/runs/<RUN_ID>/Atlas_Jobs_<YYYYMMDD-HHMMSS>_<RUN_ID>.xlsx``
plus JSON/JSONL side artifacts, and appends one line to
``output/production/run_index.jsonl``. No earlier workbook is overwritten; the
workbook is written atomically and reopen-validated. ``Atlas_Jobs_LATEST.xlsx``
is never touched here (recovery build).
"""

from __future__ import annotations

import datetime
import json
import os
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Optional, Sequence

from openpyxl import Workbook, load_workbook

from atlas.hunt.campaign import CompanyCampaign
from atlas.hunt.experience_v2 import evaluate_experience_fit
from atlas.hunt.matching import CandidateMatchDecision
from atlas.hunt.models import RunLineage, as_dict
from atlas.hunt.pipeline import HuntPipelineResult, JobEvaluation
from atlas.hunt.qualification import QualificationStatus
from atlas.hunt.role_intent import RoleIntentPolicy
from atlas.policy.loader import PolicyBundle

__all__ = ["HuntReport", "write_hunt_report", "REQUIRED_HUNT_SHEETS", "append_run_index"]

REQUIRED_HUNT_SHEETS = (
    "All_Jobs",
    "New_Companies",
    "Company_Coverage",
    "Source_Coverage",
    "Closed_or_Rejected",
    "Resume_Tailoring",
    "Recruiter_Contacts",
    "Run_Summary",
)

ALL_JOBS_COLUMNS = (
    "Record_Class", "Company", "Role_Title", "Location", "Lane", "Role_Family",
    "Stack_Anchors", "Qualification_Status", "Qualification_Reasons",
    "Experience_Min", "Experience_Max", "Experience_Preferred", "Experience_Fit",
    "Work_Mode", "Primary_Source", "Discovery_Channels", "Official_Apply_URL",
    "Official_Requisition_ID", "Freshness_Band", "Verification_Level",
    "Match_Score", "Requirements_Matched", "Missing_Requirements", "Recommendation",
)

COMPANY_COVERAGE_COLUMNS = (
    "Company", "Tier", "Group", "Official_Careers_Domain", "Source", "Route",
    "Check_Type", "Lane", "Snapshot_ID", "Pages", "Raw_Jobs", "Prefiltered_Jobs",
    "Hydrated_Jobs", "Qualified_Jobs", "Wrong_Stack_Rejected", "Role_Family_Rejected",
    "Experience_Rejected", "Location_Rejected", "Freshness_Rejected", "Manual_Verification",
    "Access_Status", "Terminal_Status", "Checked_At", "Next_Check",
)

SOURCE_COVERAGE_COLUMNS = (
    "Source_Instance", "Source_Family", "Route", "Health", "Pages", "Requests",
    "Companies", "Raw_Jobs", "Qualified_Jobs", "Limitation", "Retry_State",
    "Adapter_Version", "Parser_Version",
)

CLOSED_REJECTED_COLUMNS = (
    "Company", "Role_Title", "Lane", "Reason_Code", "Dominant_Stack",
    "Role_Family", "Evidence_Summary",
)

RESUME_TAILORING_COLUMNS = (
    "Company", "Role_Title", "Lane", "Recommendation", "Match_Score",
    "Requirements_Matched", "Missing_Requirements", "Official_Apply_URL",
)

# New companies discovered/added during the campaign (V3 dynamic-expansion
# semantics): companies not in the base seed, or added by an extension batch.
NEW_COMPANIES_COLUMNS = (
    "Company", "Group", "Origin", "Tier", "Batch_Index", "Sealed_At", "Provenance",
)

# Recruiter/HR outreach queue (V3 Outreach_Queue). Recruiter outreach is DEFERRED
# in the Search Recovery build (not part of search-quality recovery), so this
# sheet is authored header-only and truthfully carries no fabricated contacts.
RECRUITER_CONTACTS_COLUMNS = (
    "Company", "Role_Title", "Lane", "Contact_Name", "Contact_Role",
    "Contact_Source", "Contact_Confidence", "Outreach_Status", "Notes",
)


@dataclass(frozen=True)
class HuntReport:
    run_id: str
    run_dir: Path
    workbook_path: Path
    all_jobs_rows: int
    company_coverage_rows: int
    source_coverage_rows: int
    closed_rejected_rows: int
    application_packs: int
    outcome: str


def _write_sheet(ws, columns: Sequence[str], rows: Sequence[Sequence]) -> None:
    ws.append(list(columns))
    for r in rows:
        ws.append(["" if v is None else v for v in r])


def _atomic_write_workbook(wb: Workbook, path: Path) -> None:
    if path.exists():
        raise FileExistsError(f"refusing to overwrite existing workbook: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=str(path.parent))
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    # reopen-validate
    check = load_workbook(path, read_only=True)
    missing = [s for s in REQUIRED_HUNT_SHEETS if s not in check.sheetnames]
    check.close()
    if missing:
        raise ValueError(f"workbook missing required sheets: {missing}")


def append_run_index(index_path: Path, record: dict) -> None:
    """Append one JSON line to run_index.jsonl (create-safe)."""
    index_path.parent.mkdir(parents=True, exist_ok=True)
    line = json.dumps(record, sort_keys=True, ensure_ascii=False)
    with open(index_path, "a", encoding="utf-8") as fh:
        fh.write(line + "\n")
        fh.flush()
        os.fsync(fh.fileno())


def _dump_json(path: Path, payload) -> None:
    path.write_text(json.dumps(as_dict(payload), indent=2, ensure_ascii=False, sort_keys=False), encoding="utf-8")


def _dump_jsonl(path: Path, rows) -> None:
    with open(path, "w", encoding="utf-8") as fh:
        for r in rows:
            fh.write(json.dumps(as_dict(r), ensure_ascii=False) + "\n")


def write_hunt_report(
    result: HuntPipelineResult,
    campaign: CompanyCampaign,
    shortlist: Sequence[CandidateMatchDecision],
    lineage: RunLineage,
    intent: RoleIntentPolicy,
    policy: PolicyBundle,
    *,
    run_id: str,
    output_root: Path,
    outcome: str,
    now: Optional[datetime.datetime] = None,
) -> HuntReport:
    now = now or datetime.datetime.now(datetime.timezone.utc)
    stamp = now.strftime("%Y%m%d-%H%M%S")
    run_dir = Path(output_root) / "runs" / run_id
    run_dir.mkdir(parents=True, exist_ok=True)
    # Immutability guard: a run id is written exactly once. A second attempt to
    # publish into an existing run directory is refused deterministically,
    # regardless of sub-second filename timing (build spec 25).
    existing_wb = list(run_dir.glob("Atlas_Jobs_*.xlsx"))
    if existing_wb:
        raise FileExistsError(
            f"run {run_id} already has a published workbook ({existing_wb[0].name}); refusing to overwrite"
        )
    workbook_path = run_dir / f"Atlas_Jobs_{stamp}_{run_id}.xlsx"

    # index qualified evaluations by job key for the shortlist join
    eval_by_key = {
        ev.detail.canonical_key: ev
        for ev in result.evaluations
        if ev.final_status == QualificationStatus.QUALIFIED.value and ev.final_lane
    }

    all_jobs_rows: list[list] = []
    for m in shortlist:
        ev: Optional[JobEvaluation] = eval_by_key.get(m.job_key)
        if ev is None:
            continue
        d = ev.detail
        dec = ev.qualification.by_lane[ev.final_lane]
        fit = evaluate_experience_fit(
            title=d.title, experience_text=d.experience_text or d.description,
            policy=policy.experience, candidate_years=None,
        )
        all_jobs_rows.append([
            d.record_class, d.company, d.title, d.location, m.lane, dec.role_family,
            ", ".join(dec.matched_anchors), "QUALIFIED", "; ".join(dec.reasons),
            fit.mandatory_min_years, fit.mandatory_max_years, fit.preferred_years, m.experience_fit,
            d.work_mode, d.source_family, ", ".join(d.discovery_channels), d.official_url,
            d.requisition_id, m.freshness_band, d.verification_state, m.match_score,
            ", ".join(m.requirements_matched), ", ".join(m.missing_requirements), m.recommendation,
        ])

    # New_Companies: companies not in the base seed (dynamic-expansion provenance)
    # or added by an extension batch (batch_index > 0). Truthful, no fabrication.
    new_company_rows: list[list] = []
    for batch in campaign.batches:
        for c in batch.companies:
            if c.origin != "SEED" or batch.batch_index > 0:
                new_company_rows.append([
                    c.name, c.group, c.origin, c.tier, batch.batch_index,
                    batch.sealed_at, batch.reason,
                ])

    company_rows = [
        [c.company, c.tier, c.group, c.official_domain, c.source, c.route, c.check_type, c.lane,
         c.snapshot_id, c.pages, c.raw_jobs, c.prefiltered_jobs, c.hydrated_jobs, c.qualified_jobs,
         c.wrong_stack_rejected, c.role_family_rejected, c.experience_rejected, c.location_rejected,
         c.freshness_rejected, c.manual_verification, c.access_status, c.terminal_status,
         c.checked_at, c.next_check]
        for c in result.coverage
    ]
    source_rows = [
        [s.source_instance_id, s.source_family, s.route, s.health, s.pages, s.request_count,
         s.companies, s.raw_jobs, s.qualified_jobs, s.limitation, s.retry_state,
         s.adapter_version, s.parser_version]
        for s in result.source_coverage
    ]

    closed_rows: list[list] = []
    for ev in result.evaluations:
        if ev.final_status == QualificationStatus.QUALIFIED.value:
            continue
        dec = ev.qualification.primary or next(iter(ev.qualification.by_lane.values()), None)
        lane = ev.final_lane or (dec.lane if dec else "")
        closed_rows.append([
            ev.detail.company, ev.detail.title, lane, ev.final_status,
            dec.dominant_stack if dec else "", dec.role_family if dec else "",
            (ev.detail.description or "")[:180],
        ])

    resume_rows = [
        [m.company, m.title, m.lane, m.recommendation, m.match_score,
         ", ".join(m.requirements_matched), ", ".join(m.missing_requirements),
         eval_by_key[m.job_key].detail.official_url if m.job_key in eval_by_key else ""]
        for m in shortlist if m.is_apply_family
    ]

    lane_summary = {lane: {"qualified": 0, "relevant": 0} for lane in campaign.lanes}
    for ev in result.evaluations:
        if ev.final_status == QualificationStatus.QUALIFIED.value and ev.final_lane in lane_summary:
            lane_summary[ev.final_lane]["qualified"] += 1
    for m in shortlist:
        if m.lane in lane_summary:
            lane_summary[m.lane]["relevant"] += 1

    summary_rows = [["Outcome", outcome]]
    summary_rows.append(["Companies planned", campaign.company_count])
    summary_rows.append(["Company x lane obligations", campaign.obligation_count])
    summary_rows.append(["Sealed batches", len(campaign.batches)])
    summary_rows.append(["Raw jobs", result.raw_jobs])
    summary_rows.append(["Hydrated jobs", result.hydrated_jobs])
    summary_rows.append(["Qualified jobs", result.qualified_jobs])
    summary_rows.append(["Relevant (shortlist) jobs", len(shortlist)])
    summary_rows.append(["Network calls", result.network_calls])
    for lane in campaign.lanes:
        summary_rows.append([f"Lane {lane}", f"qualified={lane_summary[lane]['qualified']} shortlist={lane_summary[lane]['relevant']}"])

    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb.create_sheet("All_Jobs"), ALL_JOBS_COLUMNS, all_jobs_rows)
    _write_sheet(wb.create_sheet("New_Companies"), NEW_COMPANIES_COLUMNS, new_company_rows)
    _write_sheet(wb.create_sheet("Company_Coverage"), COMPANY_COVERAGE_COLUMNS, company_rows)
    _write_sheet(wb.create_sheet("Source_Coverage"), SOURCE_COVERAGE_COLUMNS, source_rows)
    _write_sheet(wb.create_sheet("Closed_or_Rejected"), CLOSED_REJECTED_COLUMNS, closed_rows)
    _write_sheet(wb.create_sheet("Resume_Tailoring"), RESUME_TAILORING_COLUMNS, resume_rows)
    _write_sheet(wb.create_sheet("Recruiter_Contacts"), RECRUITER_CONTACTS_COLUMNS, [])
    _write_sheet(wb.create_sheet("Run_Summary"), ("Metric", "Value"), summary_rows)
    _atomic_write_workbook(wb, workbook_path)

    # --- JSON/JSONL side artifacts ---
    _dump_json(run_dir / "run_manifest.json", {
        "run_id": run_id, "workbook": workbook_path.name, "outcome": outcome,
        "created_at": now.isoformat(), "all_jobs_rows": len(all_jobs_rows),
        "company_coverage_rows": len(company_rows), "source_coverage_rows": len(source_rows),
    })
    _dump_json(run_dir / "sealed_company_plan.json", {
        "campaign_id": campaign.campaign_id, "company_plan_hash": campaign.company_plan_hash,
        "role_policy_hash": campaign.role_policy_hash, "lanes": list(campaign.lanes),
        "companies": [as_dict(c) for c in campaign.companies],
    })
    _dump_jsonl(run_dir / "sealed_batch_history.jsonl", campaign.batches)
    _dump_json(run_dir / "query_plan.json", {
        "lanes": {k: list(intent.lane(k).query_templates) for k in intent.lane_keys()},
    })
    _dump_json(run_dir / "company_coverage.json", result.coverage)
    _dump_json(run_dir / "source_coverage.json", result.source_coverage)
    _dump_jsonl(run_dir / "raw_observations.jsonl", result.snapshots)
    _dump_jsonl(run_dir / "job_details.jsonl", result.details)
    _dump_jsonl(run_dir / "qualification_decisions.jsonl",
                [ev.qualification.by_lane[ev.final_lane] for ev in result.evaluations
                 if ev.final_status == QualificationStatus.QUALIFIED.value and ev.final_lane])
    _dump_json(run_dir / "recommendations.json", list(shortlist))
    _dump_json(run_dir / "run_lineage.json", lineage)

    append_run_index(Path(output_root) / "run_index.jsonl", {
        "run_id": run_id, "kind": lineage.run_kind, "parent_run_id": lineage.parent_run_id,
        "collection_run_id": lineage.collection_run_id, "outcome": outcome,
        "workbook": str(workbook_path), "role_policy_hash": lineage.role_policy_hash,
        "company_plan_hash": campaign.company_plan_hash, "created_at": now.isoformat(),
        "relevant_jobs": len(shortlist), "qualified_jobs": result.qualified_jobs,
    })

    return HuntReport(
        run_id=run_id, run_dir=run_dir, workbook_path=workbook_path,
        all_jobs_rows=len(all_jobs_rows), company_coverage_rows=len(company_rows),
        source_coverage_rows=len(source_rows), closed_rejected_rows=len(closed_rows),
        application_packs=len(resume_rows), outcome=outcome,
    )
