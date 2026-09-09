"""Hunt product-quality validator (build spec 24).

Fails a run when the search-quality contract is violated: missing/non-terminal
coverage, an absent lane, empty coverage sheets, a candidate-facing row without
the lane's required anchor evidence, a wrong role family in the shortlist, a
diversity-cap breach, non-reconciling counts, a package without an apply-family
recommendation, or a non-unique/mis-named workbook. It never prints private
candidate content.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from atlas.hunt.matching import APPLY_FAMILY
from atlas.hunt.role_family import DEVELOPMENT_FAMILIES
from atlas.hunt.role_intent import RoleIntentPolicy, load_role_intent_policy
from atlas.hunt.signals import signal_present

__all__ = ["ValidationIssue", "ValidationReport", "validate_run_dir"]

_TERMINAL = frozenset({"CHECKED", "ZERO", "ACCESS_LIMITED", "ERROR"})
_WORKBOOK_RE = re.compile(r"^Atlas_Jobs_\d{8}-\d{6}_.+\.xlsx$")


@dataclass
class ValidationIssue:
    code: str
    detail: str
    lane: str = ""
    company: str = ""


@dataclass
class ValidationReport:
    run_id: str
    passed: bool = True
    issues: list[ValidationIssue] = field(default_factory=list)
    all_jobs_rows: int = 0
    company_coverage_rows: int = 0
    source_coverage_rows: int = 0

    def add(self, code: str, detail: str, **kw) -> None:
        self.issues.append(ValidationIssue(code=code, detail=detail, **kw))
        self.passed = False

    def to_dict(self) -> dict:
        return {
            "run_id": self.run_id, "passed": self.passed,
            "all_jobs_rows": self.all_jobs_rows,
            "company_coverage_rows": self.company_coverage_rows,
            "source_coverage_rows": self.source_coverage_rows,
            "issues": [vars(i) for i in self.issues],
        }


def _rows(ws) -> list[dict]:
    it = ws.iter_rows(values_only=True)
    try:
        header = list(next(it))
    except StopIteration:
        return []
    out = []
    for r in it:
        out.append({header[i]: (r[i] if i < len(r) else None) for i in range(len(header))})
    return out


def _lane_anchor_ok(lane: str, anchors: str, intent: RoleIntentPolicy) -> bool:
    text = anchors or ""
    if lane not in intent.lanes:
        return True
    contract = intent.lane(lane)
    if not contract.required_anchor_groups:
        return True
    for grp in contract.required_anchor_groups:
        if not any(signal_present(text, a) for a in grp.any_of):
            return False
    return True


def validate_run_dir(run_dir: Path, *, intent: Optional[RoleIntentPolicy] = None, root: Optional[Path] = None) -> ValidationReport:
    run_dir = Path(run_dir)
    intent = intent or load_role_intent_policy(root=root)
    report = ValidationReport(run_id=run_dir.name)

    workbooks = sorted(run_dir.glob("Atlas_Jobs_*.xlsx"))
    if not workbooks:
        report.add("NO_WORKBOOK", "no Atlas_Jobs_*.xlsx workbook in run directory")
        return report
    wb_path = workbooks[-1]
    if not _WORKBOOK_RE.match(wb_path.name):
        report.add("WORKBOOK_NAME", f"workbook filename not unique/timestamped: {wb_path.name}")

    wb = load_workbook(wb_path, read_only=True)
    expected_sheets = {
        "All_Jobs", "New_Companies", "Company_Coverage", "Source_Coverage",
        "Closed_or_Rejected", "Resume_Tailoring", "Recruiter_Contacts", "Run_Summary",
    }
    for sheet in expected_sheets:
        if sheet not in wb.sheetnames:
            report.add("MISSING_SHEET", f"missing sheet {sheet}")
    extra = set(wb.sheetnames) - expected_sheets
    if extra:
        report.add("EXTRA_SHEET", f"unexpected sheet(s): {sorted(extra)}")
    if not report.passed and any(i.code in ("MISSING_SHEET", "EXTRA_SHEET") for i in report.issues):
        wb.close()
        return report

    all_jobs = _rows(wb["All_Jobs"])
    company_cov = _rows(wb["Company_Coverage"])
    source_cov = _rows(wb["Source_Coverage"])
    resume = _rows(wb["Resume_Tailoring"])
    summary = {r.get("Metric"): r.get("Value") for r in _rows(wb["Run_Summary"])}
    wb.close()

    report.all_jobs_rows = len(all_jobs)
    report.company_coverage_rows = len(company_cov)
    report.source_coverage_rows = len(source_cov)

    # --- coverage presence ---
    if not company_cov:
        report.add("EMPTY_COMPANY_COVERAGE", "Company_Coverage sheet is empty")
    if not source_cov:
        report.add("EMPTY_SOURCE_COVERAGE", "Source_Coverage sheet is empty")

    lanes_seen = {r.get("Lane") for r in company_cov if r.get("Lane")}
    for lane in intent.lane_keys():
        if lane not in lanes_seen:
            report.add("MISSING_LANE", f"lane {lane} absent from Company_Coverage", lane=lane)

    for r in company_cov:
        if r.get("Terminal_Status") not in _TERMINAL:
            report.add("NONTERMINAL_COVERAGE",
                       f"company/lane obligation not terminal: {r.get('Terminal_Status')}",
                       lane=r.get("Lane") or "", company=r.get("Company") or "")

    # --- candidate-facing row integrity ---
    per_company: dict[str, int] = {}
    per_company_lane: dict[tuple, int] = {}
    for r in all_jobs:
        company = r.get("Company") or ""
        lane = r.get("Lane") or ""
        if r.get("Qualification_Status") != "QUALIFIED":
            report.add("UNQUALIFIED_ROW", f"All_Jobs row not QUALIFIED: {company}", lane=lane, company=company)
        fam = r.get("Role_Family") or ""
        if fam and fam not in DEVELOPMENT_FAMILIES:
            report.add("WRONG_ROLE_FAMILY", f"non-development role family {fam} in shortlist", lane=lane, company=company)
        if not _lane_anchor_ok(lane, r.get("Stack_Anchors") or "", intent):
            report.add("MISSING_LANE_ANCHOR",
                       f"{lane} row lacks required anchor evidence (anchors={r.get('Stack_Anchors')!r})",
                       lane=lane, company=company)
        per_company[company] = per_company.get(company, 0) + 1
        per_company_lane[(company, lane)] = per_company_lane.get((company, lane), 0) + 1

    for company, n in per_company.items():
        if n > intent.max_display_per_company:
            report.add("COMPANY_CAP", f"company {company} has {n} shortlist rows (> {intent.max_display_per_company})", company=company)
    for (company, lane), n in per_company_lane.items():
        if n > intent.max_display_per_lane_per_company:
            report.add("COMPANY_LANE_CAP", f"{company}/{lane} has {n} rows (> {intent.max_display_per_lane_per_company})", company=company, lane=lane)
    if len(all_jobs) > intent.max_shortlist_total:
        report.add("SHORTLIST_CAP", f"shortlist has {len(all_jobs)} rows (> {intent.max_shortlist_total})")

    # --- packages only for apply-family ---
    for r in resume:
        if (r.get("Recommendation") or "") not in APPLY_FAMILY:
            report.add("NON_APPLY_PACKAGE", f"Resume_Tailoring row is not apply-family: {r.get('Recommendation')}",
                       company=r.get("Company") or "")

    # --- reconcile ---
    rel = summary.get("Relevant (shortlist) jobs")
    try:
        if rel is not None and int(rel) != len(all_jobs):
            report.add("COUNT_MISMATCH", f"Run_Summary relevant={rel} != All_Jobs rows={len(all_jobs)}")
    except (TypeError, ValueError):
        pass

    return report
