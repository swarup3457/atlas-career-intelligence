"""Immutable agentic-pilot output harness (V4 architecture s.10, prompt s.14).

Writes ONE unique, immutable run under
``output/production/agentic_pilots/<RUN_ID>/`` with the ten required artifacts
and the eight-sheet workbook. It NEVER updates ``output/production/latest`` and
refuses to overwrite an existing run directory (prompt s.1, s.14).

This is a pure structured writer: the governor / demonstration driver assembles
the row dicts (from the deterministic gates + typed evidence) and hands them
here, so the contract is testable offline without a live run.
"""

from __future__ import annotations

import datetime
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from openpyxl import Workbook

__all__ = ["AgenticReport", "write_agentic_report", "ACCEPTED_COLUMNS"]

#: Accepted-India-job columns (prompt s.14). Every accepted row carries them.
ACCEPTED_COLUMNS = [
    "Company", "Role", "Official URL", "Location", "Lane", "Role family",
    "Mandatory backend", "Frontend stack", "Mandatory total experience",
    "Preferred experience", "India decision", "Location evidence", "Stack evidence",
    "Experience evidence", "Requirements matched", "Requirements missing",
    "Recommendation", "Search model", "Semantic-review model", "Company task ID",
]

_COVERAGE_COLUMNS = [
    "Company", "Official domain", "Career entry URL", "Source/route", "Search status",
    "Genuinely searched", "Lanes accounted", "Queries", "Details opened",
    "Accepted", "Rejected", "Foreign", "Limitations",
]
_REJECTED_COLUMNS = ["Company", "Role", "Lane", "Reason", "Detail", "Location", "URL"]
_FOREIGN_COLUMNS = ["Company", "Role", "Location", "Geography decision", "URL", "Reason"]
_SOURCE_HEALTH_COLUMNS = [
    "Company", "Route", "Trusted hosts", "Observations", "Actions", "Details opened",
    "Redirect chain", "Diagnostics",
]
_USAGE_COLUMNS = [
    "Scope", "Model", "AI credits", "Input tokens", "Cached tokens", "Output tokens",
    "Reasoning tokens", "Tool calls", "Web searches", "Web fetches",
    "Browser observations", "Details opened",
]


@dataclass
class AgenticReport:
    run_id: str
    run_dir: Path
    workbook_path: Path
    outcome: str
    accepted_rows: int = 0
    rejected_rows: int = 0
    foreign_rows: int = 0
    genuinely_searched: int = 0
    artifacts: list[str] = field(default_factory=list)


def _sheet(wb: Workbook, title: str, columns: list[str], rows: list[dict]):
    ws = wb.create_sheet(title=title)
    ws.append(columns)
    for r in rows:
        ws.append([_cell(r.get(c, "")) for c in columns])
    return ws


def _cell(v) -> object:
    if v is None:
        return ""
    if isinstance(v, (list, tuple)):
        return "; ".join(str(x) for x in v)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return v


def write_agentic_report(
    *,
    run_id: str,
    output_root: Path,
    outcome: str,
    exec_summary: dict,
    company_coverage: list[dict],
    accepted: list[dict],
    rejected: list[dict],
    foreign: list[dict],
    source_health: list[dict],
    usage: dict,
    benchmark_comparison: dict,
    browser_recipes: Optional[list[dict]] = None,
    company_results: Optional[dict] = None,
    candidate_provenance: Optional[dict] = None,
    run_manifest_extra: Optional[dict] = None,
    timestamp: Optional[str] = None,
) -> AgenticReport:
    output_root = Path(output_root)
    run_dir = output_root / "agentic_pilots" / run_id
    if run_dir.exists():
        raise FileExistsError(f"refusing to overwrite existing agentic run dir: {run_dir}")
    run_dir.mkdir(parents=True, exist_ok=False)
    ts = timestamp or datetime.datetime.now().strftime("%Y%m%d-%H%M%S")

    # -- JSON artifacts ------------------------------------------------------
    (run_dir / "jobs_accepted.json").write_text(json.dumps(accepted, indent=2, ensure_ascii=False), encoding="utf-8")
    (run_dir / "jobs_rejected.json").write_text(json.dumps(rejected, indent=2, ensure_ascii=False), encoding="utf-8")
    (run_dir / "foreign_leads.json").write_text(json.dumps(foreign, indent=2, ensure_ascii=False), encoding="utf-8")
    (run_dir / "coverage.json").write_text(json.dumps(company_coverage, indent=2, ensure_ascii=False), encoding="utf-8")
    (run_dir / "browser_recipes.json").write_text(
        json.dumps(browser_recipes or [], indent=2, ensure_ascii=False), encoding="utf-8")
    (run_dir / "llm_usage.json").write_text(json.dumps(usage, indent=2, ensure_ascii=False), encoding="utf-8")
    (run_dir / "benchmark_comparison.json").write_text(
        json.dumps(benchmark_comparison, indent=2, ensure_ascii=False), encoding="utf-8")

    cr_dir = run_dir / "company_results"
    cr_dir.mkdir(exist_ok=True)
    for company, result in (company_results or {}).items():
        slug = "".join(ch if ch.isalnum() else "_" for ch in company.lower()).strip("_")
        (cr_dir / f"{slug}.json").write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")

    # -- workbook (8 sheets) -------------------------------------------------
    wb = Workbook()
    wb.remove(wb.active)
    es = wb.create_sheet("Executive_Summary")
    es.append(["Metric", "Value"])
    for k, v in exec_summary.items():
        es.append([k, _cell(v)])
    _sheet(wb, "Company_Coverage", _COVERAGE_COLUMNS, company_coverage)
    _sheet(wb, "Accepted_India_Jobs", ACCEPTED_COLUMNS, accepted)
    _sheet(wb, "Rejected_Jobs", _REJECTED_COLUMNS, rejected)
    _sheet(wb, "Foreign_Leads", _FOREIGN_COLUMNS, foreign)
    _sheet(wb, "Browser_Source_Health", _SOURCE_HEALTH_COLUMNS, source_health)
    _sheet(wb, "LLM_Usage", _USAGE_COLUMNS, usage.get("rows", []) if isinstance(usage, dict) else [])
    bc = wb.create_sheet("Benchmark_Comparison")
    bc.append(["Metric", "Native", "Atlas V4", "Note"])
    for row in benchmark_comparison.get("rows", []):
        bc.append([row.get("metric", ""), _cell(row.get("native", "")),
                   _cell(row.get("atlas_v4", "")), _cell(row.get("note", ""))])

    workbook_path = run_dir / f"Atlas_Agentic_India_Pilot_{ts}_{run_id}.xlsx"
    wb.save(workbook_path)

    # -- run manifest --------------------------------------------------------
    genuinely_searched = sum(1 for c in company_coverage if c.get("Genuinely searched") in (True, "True", "Yes"))
    manifest = {
        "run_id": run_id,
        "kind": "AGENTIC_V4",
        "outcome": outcome,
        "created_at": datetime.datetime.now(datetime.timezone.utc).isoformat(),
        "workbook": workbook_path.name,
        "companies": [c.get("Company") for c in company_coverage],
        "counts": {
            "accepted": len(accepted), "rejected": len(rejected), "foreign": len(foreign),
            "genuinely_searched": genuinely_searched, "companies": len(company_coverage),
        },
        "candidate_provenance": candidate_provenance or {},
        "update_latest": False,
        "artifacts": [
            "jobs_accepted.json", "jobs_rejected.json", "foreign_leads.json", "coverage.json",
            "browser_recipes.json", "llm_usage.json", "benchmark_comparison.json",
            "company_results/", workbook_path.name,
        ],
    }
    if run_manifest_extra:
        manifest.update(run_manifest_extra)
    (run_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")

    return AgenticReport(
        run_id=run_id, run_dir=run_dir, workbook_path=workbook_path, outcome=outcome,
        accepted_rows=len(accepted), rejected_rows=len(rejected), foreign_rows=len(foreign),
        genuinely_searched=genuinely_searched,
        artifacts=[p.name for p in run_dir.iterdir()],
    )
