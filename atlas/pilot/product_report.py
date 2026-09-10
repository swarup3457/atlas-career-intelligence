"""Immutable eight-sheet PRODUCT pilot workbook + sealed selection manifest (prompt s.9).

Writes ``output/production/product_pilots/<RUN_ID>/`` with the eight required sheets
(``Validated_Jobs``, ``Rejected_Jobs``, ``Foreign_Leads``, ``Company_Coverage``,
``Source_Coverage``, ``Evidence_Audit``, ``Usage``, ``Selection_Audit``), reuses the
deterministic :mod:`atlas.pilot.evaluate` gates for accepted/rejected/foreign routing, and
hashes every run artifact. ``Company_Coverage`` always contains all five companies, including
zero-match and blocked ones. It NEVER overwrites a prior workbook and NEVER updates ``latest``.
"""

from __future__ import annotations

import datetime
import hashlib
import json
import os
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Optional, Sequence

from openpyxl import Workbook, load_workbook

from atlas.pilot.evaluate import AcceptedJob, PilotEvaluation
from atlas.pilot.models import CompanySearchResult, SEARCHED_TERMINAL

REQUIRED_SHEETS = (
    "Validated_Jobs", "Rejected_Jobs", "Foreign_Leads", "Company_Coverage",
    "Source_Coverage", "Evidence_Audit", "Usage", "Selection_Audit",
)

VALIDATED_COLUMNS = (
    "Company", "Role_Title", "Location", "Lane", "Role_Family", "Stack_Anchors",
    "Experience_Fit", "Work_Mode", "Source_Family", "Official_Apply_URL", "Requisition_ID",
    "Posted_Date", "Updated_Date", "Geography_Decision", "Location_Evidence",
    "Supported_Stack_Evidence", "Experience_Decision", "Requirement_Evidence",
    "Match_Score", "Recommendation", "LLM_Search_Model", "Company_Search_Task_ID",
)
REJECTED_COLUMNS = ("Company", "Role_Title", "Lane", "Reason_Code", "Location", "Detail", "URL")
FOREIGN_COLUMNS = ("Company", "Role_Title", "Location", "Geography_Decision", "Reason", "URL")
COMPANY_COVERAGE_COLUMNS = (
    "Company", "Official_Domain", "Career_Entry_URL", "Route", "Source_Family", "Terminal_Status",
    "Primary_Lanes_Attempted", "Queries_Attempted", "Result_States_Observed", "Details_Opened",
    "Accepted", "Rejected", "Foreign", "Internal_Errors", "External_Errors",
    "Model", "Elapsed_Seconds", "Corrected_Credits", "Recipe_Used_Or_Proposed", "Retry_Count",
)
SOURCE_COVERAGE_COLUMNS = ("Source_Family", "Route", "Companies", "Jobs_Accepted", "Statuses")
EVIDENCE_AUDIT_COLUMNS = (
    "Company", "Role_Title", "Official_URL", "Requisition_ID", "Source_Family",
    "Verification_Level", "Evidence_Snippets",
)
USAGE_COLUMNS = ("Metric", "Value")
SELECTION_AUDIT_COLUMNS = (
    "Company", "Category", "Source_Family", "Bucket", "Score", "Lane_Affinity",
    "Career_Entry_URL", "Reason",
)

# blocker statuses that represent an external limitation (not an internal failure)
EXTERNAL_BLOCKERS = frozenset({"ACCESS_LIMITED", "AUTH_REQUIRED", "UNSUPPORTED_SITE", "NETWORK_UNAVAILABLE"})
INTERNAL_FAILURES = frozenset({"FAILED", "OFFICIAL_SOURCE_UNRESOLVED"})


@dataclass(frozen=True)
class ProductPilotReport:
    run_id: str
    run_dir: Path
    workbook_path: Path
    all_jobs_rows: int
    company_coverage_rows: int
    foreign_leads: int
    outcome: str
    artifact_hashes: dict


def _slug(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")


def _write_sheet(ws, columns: Sequence[str], rows: Sequence[Sequence]) -> None:
    ws.append(list(columns))
    for r in rows:
        ws.append(["" if v is None else v for v in r])


def _dump(path: Path, payload) -> None:
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False, sort_keys=True), encoding="utf-8")


def seal_selection_manifest(run_dir: Path, selection, product_config, *, run_id: str,
                            now: Optional[datetime.datetime] = None) -> Path:
    """Persist the sealed selection manifest BEFORE the first live browser/network action
    (prompt s.3). Refuses to overwrite an existing sealed manifest."""
    now = now or datetime.datetime.now(datetime.timezone.utc)
    run_dir.mkdir(parents=True, exist_ok=True)
    path = run_dir / "selection_manifest.json"
    if path.exists():
        return path  # already sealed for this run (idempotent on resume)
    payload = {
        "run_id": run_id,
        "sealed_at": now.isoformat(),
        "pilot_name": product_config.pilot_name,
        "selection_mode": selection.mode,
        "selection_seed": selection.seed,
        "seed_material": selection.seed_material,
        "config_source": product_config.source_path,
        "config_sha256": product_config.source_sha256,
        "companies": list(selection.names()),
        "selection": selection.to_dict(),
    }
    _dump(path, payload)
    return path


def _validated_row(a: AcceptedJob) -> list:
    m = a.match
    ev = a.evaluation
    anchors = (", ".join(ev.qualification.by_lane[a.lane].matched_anchors)
               if a.lane in ev.qualification.by_lane else "")
    return [
        a.company, a.title, a.location, a.lane, a.role_family, anchors,
        m.experience_fit, ev.detail.work_mode, ev.detail.source_family, a.official_url,
        a.requisition_id, a.posted_date or "", a.updated_date or "", a.geography_decision,
        a.location_evidence, a.supported_stack_evidence, a.experience_decision,
        a.requirement_evidence, m.match_score, m.recommendation, a.llm_search_model,
        a.company_search_task_id,
    ]


def _result_states(result: CompanySearchResult) -> str:
    states = []
    if result.status in SEARCHED_TERMINAL:
        states.append("SEARCHED")
    if result.status in EXTERNAL_BLOCKERS:
        states.append("EXTERNAL_BLOCK")
    if result.status in INTERNAL_FAILURES:
        states.append("INTERNAL_UNRESOLVED")
    if result.jobs:
        states.append(f"{len(result.jobs)}_DETAILS")
    return ", ".join(states) or result.status


def write_product_pilot_report(
    evaluation: PilotEvaluation,
    results: list[CompanySearchResult],
    usage_snapshot: dict,
    product_config,
    selection,
    *,
    run_id: str,
    output_root: Path,
    outcome: str,
    candidate_provenance: dict,
    elapsed_by_company: Optional[dict] = None,
    now: Optional[datetime.datetime] = None,
) -> ProductPilotReport:
    now = now or datetime.datetime.now(datetime.timezone.utc)
    stamp = now.strftime("%Y%m%d-%H%M%S")
    elapsed_by_company = elapsed_by_company or {}
    run_dir = Path(output_root) / "product_pilots" / run_id
    if run_dir.exists() and list(run_dir.glob("Atlas_Product_Company_5_Pilot_*.xlsx")):
        raise FileExistsError(f"run {run_id} already published; refusing to overwrite")
    run_dir.mkdir(parents=True, exist_ok=True)
    (run_dir / "company_results").mkdir(exist_ok=True)
    workbook_path = run_dir / f"Atlas_Product_Company_5_Pilot_{stamp}_{run_id}.xlsx"

    # per-company accepted/rejected/foreign counts
    by_company = {ce.company: ce for ce in evaluation.companies}

    validated_rows = [_validated_row(a) for a in evaluation.accepted]
    rejected_rows = [
        [ce.company, rj.title, rj.lane, rj.reason_code, rj.location, rj.detail[:200], rj.url]
        for ce in evaluation.companies for rj in ce.rejected
    ]
    foreign_rows = [
        [fl.get("company", ""), fl.get("title", ""), fl.get("location", ""),
         fl.get("geography_decision", ""), fl.get("reason", ""), fl.get("url", "")]
        for fl in evaluation.foreign_leads
    ]

    # Company_Coverage: ALL five companies, including zero-match / blocked (prompt s.9)
    usage_by_task = (usage_snapshot.get("by_child") or usage_snapshot.get("children") or {})
    coverage_rows = []
    for r in results:
        ce = by_company.get(r.company)
        acc = len(ce.accepted) if ce else 0
        rej = len(ce.rejected) if ce else 0
        frn = len(ce.foreign_leads) if ce else 0
        attempted = [l for l in product_config.primary_lanes
                     if r.lanes.get(l) and (r.lanes[l].attempted or r.lanes[l].board_snapshot_evaluated)]
        internal_err = "; ".join(x for x in r.limitations if "fail" in x.lower() or "error" in x.lower()) \
            if r.status in INTERNAL_FAILURES else ""
        external_err = "; ".join(r.limitations) if r.status in EXTERNAL_BLOCKERS else ""
        credits = 0
        child = usage_by_task.get(r.company) if isinstance(usage_by_task, dict) else None
        if isinstance(child, dict):
            credits = child.get("ai_credits", child.get("totals", {}).get("ai_credits", 0))
        coverage_rows.append([
            r.company, r.official_domain, r.career_entry_url, r.route, r.source_family, r.status,
            ", ".join(attempted), "; ".join(r.queries_attempted)[:400], _result_states(r),
            len(r.jobs), acc, rej, frn, internal_err[:300], external_err[:300], r.model,
            round(float(elapsed_by_company.get(r.company, 0.0)), 2), credits,
            (r.route or "NONE"), r.retries,
        ])

    # Source_Coverage
    src: dict[str, dict] = {}
    for r in results:
        fam = r.source_family or "OFFICIAL_CAREERS"
        s = src.setdefault(fam, {"route": r.route, "companies": 0, "jobs": 0, "statuses": set()})
        s["companies"] += 1
        s["statuses"].add(r.status)
    for a in evaluation.accepted:
        fam = a.evaluation.detail.source_family or "OFFICIAL_CAREERS"
        src.setdefault(fam, {"route": "", "companies": 0, "jobs": 0, "statuses": set()})["jobs"] += 1
    source_rows = [[fam, s["route"], s["companies"], s["jobs"], ", ".join(sorted(s["statuses"]))]
                   for fam, s in sorted(src.items())]

    # Evidence_Audit
    evidence_rows = []
    for a in evaluation.accepted:
        evidence_rows.append([
            a.company, a.title, a.official_url, a.requisition_id,
            a.evaluation.detail.source_family, a.evaluation.detail.verification_state,
            " || ".join(a.evaluation.detail.evidence_texts)[:600],
        ])
    for r in results:
        for url in r.evidence_urls:
            evidence_rows.append([r.company, "(company evidence)", url, "", r.source_family, r.status, ""])

    # Usage (corrected top-level aggregate only — never summed hierarchical duplicates)
    totals = usage_snapshot.get("totals", {})
    usage_rows = [
        ["Models", ", ".join(totals.get("models", []))],
        ["Input tokens", totals.get("input_tokens", 0)],
        ["Output tokens", totals.get("output_tokens", 0)],
        ["AI credits (corrected)", totals.get("ai_credits", 0)],
        ["Tool calls", totals.get("tool_calls", 0)],
        ["Companies", len(results)],
    ]
    for model, mv in sorted((usage_snapshot.get("by_model") or {}).items()):
        usage_rows.append([f"by_model:{model}:ai_credits", mv.get("ai_credits", 0)])

    # Selection_Audit
    selection_rows = [
        [c.company.name, c.company.category, c.company.resolved_source_family, c.bucket,
         round(c.score, 3), ", ".join(c.company.lane_affinity), c.company.career_entry_url, c.reason]
        for c in selection.selected
    ]

    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb.create_sheet("Validated_Jobs"), VALIDATED_COLUMNS, validated_rows)
    _write_sheet(wb.create_sheet("Rejected_Jobs"), REJECTED_COLUMNS, rejected_rows)
    _write_sheet(wb.create_sheet("Foreign_Leads"), FOREIGN_COLUMNS, foreign_rows)
    _write_sheet(wb.create_sheet("Company_Coverage"), COMPANY_COVERAGE_COLUMNS, coverage_rows)
    _write_sheet(wb.create_sheet("Source_Coverage"), SOURCE_COVERAGE_COLUMNS, source_rows)
    _write_sheet(wb.create_sheet("Evidence_Audit"), EVIDENCE_AUDIT_COLUMNS, evidence_rows)
    _write_sheet(wb.create_sheet("Usage"), USAGE_COLUMNS, usage_rows)
    _write_sheet(wb.create_sheet("Selection_Audit"), SELECTION_AUDIT_COLUMNS, selection_rows)

    if workbook_path.exists():
        raise FileExistsError(f"refusing to overwrite existing workbook: {workbook_path}")
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=str(run_dir))
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, workbook_path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)

    # reopen with openpyxl and verify sheets (prompt s.9)
    check = load_workbook(workbook_path, read_only=True)
    missing = [s for s in REQUIRED_SHEETS if s not in check.sheetnames]
    check.close()
    if missing:
        raise ValueError(f"product workbook missing required sheets: {missing}")

    # --- JSON side artifacts ---
    for r in results:
        _dump(run_dir / "company_results" / f"{_slug(r.company)}.json", r.to_dict())
    _dump(run_dir / "jobs_validated.json", [
        {"company": a.company, "title": a.title, "location": a.location, "lane": a.lane,
         "official_url": a.official_url, "requisition_id": a.requisition_id,
         "geography_decision": a.geography_decision, "experience_decision": a.experience_decision,
         "match_score": a.match.match_score, "recommendation": a.match.recommendation,
         "llm_search_model": a.llm_search_model} for a in evaluation.accepted])
    _dump(run_dir / "jobs_rejected.json", [rj.to_dict() for ce in evaluation.companies for rj in ce.rejected])
    _dump(run_dir / "foreign_leads.json", evaluation.foreign_leads)
    _dump(run_dir / "usage.json", usage_snapshot)
    _dump(run_dir / "run_manifest.json", {
        "run_id": run_id, "workbook": workbook_path.name, "outcome": outcome,
        "created_at": now.isoformat(), "pilot": product_config.pilot_name,
        "selection_mode": selection.mode, "selection_seed": selection.seed,
        "config_source": product_config.source_path, "config_sha256": product_config.source_sha256,
        "companies_planned": len(product_config.companies),
        "companies_terminal": sum(1 for r in results if r.status),
        "validated_jobs": len(validated_rows), "rejected_jobs": len(rejected_rows),
        "foreign_leads": len(foreign_rows), "candidate_provenance": candidate_provenance,
        "update_latest": product_config.update_latest, "usage_totals": totals,
    })

    # hash every artifact (prompt s.9)
    artifact_hashes: dict[str, str] = {}
    for f in sorted(run_dir.rglob("*")):
        if f.is_file() and f.name != "artifact_hashes.json":
            artifact_hashes[str(f.relative_to(run_dir)).replace("\\", "/")] = \
                hashlib.sha256(f.read_bytes()).hexdigest()
    _dump(run_dir / "artifact_hashes.json", artifact_hashes)

    return ProductPilotReport(
        run_id=run_id, run_dir=run_dir, workbook_path=workbook_path,
        all_jobs_rows=len(validated_rows), company_coverage_rows=len(coverage_rows),
        foreign_leads=len(foreign_rows), outcome=outcome, artifact_hashes=artifact_hashes,
    )


def build_product_report_fn(product_config, selection, *, run_id: str, elapsed_by_company=None):
    """Return a ``report_fn`` callable for the governor's finalize node."""
    def _fn(runtime, evaluation, results, usage_snapshot, outcome, candidate_prov):
        return write_product_pilot_report(
            evaluation, results, usage_snapshot, product_config, selection,
            run_id=run_id, output_root=runtime.output_root, outcome=outcome,
            candidate_provenance=candidate_prov, elapsed_by_company=elapsed_by_company,
        )
    return _fn


__all__ = [
    "write_product_pilot_report", "build_product_report_fn", "seal_selection_manifest",
    "ProductPilotReport", "REQUIRED_SHEETS",
]
