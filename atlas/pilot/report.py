"""Unique, immutable pilot report writer (architecture s.12/13, prompt s.13).

Writes ``output/production/llm_pilots/<RUN_ID>/`` with the eight-sheet business workbook
(``All_Jobs`` carries eight extra audit-evidence columns and only India, target-stack rows),
plus the JSON side artifacts (company_results, jobs_accepted/rejected, foreign_leads,
llm_usage, coverage, run_manifest). It NEVER overwrites an existing workbook and NEVER
updates ``latest``.
"""

from __future__ import annotations

import datetime
import json
import os
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Optional, Sequence

from openpyxl import Workbook, load_workbook

from atlas.pilot.config import PilotConfig
from atlas.pilot.evaluate import AcceptedJob, PilotEvaluation
from atlas.pilot.models import CompanySearchResult

REQUIRED_SHEETS = (
    "All_Jobs", "New_Companies", "Company_Coverage", "Source_Coverage",
    "Closed_or_Rejected", "Resume_Tailoring", "Recruiter_Contacts", "Run_Summary",
)

ALL_JOBS_COLUMNS = (
    "Record_Class", "Company", "Role_Title", "Location", "Lane", "Role_Family",
    "Stack_Anchors", "Qualification_Status", "Experience_Fit", "Work_Mode",
    "Primary_Source", "Official_Apply_URL", "Official_Requisition_ID",
    "Posted_Date", "Updated_Date", "Freshness_Band", "Verification_Level",
    "Match_Score", "Requirements_Matched", "Missing_Requirements", "Recommendation",
    # --- audit-evidence columns (prompt s.13) ---
    "Geography_Decision", "Location_Evidence", "Supported_Stack_Evidence",
    "Unsupported_Mandatory_Backend", "Experience_Decision", "Requirement_Evidence",
    "LLM_Search_Model", "Company_Search_Task_ID",
)

NEW_COMPANIES_COLUMNS = ("Company", "Cohort", "Official_Domain", "Route", "Source_Family", "Status", "Task_ID")
COMPANY_COVERAGE_COLUMNS = (
    "Company", "Official_Domain", "Career_Entry_URL", "Route", "Source_Family", "Company_Status",
    "Lane", "Attempted", "Board_Snapshot_Evaluated", "Queries", "Pages", "Candidates", "Model", "Limitations",
)
SOURCE_COVERAGE_COLUMNS = ("Source_Family", "Route", "Companies", "Jobs_Accepted", "Statuses")
CLOSED_REJECTED_COLUMNS = ("Company", "Role_Title", "Lane", "Reason_Code", "Location", "Detail", "URL")
RESUME_TAILORING_COLUMNS = (
    "Company", "Role_Title", "Lane", "Recommendation", "Match_Score",
    "Requirements_Matched", "Missing_Requirements", "Official_Apply_URL",
)
RECRUITER_CONTACTS_COLUMNS = (
    "Company", "Role_Title", "Lane", "Contact_Name", "Contact_Role", "Outreach_Status", "Notes",
)


@dataclass(frozen=True)
class PilotReport:
    run_id: str
    run_dir: Path
    workbook_path: Path
    all_jobs_rows: int
    company_coverage_rows: int
    foreign_leads: int
    outcome: str


def _slug(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")


def _write_sheet(ws, columns: Sequence[str], rows: Sequence[Sequence]) -> None:
    ws.append(list(columns))
    for r in rows:
        ws.append(["" if v is None else v for v in r])


def _atomic_write_workbook(wb: Workbook, path: Path) -> None:
    if path.exists():
        raise FileExistsError(f"refusing to overwrite existing workbook: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=str(path.parent))
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    check = load_workbook(path, read_only=True)
    missing = [s for s in REQUIRED_SHEETS if s not in check.sheetnames]
    check.close()
    if missing:
        raise ValueError(f"workbook missing required sheets: {missing}")


def _dump(path: Path, payload) -> None:
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")


def _all_jobs_row(a: AcceptedJob) -> list:
    m = a.match
    ev = a.evaluation
    return [
        "OFFICIAL_DIRECT", a.company, a.title, a.location, a.lane, a.role_family,
        ", ".join(ev.qualification.by_lane[a.lane].matched_anchors) if a.lane in ev.qualification.by_lane else "",
        "QUALIFIED", m.experience_fit, ev.detail.work_mode,
        ev.detail.source_family, a.official_url, a.requisition_id,
        a.posted_date or "", a.updated_date or "", m.freshness_band, ev.detail.verification_state,
        m.match_score, ", ".join(m.requirements_matched), ", ".join(m.missing_requirements), m.recommendation,
        a.geography_decision, a.location_evidence, a.supported_stack_evidence,
        a.unsupported_mandatory_backend, a.experience_decision, a.requirement_evidence,
        a.llm_search_model, a.company_search_task_id,
    ]


def write_pilot_report(
    evaluation: PilotEvaluation,
    results: list[CompanySearchResult],
    usage_snapshot: dict,
    config: PilotConfig,
    *,
    run_id: str,
    parent_run_id: Optional[str],
    output_root: Path,
    outcome: str,
    candidate_provenance: dict,
    now: Optional[datetime.datetime] = None,
) -> PilotReport:
    now = now or datetime.datetime.now(datetime.timezone.utc)
    stamp = now.strftime("%Y%m%d-%H%M%S")
    run_dir = Path(output_root) / "llm_pilots" / run_id
    if run_dir.exists() and list(run_dir.glob("Atlas_LLM_India_Pilot_*.xlsx")):
        raise FileExistsError(f"run {run_id} already published; refusing to overwrite")
    run_dir.mkdir(parents=True, exist_ok=True)
    (run_dir / "company_results").mkdir(exist_ok=True)
    workbook_path = run_dir / f"Atlas_LLM_India_Pilot_{stamp}_{run_id}.xlsx"

    # --- All_Jobs (India, target-stack only) ---
    all_jobs_rows = [_all_jobs_row(a) for a in evaluation.accepted]

    new_company_rows = [
        [r.company, "LLM_INDIA_OFFICIAL_SEARCH_10", r.official_domain, r.route, r.source_family, r.status, r.task_id]
        for r in results
    ]

    coverage_rows = []
    for r in results:
        for lane in config.primary_lanes:
            cov = r.lanes.get(lane)
            coverage_rows.append([
                r.company, r.official_domain, r.career_entry_url, r.route, r.source_family, r.status,
                lane, (cov.attempted if cov else False), (cov.board_snapshot_evaluated if cov else False),
                "; ".join(cov.queries) if cov else "", (cov.pages if cov else 0),
                (cov.candidates if cov else 0), r.model, "; ".join(r.limitations)[:400],
            ])

    # Source_Coverage
    src: dict[str, dict] = {}
    for r in results:
        fam = r.source_family or "OFFICIAL_CAREERS"
        s = src.setdefault(fam, {"route": r.route, "companies": 0, "jobs": 0, "statuses": set()})
        s["companies"] += 1
        s["statuses"].add(r.status)
    for a in evaluation.accepted:
        fam = a.evaluation.detail.source_family or "OFFICIAL_CAREERS"
        src.setdefault(fam, {"route": "", "companies": 0, "jobs": 0, "statuses": set()})["jobs"] += 1
    source_rows = [
        [fam, s["route"], s["companies"], s["jobs"], ", ".join(sorted(s["statuses"]))]
        for fam, s in sorted(src.items())
    ]

    closed_rows = [
        [ce.company, rj.title, rj.lane, rj.reason_code, rj.location, rj.detail[:200], rj.url]
        for ce in evaluation.companies
        for rj in ce.rejected
    ]

    resume_rows = [
        [a.company, a.title, a.lane, a.match.recommendation, a.match.match_score,
         ", ".join(a.match.requirements_matched), ", ".join(a.match.missing_requirements), a.official_url]
        for a in evaluation.accepted if a.is_recommended
    ]

    summary_rows = [
        ["Outcome", outcome],
        ["Run ID", run_id],
        ["Parent run ID", parent_run_id or ""],
        ["Companies planned", len(config.companies)],
        ["Companies terminal", sum(1 for r in results if r.status)],
        ["India jobs (All_Jobs)", len(all_jobs_rows)],
        ["Foreign leads (side artifact)", len(evaluation.foreign_leads)],
        ["Secondary-audit (GENERAL_SOFTWARE)", len(evaluation.secondary_audit)],
        ["Rejected jobs", len(evaluation.rejected)],
        ["Models", ", ".join(usage_snapshot.get("totals", {}).get("models", []))],
        ["Input tokens", usage_snapshot.get("totals", {}).get("input_tokens", 0)],
        ["Output tokens", usage_snapshot.get("totals", {}).get("output_tokens", 0)],
        ["AI credits", usage_snapshot.get("totals", {}).get("ai_credits", 0)],
        ["Candidate synthetic", candidate_provenance.get("synthetic", "?")],
        ["Update latest", config.update_latest],
    ]

    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb.create_sheet("All_Jobs"), ALL_JOBS_COLUMNS, all_jobs_rows)
    _write_sheet(wb.create_sheet("New_Companies"), NEW_COMPANIES_COLUMNS, new_company_rows)
    _write_sheet(wb.create_sheet("Company_Coverage"), COMPANY_COVERAGE_COLUMNS, coverage_rows)
    _write_sheet(wb.create_sheet("Source_Coverage"), SOURCE_COVERAGE_COLUMNS, source_rows)
    _write_sheet(wb.create_sheet("Closed_or_Rejected"), CLOSED_REJECTED_COLUMNS, closed_rows)
    _write_sheet(wb.create_sheet("Resume_Tailoring"), RESUME_TAILORING_COLUMNS, resume_rows)
    _write_sheet(wb.create_sheet("Recruiter_Contacts"), RECRUITER_CONTACTS_COLUMNS, [])
    _write_sheet(wb.create_sheet("Run_Summary"), ("Metric", "Value"), summary_rows)
    _atomic_write_workbook(wb, workbook_path)

    # --- JSON side artifacts ---
    for r in results:
        _dump(run_dir / "company_results" / f"{_slug(r.company)}.json", r.to_dict())
    _dump(run_dir / "jobs_accepted.json", [
        {
            "company": a.company, "title": a.title, "location": a.location, "lane": a.lane,
            "official_url": a.official_url, "requisition_id": a.requisition_id,
            "geography_decision": a.geography_decision, "location_evidence": a.location_evidence,
            "supported_stack_evidence": a.supported_stack_evidence,
            "unsupported_mandatory_backend": a.unsupported_mandatory_backend,
            "experience_decision": a.experience_decision, "requirement_evidence": a.requirement_evidence,
            "match_score": a.match.match_score, "recommendation": a.match.recommendation,
            "posted_date": a.posted_date, "updated_date": a.updated_date,
            "llm_search_model": a.llm_search_model, "company_search_task_id": a.company_search_task_id,
        } for a in evaluation.accepted
    ])
    _dump(run_dir / "jobs_rejected.json", [rj.to_dict() for ce in evaluation.companies for rj in ce.rejected])
    _dump(run_dir / "foreign_leads.json", evaluation.foreign_leads)
    _dump(run_dir / "secondary_audit.json", evaluation.secondary_audit)
    _dump(run_dir / "llm_usage.json", usage_snapshot)
    _dump(run_dir / "coverage.json", {
        "companies": [
            {"company": r.company, "status": r.status, "route": r.route, "source_family": r.source_family,
             "official_domain": r.official_domain, "career_entry_url": r.career_entry_url,
             "lanes": {k: v.to_dict() for k, v in r.lanes.items()}, "model": r.model,
             "tool_calls": r.tool_calls, "queries_attempted": r.queries_attempted,
             "pages_or_interactions": r.pages_or_interactions, "limitations": r.limitations}
            for r in results
        ],
    })
    _dump(run_dir / "run_manifest.json", {
        "run_id": run_id, "parent_run_id": parent_run_id, "workbook": workbook_path.name,
        "outcome": outcome, "created_at": now.isoformat(),
        "pilot": config.raw.get("pilot_name", "LLM_INDIA_OFFICIAL_SEARCH_10"),
        "config_source": config.source_path, "config_sha256": config.source_sha256,
        "companies_planned": len(config.companies), "companies_terminal": sum(1 for r in results if r.status),
        "india_jobs": len(all_jobs_rows), "foreign_leads": len(evaluation.foreign_leads),
        "candidate_provenance": candidate_provenance, "update_latest": config.update_latest,
        "usage_totals": usage_snapshot.get("totals", {}),
    })

    return PilotReport(
        run_id=run_id, run_dir=run_dir, workbook_path=workbook_path,
        all_jobs_rows=len(all_jobs_rows), company_coverage_rows=len(coverage_rows),
        foreign_leads=len(evaluation.foreign_leads), outcome=outcome,
    )


__all__ = ["write_pilot_report", "PilotReport", "REQUIRED_SHEETS", "ALL_JOBS_COLUMNS"]
