"""Idempotent historical importer for prior Atlas job workbooks (PRODUCTION R1 §1A).

Reads accepted/validated job sheets from previously generated Atlas workbooks and
records them into the reused canonical-job store as ``REVERIFY_REQUIRED`` — a prior
workbook appearance never proves a job is currently active. Reuses
``canonical_jobs`` / ``job_observations`` / ``status_history`` (no second history
database), preserves original first-seen, is idempotent per source file, and
rejects malformed / blank-URL (Drivetrain-shape) rows into an import audit.
"""

from __future__ import annotations

import datetime
import hashlib
import json
import re
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urlparse

from openpyxl import load_workbook

from atlas.persistence.sqlite import StateStore

from .history import ACTIVE_FAMILY, R1_MARKER, canonical_job_id

# Sheets that hold accepted / validated / active jobs (never Rejected_Jobs).
def _is_job_sheet(name: str) -> bool:
    low = name.strip().lower()
    if low in {"validated_jobs", "accepted_india_jobs", "all_jobs"}:
        return True
    return ("job" in low and ("valid" in low or "accept" in low))


_FIELD_PREDICATES: dict[str, list[Any]] = {
    "company": [lambda h: "company" in h],
    "title": [lambda h: h in ("title", "role", "role_title"), lambda h: "title" in h, lambda h: h == "role"],
    "url": [lambda h: "official" in h and "url" in h, lambda h: "apply_url" in h,
            lambda h: "apply" in h and "url" in h, lambda h: h == "url",
            lambda h: "canonical" in h and "url" in h],
    "location": [lambda h: "location" in h],
    "requisition": [lambda h: "requisition" in h, lambda h: h == "req"],
    "recommendation": [lambda h: "recommendation" in h, lambda h: "decision" in h,
                       lambda h: "verification_level" in h],
    "posted": [lambda h: "posted" in h, lambda h: h == "date"],
}

_TS_IN_NAME = re.compile(r"(20\d{2})[-_]?(\d{2})[-_]?(\d{2})")


def _resolve_columns(headers: list[Any]) -> dict[str, int]:
    norm = [str(h or "").strip().lower() for h in headers]
    cols: dict[str, int] = {}
    for field, predicates in _FIELD_PREDICATES.items():
        for predicate in predicates:
            index = next((i for i, header in enumerate(norm) if predicate(header)), None)
            if index is not None:
                cols[field] = index
                break
    return cols


def _https_ok(url: str) -> bool:
    if not isinstance(url, str) or not url.lower().startswith("https://"):
        return False
    return bool(urlparse(url).netloc)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _first_seen_for(path: Path, posted: str | None) -> str:
    if posted and str(posted).strip():
        return str(posted).strip()
    match = _TS_IN_NAME.search(path.name)
    if match:
        return f"{match.group(1)}-{match.group(2)}-{match.group(3)}"
    return datetime.datetime.fromtimestamp(path.stat().st_mtime, datetime.timezone.utc).isoformat()


def import_historical_workbooks(
    store: StateStore,
    *,
    search_roots: Iterable[Path],
    manifest_path: Path | None = None,
) -> dict[str, Any]:
    """Import prior Atlas workbooks into cumulative history as REVERIFY_REQUIRED."""
    files_examined: list[str] = []
    files_accepted: list[str] = []
    source_sha: dict[str, str] = {}
    rows_examined = rows_imported = duplicates = rows_rejected = 0
    rejection_reasons: dict[str, int] = {}

    def _reject(reason: str) -> None:
        nonlocal rows_rejected
        rows_rejected += 1
        rejection_reasons[reason] = rejection_reasons.get(reason, 0) + 1

    workbooks: list[Path] = []
    for root in search_roots:
        root = Path(root)
        if root.exists():
            workbooks.extend(sorted(root.rglob("*.xlsx")))

    for path in workbooks:
        files_examined.append(str(path))
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception:  # unreadable/locked workbook is skipped, not fatal
            _reject("UNREADABLE_WORKBOOK")
            continue
        job_sheets = [name for name in wb.sheetnames if _is_job_sheet(name)]
        if not job_sheets:
            wb.close()
            continue
        sha = _sha256(path)
        source_sha[str(path)] = sha
        files_accepted.append(str(path))

        for sheet_name in job_sheets:
            rows = list(wb[sheet_name].iter_rows(values_only=True))
            if not rows:
                continue
            cols = _resolve_columns(list(rows[0]))
            for raw in rows[1:]:
                rows_examined += 1
                if not any(cell not in (None, "") for cell in raw):
                    _reject("EMPTY_ROW")
                    continue

                def cell(field: str) -> str:
                    index = cols.get(field)
                    if index is None or index >= len(raw):
                        return ""
                    return str(raw[index] or "").strip()

                company, title, location, url = cell("company"), cell("title"), cell("location"), cell("url")
                if not company:
                    _reject("MISSING_COMPANY"); continue
                if not title:
                    _reject("MISSING_TITLE"); continue
                if not location:
                    _reject("MISSING_LOCATION"); continue
                if not _https_ok(url):
                    _reject("BLANK_OR_NON_HTTPS_URL"); continue

                requisition = cell("requisition")
                cid = canonical_job_id(url, company, requisition)
                first_seen = _first_seen_for(path, cell("posted"))
                existing = store.get_canonical_job(cid)

                if existing is None:
                    store.upsert_canonical_job(
                        cid, company=company, job_id=requisition, role=title, location=location,
                        status="REVERIFY_REQUIRED", first_seen=first_seen, last_seen=first_seen,
                        source_url=url, official_apply_url=url,
                        content_hash=f"import:{sha[:16]}",
                        payload={
                            R1_MARKER: "IMPORTED_HISTORICAL", "prior_decision": cell("recommendation"),
                            "imported_from": path.name, "source_sha256": sha,
                        },
                    )
                    store.record_status_change(
                        history_id=f"import:{sha[:16]}:{cid}:REVERIFY_REQUIRED", canonical_id=cid,
                        to_status="REVERIFY_REQUIRED", from_status=None, reason="historical_import",
                        context={"imported_from": path.name},
                    )
                elif existing["current_status"] in ACTIVE_FAMILY:
                    # A live-verified active job is never downgraded by an old workbook.
                    pass

                added = store.add_observation(
                    observation_id=f"import:{sha[:16]}:{cid}", canonical_id=cid, record_id=path.name,
                    source_file=str(path), sheet_name=sheet_name, discovery_source="historical_import",
                    observed_status="REVERIFY_REQUIRED", observed_at=first_seen, content_hash=sha[:16],
                    provenance={"prior_decision": cell("recommendation"), "source_sha256": sha},
                )
                if added:
                    rows_imported += 1
                else:
                    duplicates += 1
        wb.close()

    manifest = {
        "generated_at_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(),
        "files_examined": len(files_examined),
        "files_accepted": len(files_accepted),
        "rows_examined": rows_examined,
        "rows_imported": rows_imported,
        "duplicates": duplicates,
        "rows_rejected": rows_rejected,
        "rejection_reasons": rejection_reasons,
        "source_file_sha256": source_sha,
    }
    if manifest_path is not None:
        Path(manifest_path).parent.mkdir(parents=True, exist_ok=True)
        Path(manifest_path).write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return manifest
