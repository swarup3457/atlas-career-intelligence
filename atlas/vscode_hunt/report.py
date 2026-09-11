from __future__ import annotations

import json
import datetime
import uuid
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from atlas.reporting.trust_boundary import partition_jobs

SHEETS = ("Validated_Jobs", "Rejected_Jobs", "Foreign_Leads", "Company_Coverage", "Lane_Coverage", "Evidence_Audit", "Session_Audit", "Run_Summary")

# Full V3.2 audit/checkpoint/final workbook: 16 immutable sheets.
V32_SHEETS = (
    "Validated_Jobs", "Rejected_Jobs", "Foreign_Leads", "Raw_Job_Leads",
    "Discovery_Providers", "Company_Candidates", "Official_Verification",
    "Selection_Audit", "Company_Coverage", "Lane_Coverage", "Source_Health",
    "Evidence_Audit", "Session_Audit", "Errors", "Discovery_Funnel", "Run_Summary",
)


def _safe_name(value: str) -> str:
    return "".join("_" if char in '<>:"/\\|?*' else char for char in str(value))


def build_workbook(root: Path, run_id: str, conn: Any) -> Path:
    timestamp = datetime.datetime.now(datetime.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    output = Path(root) / run_id / f"Atlas_VSCode_Live_Product3_{timestamp}_{run_id}.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in SHEETS:
        workbook.create_sheet(sheet)
    tasks = conn.execute("SELECT * FROM vscode_hunt_tasks WHERE run_id=? ORDER BY task_id", (run_id,)).fetchall()
    attempts = conn.execute("SELECT * FROM vscode_hunt_attempts WHERE task_id IN (SELECT task_id FROM vscode_hunt_tasks WHERE run_id=?) ORDER BY created_at", (run_id,)).fetchall()
    workbook["Validated_Jobs"].append(["company", "title", "location", "official_url", "requisition", "role_lane", "stack", "mandatory_experience", "posted_date", "recommendation", "strengths", "gaps", "verification_status"])
    workbook["Rejected_Jobs"].append(["company", "title", "location", "url", "lane", "reason_code", "detail"])
    workbook["Foreign_Leads"].append(["company", "title", "location", "url", "reason"])
    workbook["Company_Coverage"].append(["task_id", "company_id", "company_name", "official_domain", "status", "attempts", "details_opened", "browser_errors"])
    for task in tasks:
        task_attempts = [a for a in attempts if a["task_id"] == task["task_id"]]
        results = []
        for attempt in task_attempts:
            if attempt["result_path"] and Path(attempt["result_path"]).exists():
                results.append(json.loads(Path(attempt["result_path"]).read_text(encoding="utf-8")))
        final = results[-1] if results else {}
        workbook["Company_Coverage"].append([task["task_id"], task["company_id"], task["company_name"], task["official_domain"], task["status"], len(task_attempts), len(final.get("detail_urls", [])), len(final.get("browser_errors", []))])
        validated_jobs, proposal_rejections = partition_jobs(final.get("jobs", []), official_domain=task["official_domain"], company=task["company_name"])
        for vj in validated_jobs:
            ev = vj.evidence
            stack = vj.source.get("stack")
            workbook["Validated_Jobs"].append([task["company_name"], ev.title, ev.location, ev.official_url, ev.requisition_id, vj.lane, ", ".join(stack) if isinstance(stack, list) else (stack or ""), ev.experience_text, ev.posted_date, vj.recommendation, "", "", vj.verification_status])
        for rejection in final.get("rejections", []):
            workbook["Rejected_Jobs"].append([task["company_name"], rejection.get("title", ""), rejection.get("location", ""), rejection.get("url", ""), rejection.get("lane", ""), rejection.get("reason_code", ""), rejection.get("detail", rejection.get("reason", ""))])
        for rej in proposal_rejections:
            workbook["Rejected_Jobs"].append([task["company_name"], rej.title, rej.location, rej.url, rej.lane, rej.reason_code, rej.detail])
        for foreign in final.get("foreign_leads", []):
            workbook["Foreign_Leads"].append([task["company_name"], foreign.get("title", ""), foreign.get("location", ""), foreign.get("url", ""), foreign.get("reason", "NON_INDIA_LOCATION")])
    workbook["Lane_Coverage"].append(["task_id", "company", "lane", "result_state", "terminal"])
    for task in tasks:
        paths = [a["result_path"] for a in attempts if a["task_id"] == task["task_id"] and a["result_path"] and Path(a["result_path"]).exists()]
        final = json.loads(Path(paths[-1]).read_text(encoding="utf-8")) if paths else {}
        states = {str(s.get("lane")): str(s.get("state")) for s in final.get("result_states", []) if isinstance(s, dict)}
        for lane in json.loads(task["lanes_json"]):
            workbook["Lane_Coverage"].append([task["task_id"], task["company_name"], lane, states.get(lane, "NOT_OBSERVED"), task["status"] == "COMPLETE"])
    workbook["Evidence_Audit"].append(["task_id", "attempt_id", "result_hash", "detail_urls", "quotes", "source_health", "browser_errors"])
    for attempt in attempts:
        path = attempt["result_path"]
        result = json.loads(Path(path).read_text(encoding="utf-8")) if path and Path(path).exists() else {}
        workbook["Evidence_Audit"].append([attempt["task_id"], attempt["attempt_id"], attempt["result_hash"], len(result.get("detail_urls", [])), len(result.get("evidence_quotes", [])), json.dumps(result.get("source_health", {})), json.dumps(result.get("browser_errors", []))])
    workbook["Session_Audit"].append(["attempt_id", "task_id", "backend", "attempt_number", "action", "model", "browser_used", "error"])
    for attempt in attempts:
        path = attempt["result_path"]
        result = json.loads(Path(path).read_text(encoding="utf-8")) if path and Path(path).exists() else {}
        errors = result.get("browser_errors", [])
        error_text = "; ".join(error if isinstance(error, str) else json.dumps(error, sort_keys=True) for error in errors)
        workbook["Session_Audit"].append([attempt["attempt_id"], attempt["task_id"], attempt["backend"], attempt["attempt_number"], attempt["status"], result.get("model", "sonnet"), bool(result.get("browser_evidence", True)), error_text])
    workbook["Run_Summary"].append(["run_id", "tasks", "terminal_tasks"])
    terminal = sum(1 for task in tasks if task["status"] == "COMPLETE")
    workbook["Run_Summary"].append([run_id, len(tasks), terminal])
    workbook.save(output)
    load_workbook(output).close()
    return output


def build_v32_workbook(root: Path, run_id: str, conn: Any, *, kind: str = "Audit", sequence: int | None = None, context: dict | None = None) -> Path:
    """Build an immutable 16-sheet V3.2 audit / checkpoint / final workbook.

    Every sealed company and every canonical lane is always represented, even with
    no validated jobs. Never overwrites an existing workbook.
    """
    context = context or {}
    kind = str(kind).capitalize()
    timestamp = datetime.datetime.now(datetime.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    out_dir = Path(root) / "production" / "vscode_discovery_v32" / _safe_name(run_id)
    out_dir.mkdir(parents=True, exist_ok=True)
    if kind == "Checkpoint":
        name = f"Atlas_VSCode_DiscoveryV32_Checkpoint_{int(sequence or 1):03d}_{timestamp}_{run_id}.xlsx"
    else:
        name = f"Atlas_VSCode_DiscoveryV32_{kind}_{timestamp}_{run_id}.xlsx"
    output = out_dir / _safe_name(name)
    if output.exists():  # immutability guard: never overwrite a prior workbook
        output = out_dir / _safe_name(f"{output.stem}_{uuid.uuid4().hex[:6]}.xlsx")

    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in V32_SHEETS:
        workbook.create_sheet(sheet)

    tasks = conn.execute("SELECT * FROM vscode_hunt_tasks WHERE run_id=? ORDER BY company_name", (run_id,)).fetchall()
    attempts = conn.execute("SELECT * FROM vscode_hunt_attempts WHERE task_id IN (SELECT task_id FROM vscode_hunt_tasks WHERE run_id=?) ORDER BY created_at", (run_id,)).fetchall()
    commits = conn.execute("SELECT * FROM vscode_result_commits WHERE run_id=?", (run_id,)).fetchall()
    commit_by_attempt = {row["attempt_id"]: row for row in commits}
    reused = bool(context.get("discovery_reused_from_parent_run", True))

    def _result_for(task_id: str) -> dict:
        paths = [a["result_path"] for a in attempts if a["task_id"] == task_id and a["result_path"] and Path(a["result_path"]).exists()]
        if not paths:
            return {}
        try:
            return json.loads(Path(paths[-1]).read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return {}

    workbook["Validated_Jobs"].append(["company", "title", "location", "official_url", "requisition", "role_lane", "stack", "mandatory_experience", "posted_date", "recommendation", "verification_status"])
    workbook["Rejected_Jobs"].append(["company", "title", "location", "url", "lane", "reason_code", "detail"])
    workbook["Foreign_Leads"].append(["company", "title", "location", "url", "reason"])
    workbook["Raw_Job_Leads"].append(["company", "source", "title", "url", "lane_hint", "discovery_reused_from_parent_run"])
    workbook["Discovery_Providers"].append(["provider", "role", "leads_contributed", "status", "discovery_reused_from_parent_run"])
    workbook["Company_Candidates"].append(["company_id", "company_name", "official_domain", "source", "selected", "discovery_reused_from_parent_run"])
    workbook["Official_Verification"].append(["company_id", "company_name", "official_domain", "careers_url", "verification_status"])
    workbook["Selection_Audit"].append(["company_id", "company_name", "sealed", "reason", "discovery_reused_from_parent_run"])
    workbook["Company_Coverage"].append(["task_id", "company_id", "company_name", "official_domain", "status", "attempts", "details_opened", "browser_errors", "lanes"])
    workbook["Lane_Coverage"].append(["task_id", "company", "lane", "result_state", "terminal"])
    workbook["Source_Health"].append(["company", "source", "status", "detail_pages_usable", "note"])
    workbook["Evidence_Audit"].append(["task_id", "attempt_id", "result_hash", "detail_urls", "quotes", "source_health", "browser_errors"])
    workbook["Session_Audit"].append(["attempt_id", "task_id", "company", "backend", "attempt_number", "status", "commit_id", "error"])
    workbook["Errors"].append(["stage", "company", "classification", "detail", "retryable"])
    workbook["Discovery_Funnel"].append(["stage", "count", "note"])
    workbook["Run_Summary"].append(["key", "value"])

    terminal_states = {"COMPLETE", "EXTERNAL_ACCESS_LIMITED", "NEEDS_REPAIR"}
    for task in tasks:
        result = _result_for(task["task_id"])
        lanes = json.loads(task["lanes_json"])
        task_attempts = [a for a in attempts if a["task_id"] == task["task_id"]]
        browser_errors = result.get("browser_errors", []) if isinstance(result, dict) else []
        workbook["Company_Coverage"].append([task["task_id"], task["company_id"], task["company_name"], task["official_domain"], task["status"], len(task_attempts), len(result.get("detail_urls", [])), len(browser_errors), ", ".join(lanes)])
        workbook["Official_Verification"].append([task["company_id"], task["company_name"], task["official_domain"], task["careers_url"] or "", result.get("verification_status", "PENDING_LIVE_VERIFICATION") if result else "PENDING_LIVE_VERIFICATION"])
        workbook["Selection_Audit"].append([task["company_id"], task["company_name"], True, "sealed_cohort_member", reused])
        workbook["Company_Candidates"].append([task["company_id"], task["company_name"], task["official_domain"], "parent_run", True, reused])
        validated_jobs, proposal_rejections = partition_jobs(result.get("jobs", []), official_domain=task["official_domain"], company=task["company_name"])
        for vj in validated_jobs:
            ev = vj.evidence
            stack = vj.source.get("stack")
            workbook["Validated_Jobs"].append([task["company_name"], ev.title, ev.location, ev.official_url, ev.requisition_id, vj.lane, ", ".join(stack) if isinstance(stack, list) else (stack or ""), ev.experience_text, ev.posted_date, vj.recommendation, vj.verification_status])
        for rej in result.get("rejections", []) or []:
            workbook["Rejected_Jobs"].append([task["company_name"], rej.get("title", ""), rej.get("location", ""), rej.get("url", ""), rej.get("lane", ""), rej.get("reason_code", ""), rej.get("detail", rej.get("reason", ""))])
        for rej in proposal_rejections:
            workbook["Rejected_Jobs"].append([task["company_name"], rej.title, rej.location, rej.url, rej.lane, rej.reason_code, rej.detail])
        for foreign in result.get("foreign_leads", []) or []:
            workbook["Foreign_Leads"].append([task["company_name"], foreign.get("title", ""), foreign.get("location", ""), foreign.get("url", ""), foreign.get("reason", "NON_INDIA_LOCATION")])
        states = {}
        raw_states = result.get("result_states")
        if isinstance(raw_states, list):
            states = {str(s.get("lane")): str(s.get("state")) for s in raw_states if isinstance(s, dict)}
        elif isinstance(raw_states, dict):
            states = {str(k): (v if isinstance(v, str) else json.dumps(v)) for k, v in raw_states.items()}
        for lane in lanes:
            workbook["Lane_Coverage"].append([task["task_id"], task["company_name"], lane, states.get(lane, "NOT_SEARCHED"), task["status"] in terminal_states])
        source_health = result.get("source_health") if isinstance(result, dict) else None
        if source_health:
            workbook["Source_Health"].append([task["company_name"], source_health.get("source", ""), source_health.get("status", ""), bool(source_health.get("detail_pages_usable", False)), json.dumps(source_health)])
        else:
            workbook["Source_Health"].append([task["company_name"], "", "NOT_OBSERVED", False, "no live search performed yet"])

    for attempt in attempts:
        path = attempt["result_path"]
        result = json.loads(Path(path).read_text(encoding="utf-8")) if path and Path(path).exists() else {}
        company = next((t["company_name"] for t in tasks if t["task_id"] == attempt["task_id"]), "")
        commit = commit_by_attempt.get(attempt["attempt_id"])
        errors = result.get("browser_errors", [])
        error_text = "; ".join(error if isinstance(error, str) else json.dumps(error, sort_keys=True) for error in errors)
        workbook["Session_Audit"].append([attempt["attempt_id"], attempt["task_id"], company, attempt["backend"], attempt["attempt_number"], attempt["status"], commit["commit_id"] if commit else "", error_text])
        workbook["Evidence_Audit"].append([attempt["task_id"], attempt["attempt_id"], attempt["result_hash"] or "", len(result.get("detail_urls", [])), len(result.get("evidence_quotes", [])), json.dumps(result.get("source_health", {})), json.dumps(result.get("browser_errors", []))])

    for err in context.get("errors", []) or []:
        workbook["Errors"].append([err.get("stage", ""), err.get("company", ""), err.get("classification", ""), err.get("detail", ""), str(err.get("retryable", ""))])

    workbook["Discovery_Providers"].append(["parent_run_v31", "discovery", 0, "REUSED_FROM_PARENT" if reused else "NONE", reused])
    workbook["Raw_Job_Leads"].append(["(none persisted durably)", "parent_run_v31", "", "", "", reused])
    workbook["Discovery_Funnel"].append(["sealed_companies", len(tasks), "V3.2 recovery cohort"])
    workbook["Discovery_Funnel"].append(["discovery_reused_from_parent_run", int(reused), "parent discovery evidence not persisted durably in state"])
    workbook["Discovery_Funnel"].append(["durable_result_commits", len(commits), "committed worker results"])

    terminal = sum(1 for task in tasks if task["status"] in terminal_states)
    summary = {
        "run_id": run_id,
        "workbook_kind": kind,
        "sheets": len(V32_SHEETS),
        "task_contract_version": context.get("task_contract_version", ""),
        "companies": len(tasks),
        "terminal_tasks": terminal,
        "attempts": len(attempts),
        "durable_result_commits": len(commits),
        "phase_status": context.get("status", ""),
        "discovery_reused_from_parent_run": reused,
        "generated_at_utc": timestamp,
        "note": context.get("note", ""),
    }
    for key, value in summary.items():
        workbook["Run_Summary"].append([key, str(value) if isinstance(value, bool) else value])

    workbook.save(output)
    load_workbook(output).close()
    return output
