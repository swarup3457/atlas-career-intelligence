"""Agentic V4 — immutable output harness contract (prompt s.14, s.20)."""

from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from atlas.pilot.agentic_report import ACCEPTED_COLUMNS, write_agentic_report

_EXPECTED_SHEETS = [
    "Executive_Summary", "Company_Coverage", "Accepted_India_Jobs", "Rejected_Jobs",
    "Foreign_Leads", "Browser_Source_Health", "LLM_Usage", "Benchmark_Comparison",
]
_EXPECTED_ARTIFACTS = {
    "jobs_accepted.json", "jobs_rejected.json", "foreign_leads.json", "coverage.json",
    "browser_recipes.json", "llm_usage.json", "benchmark_comparison.json",
    "run_manifest.json", "company_results",
}


def _write(tmp_path, run_id="AGENTIC_TEST"):
    return write_agentic_report(
        run_id=run_id, output_root=tmp_path / "output" / "production", outcome="PARTIAL",
        exec_summary={"Companies assigned": 10, "Genuinely searched": 0, "Outcome": "PARTIAL"},
        company_coverage=[{"Company": "Fiserv", "Search status": "SEARCHED_COMPLETE_NO_MATCHES",
                           "Genuinely searched": True, "Accepted": 0, "Rejected": 4}],
        accepted=[{c: f"v-{c}" for c in ACCEPTED_COLUMNS}],
        rejected=[{"Company": "Fiserv", "Role": "Solutions Architecture - Advisor II",
                   "Lane": "", "Reason": "REJECT_ROLE_FAMILY", "Detail": "architecture/advisor"}],
        foreign=[{"Company": "Fiserv", "Role": "US role", "Location": "New York", "Geography decision": "FOREIGN_EXCLUDED"}],
        source_health=[{"Company": "Fiserv", "Route": "OFFICIAL_ATS_WORKDAY", "Observations": 3}],
        usage={"rows": [{"Scope": "total", "Model": "deterministic", "AI credits": 0}], "totals": {}},
        benchmark_comparison={"rows": [{"metric": "companies searched", "native": 10, "atlas_v4": 0, "note": "offline"}]},
        browser_recipes=[{"company": "Fiserv", "route": "OFFICIAL_ATS_WORKDAY"}],
        company_results={"Fiserv": {"company": "Fiserv", "status": "SEARCHED_COMPLETE_NO_MATCHES"}},
        candidate_provenance={"synthetic": True},
    )


def test_writes_all_artifacts_and_eight_sheets(tmp_path):
    rep = _write(tmp_path)
    names = {p.name for p in rep.run_dir.iterdir()}
    assert _EXPECTED_ARTIFACTS <= names
    assert rep.workbook_path.exists()
    wb = load_workbook(rep.workbook_path)
    assert wb.sheetnames == _EXPECTED_SHEETS
    # accepted sheet carries every required column
    acc = wb["Accepted_India_Jobs"]
    header = [c.value for c in next(acc.iter_rows(max_row=1))]
    assert header == ACCEPTED_COLUMNS


def test_company_results_written(tmp_path):
    rep = _write(tmp_path)
    cr = rep.run_dir / "company_results" / "fiserv.json"
    assert cr.exists()
    assert json.loads(cr.read_text())["status"] == "SEARCHED_COMPLETE_NO_MATCHES"


def test_manifest_does_not_update_latest(tmp_path):
    rep = _write(tmp_path)
    manifest = json.loads((rep.run_dir / "run_manifest.json").read_text())
    assert manifest["update_latest"] is False
    assert manifest["kind"] == "AGENTIC_V4"
    # no 'latest' directory is created anywhere under output/production
    assert not (tmp_path / "output" / "production" / "latest").exists()


def test_refuses_to_overwrite_existing_run(tmp_path):
    _write(tmp_path, run_id="DUP")
    with pytest.raises(FileExistsError):
        _write(tmp_path, run_id="DUP")
