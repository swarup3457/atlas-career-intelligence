from __future__ import annotations

import json
from pathlib import Path

from atlas.discovery.models import JobLead
from atlas.discovery.report import build_discovery_workbook, recommendation_tier, reconcile_metrics


def test_identity_precedence_and_cross_provider_dedup() -> None:
    canonical = JobLead("a", "api", "q", "Role", "Co", "India", "2026-09-10", "https://source/1", canonical_url="https://Official/Job/1", provider_id="x")
    provider = JobLead("freehire", "api", "q", "Role", "Co", "India", "2026-09-10", "https://source/1", provider_id="x")
    fallback = JobLead("a", "api", "q", "Role", "Co", "India", "2026-09-10", "https://source/1")
    assert canonical.identity == "url:https://official/job/1"
    assert provider.identity == "provider:freehire:x"
    assert fallback.identity.startswith("lead:")
    assert canonical.to_record()["identity"] == canonical.identity


def test_recommendation_tiers_preserve_hard_experience_rejection() -> None:
    assert recommendation_tier({"lane":"JAVA_BACKEND", "stack":["Java","Spring"], "experience_text":"2+ years"}) == "STRONG_MATCH"
    assert recommendation_tier({"lane":"JAVA_BACKEND", "stack":["Java"], "experience_text":"3-4 years"}) in {"STRONG_MATCH", "GOOD_MATCH"}
    assert recommendation_tier({"lane":"JAVA_BACKEND", "stack":["Java"], "experience_text":"6+ years"}) == "STRETCH"


def test_reconciliation_invariants() -> None:
    assert reconcile_metrics(raw=160, deduped=157, queued=135, verified_companies=6, selected_companies=3, primary_attempts=3, correction_attempts=2, total_attempts=5, validated=4, rejected=8, foreign=1) == []
    assert reconcile_metrics(raw=1, deduped=2, queued=1, verified_companies=1, selected_companies=2, primary_attempts=1, correction_attempts=0, total_attempts=2, validated=0, rejected=0, foreign=0)


def test_discovery_workbook_metrics_and_attempts(tmp_path: Path) -> None:
    evidence = tmp_path / "evidence"
    live = tmp_path / "live" / "run"
    evidence.mkdir(); live.mkdir(parents=True)
    raw = [{"provider":"freehire","mechanism":"public_api","title":f"Java {index}","company":"Co","location":"Bengaluru, India","posted_date":"2026-09-10","source_url":f"https://jobs/{index}","canonical_url":f"https://jobs/{index}","provider_id":f"p{index}","description_available":True,"prefilter_status":"QUEUED"} for index in range(157)]
    raw.extend(raw[:3])
    (evidence / "freehire_raw.json").write_text(json.dumps(raw), encoding="utf-8")
    (evidence / "freehire_queued.json").write_text(json.dumps(raw[:135]), encoding="utf-8")
    (evidence / "freehire_health.json").write_text("[]", encoding="utf-8")
    (evidence / "discovery_queries.json").write_text(json.dumps({"queries":["java"]}), encoding="utf-8")
    (evidence / "official_verification.json").write_text("[]", encoding="utf-8")
    result = {"task_id":"run::co","attempt_id":"correction-1","company":"Co","company_id":"co","official_domain":"co.com","jobs":[],"rejections":[],"foreign_leads":[],"detail_urls":[],"evidence_quotes":[],"source_health":{},"result_states":{},"completion_claim":True}
    (live / "result.json").write_text(json.dumps({**result, "attempt_id": "primary-1"}), encoding="utf-8")
    (live / "correction-result.json").write_text(json.dumps(result), encoding="utf-8")
    output = build_discovery_workbook(evidence_root=evidence, live_root=live, run_id="run", output_root=tmp_path / "out")
    from openpyxl import load_workbook
    wb = load_workbook(output, read_only=True)
    summary = list(wb["Run_Summary"].iter_rows(min_row=2, values_only=True))[0]
    assert summary[1:4] == (160, 157, 135)
    assert len(list(wb["Session_Audit"].iter_rows())) == 3
    wb.close()
