"""Phase 0.9 — end-to-end pipeline test (creates final outputs + timings).

Runs the callable ``run_phase09`` pipeline into a temp output dir (with a
small expansion set for speed), and asserts it produces the atomic XLSX +
JSON + audit artifacts, leaves the immutable fixture byte-identical, and
records real per-stage timings. Fully offline, NullController.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from atlas.controllers.base import NullController
from atlas.data_integrity.pipeline import run_phase09
from atlas.data_integrity.report import REPORT_SHEETS
from tests.conftest import EXPECTED_FIXTURE_SHA256

pytestmark = [pytest.mark.integration]


def test_run_phase09_creates_outputs_and_preserves_fixture(real_fixture, tmp_path):
    result = run_phase09(
        real_fixture,
        tmp_path / "phase09",
        controller=NullController(),
        expansion_sizes=(500,),  # keep this default-suite test fast
    )

    # 1) fixture never mutated
    assert result.fixture_unchanged
    assert result.source_sha256_before == EXPECTED_FIXTURE_SHA256
    assert result.source_sha256_after == EXPECTED_FIXTURE_SHA256

    # 2) all three artifacts exist at the returned paths
    xlsx = Path(result.report_xlsx)
    js = Path(result.report_json)
    audit = Path(result.audit_json)
    assert xlsx.exists() and js.exists() and audit.exists()

    # 3) the workbook has the full, stable sheet schema
    wb = load_workbook(xlsx)
    for sheet in REPORT_SHEETS:
        assert sheet in wb.sheetnames

    # 4) the JSON report is well-formed and internally consistent
    doc = json.loads(js.read_text(encoding="utf-8"))
    assert doc["audit"]["sha256"] == EXPECTED_FIXTURE_SHA256
    assert doc["ingestion"]["record_count"] == 82
    assert doc["reconciliation"]["created"] == 25

    # 5) real per-stage timings were captured
    assert result.timings["ingest_seconds"] >= 0.0
    assert "reconcile_seconds" in result.timings
    assert "report_write_seconds" in result.timings

    # 6) every scenario the pipeline ran passed
    assert result.all_scenarios_passed
    assert result.counts["canonical_jobs"] == 25


def test_pipeline_report_write_results_are_atomic(real_fixture, tmp_path):
    result = run_phase09(
        real_fixture, tmp_path / "phase09", expansion_sizes=(500,)
    )
    for key in ("xlsx", "json", "audit"):
        wr = result.write_results[key]
        assert wr["validated"] is True
        assert wr["locked"] is False
        # no leftover temp files
    out_dir = tmp_path / "phase09"
    assert not list(out_dir.glob("*.part.*"))
