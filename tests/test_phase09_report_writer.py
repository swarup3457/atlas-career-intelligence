"""Phase 0.9 — atomic report writer reliability tests.

Covers the temp -> reopen/validate -> os.replace protocol, the deterministic
locked-file alternate policy, incomplete-temp crash recovery, and the JSON
writer. All offline, tmp_path only.
"""

from __future__ import annotations

import json
import os
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from atlas.data_integrity import report_writer as rw
from atlas.reporting.excel import ExcelReporter

pytestmark = pytest.mark.unit


def _sample_workbook():
    return ExcelReporter().build_workbook(
        [("Summary", ["metric", "value"], [{"metric": "x", "value": 1}]),
         ("Findings", ["code"], [{"code": "NONE"}])]
    )


# ---------------------------------------------------------------------------
# Happy path + no leftover temp
# ---------------------------------------------------------------------------
def test_atomic_workbook_write_creates_valid_file(tmp_path):
    final = tmp_path / "report.xlsx"
    res = rw.write_workbook_atomic(final, _sample_workbook(), expected_sheets=("Summary", "Findings"))
    assert Path(res.written_path) == final
    assert final.exists()
    assert not res.locked
    # temp file must not linger
    assert not (tmp_path / "report.part.xlsx").exists()
    # file is a real, reopenable workbook
    wb = load_workbook(final)
    assert set(wb.sheetnames) == {"Summary", "Findings"}


def test_atomic_json_write_roundtrips(tmp_path):
    final = tmp_path / "report.json"
    payload = {"b": 2, "a": 1, "nested": {"x": [1, 2, 3]}}
    res = rw.write_json_atomic(final, payload)
    assert Path(res.written_path) == final
    assert json.loads(final.read_text(encoding="utf-8")) == payload
    assert not (tmp_path / "report.part.json").exists()


# ---------------------------------------------------------------------------
# Validation: a bad temp never becomes the final file
# ---------------------------------------------------------------------------
def test_validation_failure_discards_temp_and_raises(tmp_path):
    final = tmp_path / "report.xlsx"
    with pytest.raises(rw.ReportWriteError):
        rw.write_workbook_atomic(
            final, _sample_workbook(), expected_sheets=("Summary", "DoesNotExist")
        )
    # neither the final nor the temp exists afterwards
    assert not final.exists()
    assert not (tmp_path / "report.part.xlsx").exists()


def test_json_validation_failure_on_unserializable(tmp_path):
    final = tmp_path / "bad.json"
    with pytest.raises(rw.ReportWriteError):
        rw.write_json_atomic(final, {"x": object()})
    assert not final.exists()


# ---------------------------------------------------------------------------
# Locked final -> deterministic alternate
# ---------------------------------------------------------------------------
def test_locked_final_uses_deterministic_alternate(tmp_path, monkeypatch):
    final = tmp_path / "report.xlsx"
    # Pre-create the final so we can assert the alternate is distinct.
    final.write_bytes(b"existing")

    real_replace = os.replace

    def fake_replace(src, dst, *a, **k):
        if Path(dst) == final:
            raise PermissionError("locked by another process")
        return real_replace(src, dst, *a, **k)

    monkeypatch.setattr(rw.os, "replace", fake_replace)

    res = rw.write_workbook_atomic(final, _sample_workbook())
    assert res.locked is True
    assert res.used_alternate is True
    expected_alt = tmp_path / "report.locked.xlsx"
    assert Path(res.written_path) == expected_alt
    assert expected_alt.exists()
    # deterministic: same alternate path chosen again on a repeat locked write
    res2 = rw.write_workbook_atomic(final, _sample_workbook())
    assert Path(res2.written_path) == expected_alt


# ---------------------------------------------------------------------------
# Incomplete temp crash recovery
# ---------------------------------------------------------------------------
def test_recover_incomplete_temp_files(tmp_path):
    # simulate a crash: a leftover .part temp with garbage
    leftover = tmp_path / "report.part.xlsx"
    leftover.write_bytes(b"partial-garbage")
    other = tmp_path / "data.part.json"
    other.write_text("{")
    removed = rw.recover_incomplete_temps(tmp_path)
    assert str(leftover) in removed
    assert str(other) in removed
    assert not leftover.exists()
    assert not other.exists()


def test_write_recovers_stale_temp_first(tmp_path):
    final = tmp_path / "report.xlsx"
    stale = tmp_path / "report.part.xlsx"
    stale.write_bytes(b"stale")
    res = rw.write_workbook_atomic(final, _sample_workbook())
    assert str(stale) in res.recovered
    assert final.exists()
    assert not stale.exists()


def test_recover_targets_single_artifact(tmp_path):
    keep = tmp_path / "other.part.json"
    keep.write_text("{")
    target = tmp_path / "report.part.xlsx"
    target.write_bytes(b"x")
    removed = rw.recover_incomplete_temps(tmp_path, final_name="report.xlsx")
    assert str(target) in removed
    assert keep.exists()  # untouched — different artifact
