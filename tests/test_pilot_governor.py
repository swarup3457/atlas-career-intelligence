"""Failing-first end-to-end pilot governor test (offline, deterministic, fake client).

Exercises the full corrected pipeline: LangGraph governor -> constrained tools -> India +
stack + experience gates -> unique 8-sheet workbook with audit columns -> product validator,
plus 10/10 terminal, retry isolation, restart resume, usage aggregation, and history gate.
"""

from __future__ import annotations

import json

import pytest

from atlas.pilot.config import PilotConfig
from atlas.pilot.governor import PilotRuntime, run_pilot
from atlas.pilot.report import ALL_JOBS_COLUMNS, REQUIRED_SHEETS
from atlas.pilot.worker import LlmCompanySearchWorker
from atlas.sources.http_client import HttpResponse

_CAREERS_HTML = ("<!doctype html><html><body>Careers. board: "
                 "https://boards.greenhouse.io/{tok} openings</body></html>")

_JOBS = {
    "acme": [
        {"id": 1, "title": "Java Backend Engineer", "location": {"name": "Bengaluru, India"},
         "absolute_url": "https://boards.greenhouse.io/acme/jobs/1", "requisition_id": "A1",
         "first_published": "2026-09-01", "updated_at": "2026-09-05",
         "content": "Java Spring Boot REST APIs Hibernate MySQL microservices. 2-3 years."},
        {"id": 2, "title": "Senior Backend Engineer (Ruby on Rails)", "location": {"name": "Bengaluru, India"},
         "absolute_url": "https://boards.greenhouse.io/acme/jobs/2", "requisition_id": "A2",
         "content": "Ruby on Rails PostgreSQL. 8+ years."},
        {"id": 3, "title": "Java Engineer", "location": {"name": "San Francisco, CA"},
         "absolute_url": "https://boards.greenhouse.io/acme/jobs/3", "requisition_id": "A3",
         "content": "Java Spring Boot. 2 years."},
        {"id": 4, "title": "Full Stack Engineer", "location": {"name": "Hyderabad, India"},
         "absolute_url": "https://boards.greenhouse.io/acme/jobs/4", "requisition_id": "A4",
         "content": "Node.js backend React frontend TypeScript. 2 years."},
    ],
    "globex": [
        {"id": 9, "title": ".NET Developer", "location": {"name": "Pune, India"},
         "absolute_url": "https://boards.greenhouse.io/globex/jobs/9", "requisition_id": "G9",
         "first_published": "2026-09-03", "updated_at": "2026-09-07",
         "content": "C# ASP.NET Core Entity Framework SQL Server. 2-3 years."},
    ],
}


def _fake_client_factory():
    class Fake:
        def fetch(self, request):
            import re
            url = request.url
            for tok in ("acme", "globex"):
                if url.endswith(f"/careers") and tok in url:
                    return HttpResponse.build(200, body=_CAREERS_HTML.format(tok=tok).encode(),
                                              headers={"Content-Type": "text/html"}, url=url)
            m = re.search(r"/v1/boards/(\w+)/jobs", url)
            if m and "/jobs/" not in url:
                tok = m.group(1)
                return HttpResponse.build(200, body=json.dumps({"jobs": _JOBS.get(tok, [])}).encode(),
                                          headers={"Content-Type": "application/json"}, url=url)
            d = re.search(r"/v1/boards/(\w+)/jobs/(\d+)", url)
            if d:
                tok, jid = d.group(1), int(d.group(2))
                job = next((j for j in _JOBS.get(tok, []) if j["id"] == jid), None)
                if job:
                    return HttpResponse.build(200, body=json.dumps(job).encode(),
                                              headers={"Content-Type": "application/json"}, url=url)
            return HttpResponse.build(404, body=b"{}", headers={"Content-Type": "application/json"}, url=url)
    return Fake()


def _config(companies=("Acme", "Globex")):
    lanes = ("JAVA_BACKEND", "JAVA_FULLSTACK", "REACT_FRONTEND", "DOTNET",
             "ENTERPRISE_HR_PAYROLL_INTEGRATION")
    qt = {l: (l.split("_")[0].title() + " Developer", "Software Engineer") for l in lanes}
    return PilotConfig(
        companies=companies, primary_lanes=lanes, secondary_audit_lanes=("GENERAL_SOFTWARE",),
        query_templates=qt, location_variants=("India", "Bengaluru", "Hyderabad", "Pune"),
        forbidden_locations=("United States", "San Francisco"),
        unsupported_mandatory_backend=("Python", "Node.js", "Ruby"),
        experience_bands=("0-2", "2", "2-3", "3"),
        company_search_preferred_model="claude-sonnet-5", escalation_preferred_model="claude-opus-4.8",
        engineering_model="claude-opus-4.8", concurrency_company_agents=2,
        max_search_rounds_per_company=2, max_escalations_per_company=1,
        max_job_details_per_company=30, max_pages_or_load_more_per_query=2, update_latest=False,
        raw={"pilot_name": "TEST_PILOT"},
    )


@pytest.fixture
def profile(tmp_path):
    from atlas.candidate.search_profile import load_candidate_search_profile
    return load_candidate_search_profile(live=True, private_out=tmp_path / "p.json")


def _runtime(tmp_path, profile, companies=("Acme", "Globex"), run_id="TESTRUN"):
    config = _config(companies)
    worker = LlmCompanySearchWorker(config=config, mode="deterministic",
                                    profile_summary=profile.redacted_dict(),
                                    http_client_factory=_fake_client_factory)
    return PilotRuntime(config=config, worker=worker, profile=profile,
                        output_root=tmp_path, run_id=run_id)


def _patch_hints(monkeypatch):
    import atlas.pilot.discovery as disc
    import atlas.pilot.governor as gov
    hints = dict(disc.OFFICIAL_DOMAIN_HINTS)
    hints.update({"acme": "acme.example", "globex": "globex.example"})
    entries = dict(disc.CAREERS_ENTRY_HINTS)
    entries.update({"acme": "https://www.acme.example/careers",
                    "globex": "https://www.globex.example/careers"})
    monkeypatch.setattr(disc, "OFFICIAL_DOMAIN_HINTS", hints)
    monkeypatch.setattr(disc, "CAREERS_ENTRY_HINTS", entries)
    monkeypatch.setattr(gov, "OFFICIAL_DOMAIN_HINTS", hints)
    monkeypatch.setattr(gov, "CAREERS_ENTRY_HINTS", entries)


def test_pilot_end_to_end_india_only_with_audit_columns(tmp_path, profile, monkeypatch):
    from openpyxl import load_workbook

    _patch_hints(monkeypatch)
    rt = _runtime(tmp_path, profile)
    outcome = run_pilot(rt)

    assert outcome.companies_terminal == 2  # all planned companies terminal
    assert outcome.india_jobs >= 2          # acme Java + globex .NET
    assert outcome.report is not None
    wb = load_workbook(outcome.report.workbook_path, read_only=True)
    assert set(wb.sheetnames) == set(REQUIRED_SHEETS)
    ws = wb["All_Jobs"]
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    for audit_col in ("Geography_Decision", "Location_Evidence", "Supported_Stack_Evidence",
                      "Unsupported_Mandatory_Backend", "Experience_Decision", "Requirement_Evidence",
                      "LLM_Search_Model", "Company_Search_Task_ID"):
        assert audit_col in header
    gi = header.index("Geography_Decision")
    li_ev = header.index("Location_Evidence")
    ui = header.index("Unsupported_Mandatory_Backend")
    li = header.index("Location")
    for row in ws.iter_rows(min_row=2, values_only=True):
        # Section 13: the Geography_Decision audit column must read INDIA_ELIGIBLE, with the
        # fine-grained typed class preserved in Location_Evidence.
        assert row[gi] == "INDIA_ELIGIBLE", row
        assert any(t in (row[li_ev] or "") for t in
                   ("INDIA_PRIMARY", "INDIA_SECONDARY", "REMOTE_INDIA", "INDIA_WIDE")), row
        assert not row[ui], f"unsupported backend in main row: {row}"
        assert "San Francisco" not in (row[li] or "")
    wb.close()
    assert outcome.validation.passed, outcome.validation.failures


def test_foreign_and_wrong_stack_excluded_from_main(tmp_path, profile, monkeypatch):
    _patch_hints(monkeypatch)
    rt = _runtime(tmp_path, profile)
    outcome = run_pilot(rt)
    run_dir = outcome.report.run_dir
    accepted = json.loads((run_dir / "jobs_accepted.json").read_text())
    foreign = json.loads((run_dir / "foreign_leads.json").read_text())
    titles = {a["title"] for a in accepted}
    assert "Java Engineer" not in titles  # San Francisco Java -> foreign, not main
    assert any(f["geography_decision"] == "FOREIGN_EXCLUDED" for f in foreign)
    # Node.js + React full stack must not appear as a React main row
    assert not any("Full Stack" in a["title"] and a["lane"] == "REACT_FRONTEND" for a in accepted)


def test_ten_of_ten_terminal_required(tmp_path, profile, monkeypatch):
    _patch_hints(monkeypatch)
    # a company with no source resolves truthfully but still counts as terminal
    rt = _runtime(tmp_path, profile, companies=("Acme", "Globex", "NoSourceCo"), run_id="TR10")
    outcome = run_pilot(rt)
    assert outcome.companies_terminal == 3
    cov = json.loads((outcome.report.run_dir / "coverage.json").read_text())
    statuses = {c["company"]: c["status"] for c in cov["companies"]}
    assert statuses["NoSourceCo"] == "OFFICIAL_SOURCE_UNRESOLVED"


def test_restart_resume_skips_completed(tmp_path, profile, monkeypatch):
    _patch_hints(monkeypatch)
    # first run publishes partials + workbook
    rt1 = _runtime(tmp_path, profile, run_id="RESUME1")
    run_pilot(rt1)
    run_dir = tmp_path / "llm_pilots" / "RESUME1"
    partial_dir = run_dir / "_partial"
    assert (partial_dir / "acme.json").exists()
    # simulate a process that died AFTER completing companies but BEFORE (re)publishing:
    # remove the published workbook so finalize can publish once on resume.
    for wb in run_dir.glob("Atlas_LLM_India_Pilot_*.xlsx"):
        wb.unlink()

    class BoomWorker(LlmCompanySearchWorker):
        def search_company(self, task, usage):  # noqa: ARG002
            raise AssertionError(f"should not re-run completed company {task.company}")

    config = _config()
    boom = BoomWorker(config=config, mode="deterministic", http_client_factory=_fake_client_factory)
    rt2 = PilotRuntime(config=config, worker=boom, profile=profile, output_root=tmp_path, run_id="RESUME1")
    outcome = run_pilot(rt2)  # must NOT raise (completed companies loaded from partials)
    assert outcome.companies_terminal == 2
    assert outcome.india_jobs >= 2


def test_usage_totals_reconcile(tmp_path, profile, monkeypatch):
    _patch_hints(monkeypatch)
    rt = _runtime(tmp_path, profile)
    run_pilot(rt)
    usage = json.loads((rt.output_root / "llm_pilots" / "TESTRUN" / "llm_usage.json").read_text())
    totals = usage["totals"]
    by_model = usage["by_model"]
    for f in ("input_tokens", "output_tokens"):
        assert sum(m[f] for m in by_model.values()) == totals[f]


def test_cli_summary_emits_absolute_resolvable_paths(tmp_path, profile, monkeypatch):
    # The machine/workbook gate must be able to resolve the workbook + run_dir from ANY CWD,
    # so the CLI summary must emit ABSOLUTE paths (pass-2 output_gate contract fix).
    import os
    from pathlib import Path

    from atlas.pilot.cli import _abs, _summarize

    _patch_hints(monkeypatch)
    rt = _runtime(tmp_path, profile, run_id="ABSPATHS")
    outcome = run_pilot(rt)
    summary = _summarize(rt, outcome, "")
    assert os.path.isabs(summary["workbook"]), summary["workbook"]
    assert os.path.isabs(summary["run_dir"]), summary["run_dir"]
    # resolvable from an unrelated CWD (simulate the validator running elsewhere)
    assert Path(summary["workbook"]).exists()
    assert list(Path(summary["run_dir"]).glob("Atlas_LLM_India_Pilot_*.xlsx"))
    # relative variant mirrors the report's raw path (relative in real runs where
    # output_root is 'output/production'; absolute here only because tmp_path is absolute)
    assert summary["workbook_relative"] == str(outcome.report.workbook_path).replace("\\", "/")
    assert _abs(None) is None


class _StubReport:
    def __init__(self, root):
        self.run_dir = root
        self.workbook_path = root / "stub.xlsx"
        self.all_jobs_rows = 0
        self.company_coverage_rows = 0
        self.foreign_leads = 0
        self.outcome = "COMPLETE_NO_MATCHES"
