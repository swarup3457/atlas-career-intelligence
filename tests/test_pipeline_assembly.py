"""End-to-end deterministic assembly test (PRODUCTION R1 §9–§10, §15)."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from atlas.persistence.sqlite import StateStore
from atlas.vscode_hunt.pipeline import assemble_run_outputs

VALID_JOB = {
    "proposed_decision": "accept", "title": "Java Backend Engineer",
    "official_url": "https://jobs.lever.co/acme/eng-123", "location": "Bengaluru, India",
    "lane": "JAVA_BACKEND",
    "description": "We are hiring a Java Backend Engineer to build Spring Boot microservices in Java. Requires strong Java, JVM and Spring experience. 2+ years experience.",
    "experience_text": "2+ years", "requisition_id": "REQ-123", "posted_date": "2026-09-01",
    "stack": ["Java", "Spring Boot"], "evidence_snippets": ["spring boot microservices"],
}
DRIVETRAIN_JOB = {
    "proposed_decision": "accept", "title": "Software Engineer", "official_url": "", "url": "",
    "location": "", "lane": "JAVA_BACKEND", "stack": [], "experience_text": "", "description": "",
}


def test_assemble_run_outputs_end_to_end(tmp_path: Path) -> None:
    evidence = tmp_path / "evidence"
    live = tmp_path / "live" / "run"
    evidence.mkdir()
    live.mkdir(parents=True)
    for name in ("freehire_raw.json", "freehire_queued.json", "freehire_health.json", "official_verification.json"):
        (evidence / name).write_text("[]", encoding="utf-8")
    (evidence / "discovery_queries.json").write_text(json.dumps({"queries": []}), encoding="utf-8")
    (live / "result.json").write_text(json.dumps({
        "task_id": "run::acme", "attempt_id": "primary-1", "company": "Acme", "company_id": "acme",
        "official_domain": "acme.com", "jobs": [VALID_JOB, DRIVETRAIN_JOB], "rejections": [],
        "foreign_leads": [], "detail_urls": [], "evidence_quotes": [], "source_health": {},
        "result_states": {}, "completion_claim": True,
    }), encoding="utf-8")

    with StateStore(tmp_path / "state.sqlite") as store:
        out = assemble_run_outputs(store=store, run_id="run", evidence_root=evidence, live_root=live, output_root=tmp_path / "out")

    assert out["validated_count"] == 1
    assert out["history"]["new"] == 1

    master = load_workbook(out["master_workbook"], read_only=True)
    active = list(master["Active_Verified_Jobs"].iter_rows(min_row=2, values_only=True))
    master.close()
    assert len(active) == 1 and active[0][2] == "Java Backend Engineer"

    run_wb = load_workbook(out["run_workbook"], read_only=True)
    validated = list(run_wb["Validated_Jobs"].iter_rows(min_row=2, values_only=True))
    run_wb.close()
    assert len(validated) == 1
