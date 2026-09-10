"""Failing-first tests for the five-company PRODUCT pilot (prompt s.11: execution/validation/report)."""

from __future__ import annotations

import datetime
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from atlas.candidate.search_profile import load_candidate_search_profile
from atlas.company.product_pool import load_product_pool, select_fixed_cohort
from atlas.hunt.role_intent import load_role_intent_policy
from atlas.pilot.evaluate import evaluate_pilot
from atlas.pilot.models import CompanySearchResult, CompanyStatus, JobDetailEvidence, LaneCoverage
from atlas.pilot.product_config import load_product_pilot_config
from atlas.pilot.product_report import (
    REQUIRED_SHEETS,
    seal_selection_manifest,
    write_product_pilot_report,
)
from atlas.policy.loader import load_policy

IMPORT = Path(r"C:\Atlas-Agent-Import")
CONFIG = IMPORT / "Atlas_Product_Company_5_Pilot_Config_20260910.yaml"
POOL = IMPORT / "Atlas_Product_Company_Pool_Seed_20260910.yaml"
LANES = ("JAVA_BACKEND", "JAVA_FULLSTACK", "REACT_FRONTEND", "DOTNET", "ENTERPRISE_HR_PAYROLL_INTEGRATION")
TODAY = datetime.date(2026, 9, 10)


def _profile():
    return load_candidate_search_profile(live=False, allow_synthetic=True)


def _job(**kw) -> JobDetailEvidence:
    base = dict(title="Java Backend Developer", company="ServiceNow",
                location="Bengaluru, Karnataka, India", work_mode="Hybrid",
                description="Build backend services with Java, Spring Boot, Hibernate, REST APIs, "
                            "Microservices and SQL.",
                mandatory_requirements=("Java", "Spring Boot", "REST APIs"),
                experience_text="2 to 3 years of experience",
                official_url="https://careers.servicenow.com/jobs/REQ123",
                requisition_id="REQ123", source_family="CUSTOM_CAREERS",
                evidence_snippets=("Java, Spring Boot", "Bengaluru, India"))
    base.update(kw)
    return JobDetailEvidence(**base)


def _result(company, jobs, status=CompanyStatus.COMPLETE.value) -> CompanySearchResult:
    return CompanySearchResult(
        company=company, official_domain=f"{company.lower()}.com",
        career_entry_url=f"https://{company.lower()}.com/careers", route="STRUCTURED_ATS",
        source_family="CUSTOM_CAREERS", status=status,
        lanes={l: LaneCoverage(lane=l, attempted=True) for l in LANES},
        jobs=jobs, model="deterministic", task_id=f"t::{company}",
    )


# --------------------------------------------------------------------------- config

def test_product_config_seals_five_companies_concurrency_one_one_retry():
    cfg = load_product_pilot_config(CONFIG)
    assert cfg.companies == ("Microsoft", "Google", "ServiceNow", "Workday", "Atlassian")
    assert cfg.pilot_config.concurrency_company_agents == 1
    assert cfg.max_internal_retries_per_company == 1
    # one internal retry => two search rounds; NO opus escalation in this pilot
    assert cfg.pilot_config.max_search_rounds_per_company == 2
    assert cfg.pilot_config.max_escalations_per_company == 0
    assert cfg.update_latest is False


# --------------------------------------------------------------------------- validation gates

def test_java_backend_india_positive_control():
    pe = evaluate_pilot([_result("ServiceNow", [_job()])],
                        load_role_intent_policy(), load_policy(), _profile(), today=TODAY)
    assert any(a.company == "ServiceNow" and a.lane == "JAVA_BACKEND" for a in pe.accepted)
    assert not pe.foreign_leads


def test_foreign_location_routed_to_foreign_leads_not_validated():
    job = _job(location="London, United Kingdom", official_url="https://careers.servicenow.com/jobs/UK9")
    pe = evaluate_pilot([_result("ServiceNow", [job])],
                        load_role_intent_policy(), load_policy(), _profile(), today=TODAY)
    assert not pe.accepted
    assert any(fl["company"] == "ServiceNow" for fl in pe.foreign_leads)


def test_card_only_no_title_rejected():
    job = _job(title="", description="", mandatory_requirements=())
    pe = evaluate_pilot([_result("ServiceNow", [job])],
                        load_role_intent_policy(), load_policy(), _profile(), today=TODAY)
    assert not pe.accepted
    assert any(rj.reason_code == "REJECT_NON_JOB_CAPTURE" for rj in pe.rejected)


def test_mandatory_four_plus_years_rejected():
    job = _job(experience_text="Minimum 5 years of hands-on experience required",
               description="Java Spring Boot senior role. Minimum 5 years required.")
    pe = evaluate_pilot([_result("ServiceNow", [job])],
                        load_role_intent_policy(), load_policy(), _profile(), today=TODAY)
    assert not any(a.title == job.title for a in pe.accepted)
    assert any(rj.title == job.title for rj in pe.rejected)


def test_unsupported_backend_fullstack_not_accepted_as_java():
    # React present but backend is Python/Node (no Java/.NET) -> not a supported full stack accept
    job = _job(title="Full Stack Engineer",
               description="Build UIs in React with a Python and Node.js backend.",
               mandatory_requirements=("React", "Python", "Node.js"))
    pe = evaluate_pilot([_result("ServiceNow", [job])],
                        load_role_intent_policy(), load_policy(), _profile(), today=TODAY)
    assert not any(a.lane in ("JAVA_FULLSTACK", "DOTNET") and a.company == "ServiceNow"
                   for a in pe.accepted)


def test_java_8_is_a_version_not_eight_years():
    job = _job(experience_text="2-3 years", description="Java 8, Spring Boot, REST APIs, Microservices.")
    pe = evaluate_pilot([_result("ServiceNow", [_job(), job])],
                        load_role_intent_policy(), load_policy(), _profile(), today=TODAY)
    # the Java 8 role must not be rejected for a bogus 8-year experience gate
    assert not any(rj.title == job.title and "experience" in rj.detail.lower() for rj in pe.rejected)


# --------------------------------------------------------------------------- report

def _five_results():
    return [
        _result("Microsoft", []),
        _result("Google", [], status=CompanyStatus.ACCESS_LIMITED.value),   # blocked
        _result("ServiceNow", [_job()]),                                    # one accept
        _result("Workday", [], status=CompanyStatus.TRUSTED_ZERO.value),    # zero-match
        _result("Atlassian", [_job(company="Atlassian", location="London, UK",
                                    official_url="https://atlassian.com/jobs/UK1")]),  # foreign
    ]


def _write(tmp_path, run_id="RPT_TEST"):
    cfg = load_product_pilot_config(CONFIG)
    pool = load_product_pool(POOL)
    selection = select_fixed_cohort(pool, cfg.companies, seed=cfg.selection_seed,
                                    target_lanes=cfg.primary_lanes, run_id=run_id)
    results = _five_results()
    pe = evaluate_pilot(results, load_role_intent_policy(), load_policy(), _profile(), today=TODAY)
    usage = {"totals": {"ai_credits": 12.0, "input_tokens": 100, "output_tokens": 50,
                        "tool_calls": 30, "models": ["deterministic"]}, "by_model": {}}
    report = write_product_pilot_report(
        pe, results, usage, cfg, selection, run_id=run_id,
        output_root=tmp_path, outcome="PARTIAL", candidate_provenance={"synthetic": True})
    return report, selection


def test_eight_sheets_and_all_five_companies_in_coverage(tmp_path):
    report, _ = _write(tmp_path)
    wb = load_workbook(report.workbook_path, read_only=True)
    assert set(REQUIRED_SHEETS) <= set(wb.sheetnames)
    cov = wb["Company_Coverage"]
    companies = {row[0] for row in cov.iter_rows(min_row=2, max_col=1, values_only=True)}
    wb.close()
    assert {"Microsoft", "Google", "ServiceNow", "Workday", "Atlassian"} <= companies


def test_selection_audit_reconciles_with_cohort(tmp_path):
    report, selection = _write(tmp_path)
    wb = load_workbook(report.workbook_path, read_only=True)
    audit = wb["Selection_Audit"]
    audited = [row[0] for row in audit.iter_rows(min_row=2, max_col=1, values_only=True)]
    wb.close()
    assert audited == list(selection.names())


def test_counts_reconcile_and_hashes_written(tmp_path):
    report, _ = _write(tmp_path)
    assert report.all_jobs_rows == 1        # only ServiceNow accepted
    assert report.foreign_leads == 1        # Atlassian foreign
    hashes = json.loads((report.run_dir / "artifact_hashes.json").read_text(encoding="utf-8"))
    assert hashes and (report.workbook_path.name in hashes)


def test_workbook_is_unique_no_overwrite(tmp_path):
    _write(tmp_path, run_id="UNIQ")
    with pytest.raises(FileExistsError):
        _write(tmp_path, run_id="UNIQ")


def test_manifest_sealed_before_run_and_no_latest(tmp_path):
    cfg = load_product_pilot_config(CONFIG)
    pool = load_product_pool(POOL)
    selection = select_fixed_cohort(pool, cfg.companies, seed=cfg.selection_seed, run_id="SEAL")
    run_dir = tmp_path / "product_pilots" / "SEAL"
    manifest = seal_selection_manifest(run_dir, selection, cfg, run_id="SEAL")
    assert manifest.exists()
    payload = json.loads(manifest.read_text(encoding="utf-8"))
    assert payload["companies"] == list(cfg.companies)
    assert payload["selection_mode"] == "FIXED_REPRODUCIBLE_BENCHMARK"
    # sealing does not create/update any 'latest' pointer
    assert not (tmp_path / "latest").exists()


def test_resume_skips_completed_company(tmp_path):
    from atlas.pilot.governor import PilotRuntime
    from atlas.pilot.worker import LlmCompanySearchWorker

    cfg = load_product_pilot_config(CONFIG)
    worker = LlmCompanySearchWorker(config=cfg.pilot_config, mode="deterministic")
    runtime = PilotRuntime(config=cfg.pilot_config, worker=worker, profile=_profile(),
                           output_root=tmp_path, run_id="RESUME", subdir="product_pilots")
    done = _result("ServiceNow", [_job()])
    runtime.save_partial(done)
    # a completed company is reloaded from the partial store (never re-searched)
    reload = runtime.load_partial("ServiceNow")
    assert reload is not None and reload.company == "ServiceNow"
    assert runtime.load_partial("Microsoft") is None
