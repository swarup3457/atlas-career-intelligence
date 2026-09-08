"""Phase 1E/F Stage 5 — stable production output contract."""

from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from atlas.config import load_settings
from atlas.reporting.mapping import REQUIRED_SHEETS
from atlas.reporting.production_output import (
    LATEST_WORKBOOK,
    LatestPaths,
    ProductionOutputPublisher,
    ProductionRunPaths,
    RunAlreadyPublishedError,
    recommendations_payload,
    sheet_data_from_evaluations,
)


@pytest.fixture
def settings(tmp_path):
    return load_settings(production_output_root=str(tmp_path / "production"))


def _fake_eval(job_key, rec, company="Acme", fit=80):
    class E:
        def to_dict(self):
            return {"job_key": job_key, "company": company, "title": "Java Dev",
                    "lane": "JAVA_BACKEND", "eligibility": "ELIGIBLE", "verification": "VERIFIED_OFFICIAL",
                    "freshness": "0-7 days", "candidate_fit": fit, "recommendation": rec,
                    "confidence": 0.7, "reasoning_model": "none", "reasoning_version": "1EF.rank.1",
                    "strengths": ["Java"], "gaps": [], "supported_evidence_ids": ["prof::Java"]}
    return E()


def _publish(publisher, run_id="RUN_A", status="COMPLETE", evals=None):
    evals = evals if evals is not None else [_fake_eval("k1", "STRONG_APPLY")]
    sheets = sheet_data_from_evaluations(evals, {})
    return publisher.publish(
        run_id, status=status, sheet_data=sheets,
        coverage={"companies": 5}, source_health={"linkedin": "HEALTHY"},
        recommendations=recommendations_payload(evals),
        portal_leads={"count": 3}, verification_summary={"verified": 1},
    )


def test_exact_directory_contract(settings):
    pub = ProductionOutputPublisher(settings)
    res = _publish(pub)
    paths = ProductionRunPaths(pub.root, "RUN_A")
    assert paths.workbook.is_file()
    assert paths.manifest.is_file()
    for name in ("coverage.json", "source_health.json", "recommendations.json",
                 "portal_leads.json", "verification_summary.json"):
        assert paths.side_file(name).is_file()
    assert paths.application_packs_dir.is_dir()
    assert res.status == "COMPLETE"


def test_workbook_has_eight_sheets(settings):
    pub = ProductionOutputPublisher(settings)
    _publish(pub)
    wb = load_workbook(str(ProductionRunPaths(pub.root, "RUN_A").workbook), read_only=True)
    assert tuple(wb.sheetnames) == REQUIRED_SHEETS
    assert len(wb.sheetnames) == 8
    wb.close()


def test_manifest_hashes_present_and_correct(settings):
    import hashlib
    pub = ProductionOutputPublisher(settings)
    _publish(pub)
    paths = ProductionRunPaths(pub.root, "RUN_A")
    manifest = json.loads(paths.manifest.read_text(encoding="utf-8"))
    assert manifest["status"] == "COMPLETE"
    assert manifest["report_valid"] is True
    # every listed file hash matches the file on disk
    for rel, digest in manifest["files"].items():
        data = (paths.run_dir / rel).read_bytes()
        assert hashlib.sha256(data).hexdigest() == digest
    assert manifest["workbook"] in manifest["files"]


def test_latest_atomic_update(settings):
    pub = ProductionOutputPublisher(settings)
    res = _publish(pub, run_id="RUN_A")
    assert res.latest_updated is True
    latest = LatestPaths(pub.root)
    assert latest.workbook.is_file()
    assert latest.latest_run_json.is_file()
    assert latest.latest_run_txt.read_text(encoding="utf-8").strip() == "RUN_A"
    assert json.loads(latest.latest_run_json.read_text())["run_id"] == "RUN_A"


def test_run_immutable_after_complete(settings):
    pub = ProductionOutputPublisher(settings)
    _publish(pub, run_id="RUN_A")
    with pytest.raises(RunAlreadyPublishedError):
        _publish(pub, run_id="RUN_A")


def test_partial_run_does_not_update_latest(settings):
    pub = ProductionOutputPublisher(settings)
    # publish a COMPLETE run first so latest points at it
    _publish(pub, run_id="RUN_A")
    # a later PARTIAL run must NOT overwrite the latest pointer
    res = _publish(pub, run_id="RUN_B", status="PARTIAL")
    assert res.status == "PARTIAL"
    assert res.latest_updated is False
    assert LatestPaths(pub.root).latest_run_txt.read_text().strip() == "RUN_A"


def test_relevant_count_is_after_ranking(settings):
    evals = [
        _fake_eval("k1", "STRONG_APPLY"),
        _fake_eval("k2", "MONITOR"),
        _fake_eval("k3", "NOT_EVALUATED"),
        _fake_eval("k4", "REJECT"),
    ]
    payload = recommendations_payload(evals)
    assert payload["total"] == 4
    assert payload["relevant_after_ranking"] == 1  # only the STRONG_APPLY
    assert payload["not_evaluated"] == 1


def test_not_evaluated_rows_stay_visible(settings):
    evals = [_fake_eval("k1", "STRONG_APPLY"), _fake_eval("k2", "NOT_EVALUATED")]
    pub = ProductionOutputPublisher(settings)
    _publish(pub, run_id="RUN_A", evals=evals)
    wb = load_workbook(str(ProductionRunPaths(pub.root, "RUN_A").workbook), read_only=True)
    ws = wb["All_Jobs"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    # header + 2 data rows (NOT_EVALUATED job remains visible)
    assert len(rows) == 3


def test_invalid_status_complete_downgraded_when_report_invalid(settings, monkeypatch):
    pub = ProductionOutputPublisher(settings)

    def boom(*a, **k):
        raise RuntimeError("simulated writer failure")

    monkeypatch.setattr("atlas.reporting.production_output.write_report", boom)
    res = _publish(pub, run_id="RUN_C")
    assert res.status == "PARTIAL"          # COMPLETE downgraded because workbook invalid
    assert res.report_valid is False
    assert res.latest_updated is False


def test_list_and_show_and_latest(settings):
    pub = ProductionOutputPublisher(settings)
    _publish(pub, run_id="RUN_A")
    _publish(pub, run_id="RUN_B", status="PARTIAL")
    runs = {r["run_id"]: r for r in pub.list_runs()}
    assert set(runs) == {"RUN_A", "RUN_B"}
    assert pub.show_run("RUN_A")["status"] == "COMPLETE"
    assert pub.latest()["run_id"] == "RUN_A"


def test_outputs_cli_prints_paths(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("ATLAS_PRODUCTION_OUTPUT_ROOT", str(tmp_path / "prod"))
    from atlas.cli import main
    from atlas.config import load_settings

    pub = ProductionOutputPublisher(load_settings())
    _publish(pub, run_id="RUN_CLI")

    assert main(["outputs", "list"]) == 0
    out = capsys.readouterr().out
    assert "RUN_CLI" in out and "Production root:" in out

    assert main(["outputs", "latest", "--json"]) == 0
    assert "RUN_CLI" in capsys.readouterr().out

    assert main(["outputs", "show", "--run-id", "RUN_CLI"]) == 0
    show_out = capsys.readouterr().out
    assert "Workbook:" in show_out and "RUN_CLI" in show_out

    # open-latest without --live must NOT open anything, only report
    assert main(["outputs", "open-latest"]) == 0
    assert "would open" in capsys.readouterr().out
