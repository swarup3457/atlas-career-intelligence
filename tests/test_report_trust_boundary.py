"""Regression tests for the report trust boundary (PRODUCTION R1 §4).

A worker acceptance proposal may only become a ``Validated_Jobs`` row after it
independently passes ``validate_job_evidence``. The sanitized Drivetrain shape
(blank URL / stack / experience) must NEVER qualify, and ``PYTHON_VALIDATED`` /
recommendation must never be hardcoded from a bare proposal.
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from atlas.discovery.report import build_discovery_workbook
from atlas.persistence.sqlite import StateStore
from atlas.reporting.trust_boundary import partition_jobs
from atlas.vscode_hunt.models import Backend
from atlas.vscode_hunt.report import build_v32_workbook
from atlas.vscode_hunt.service import VscodeHuntService

# A fully-grounded, India-eligible, official-ATS Java backend job that MUST pass.
VALID_JOB = {
    "proposed_decision": "accept",
    "title": "Java Backend Engineer",
    "official_url": "https://jobs.lever.co/acme/eng-123",
    "location": "Bengaluru, India",
    "lane": "JAVA_BACKEND",
    "description": (
        "We are hiring a Java Backend Engineer to build Spring Boot microservices "
        "in Java. Requires strong Java, JVM and Spring experience. 2+ years experience."
    ),
    "experience_text": "2+ years",
    "requisition_id": "REQ-123",
    "posted_date": "2026-09-01",
    "stack": ["Java", "Spring Boot"],
    "evidence_snippets": ["spring boot microservices"],
}

# The exact sanitized Drivetrain failure shape: blank URL / stack / experience.
DRIVETRAIN_JOB = {
    "proposed_decision": "accept",
    "title": "Software Engineer",
    "official_url": "",
    "url": "",
    "location": "",
    "lane": "JAVA_BACKEND",
    "stack": [],
    "experience_text": "",
    "description": "",
}


def test_partition_rejects_drivetrain_blank_shape() -> None:
    validated, rejected = partition_jobs([DRIVETRAIN_JOB], official_domain="drivetrain.com", company="Drivetrain")
    assert validated == []
    assert len(rejected) == 1
    assert rejected[0].reason_code == "NON_OFFICIAL_URL"


def test_partition_accepts_valid_job_with_real_recommendation() -> None:
    validated, rejected = partition_jobs([VALID_JOB], official_domain="acme.com", company="Acme")
    assert rejected == []
    assert len(validated) == 1
    job = validated[0]
    assert job.verification_status == "PYTHON_VALIDATED"
    assert job.recommendation in {"APPLY_NOW", "APPLY_AFTER_TAILORING", "STRETCH"}
    assert job.recommendation != "MANUAL_REVIEW"


def test_partition_skips_non_acceptance_proposals() -> None:
    reject_proposal = {**VALID_JOB, "proposed_decision": "reject"}
    validated, rejected = partition_jobs([reject_proposal], official_domain="acme.com", company="Acme")
    assert validated == []
    assert rejected == []


def test_discovery_workbook_only_validated_jobs_reach_validated_sheet(tmp_path: Path) -> None:
    evidence = tmp_path / "evidence"
    live = tmp_path / "live" / "run"
    evidence.mkdir()
    live.mkdir(parents=True)
    for name in ("freehire_raw.json", "freehire_queued.json", "freehire_health.json", "official_verification.json"):
        (evidence / name).write_text("[]", encoding="utf-8")
    (evidence / "discovery_queries.json").write_text(json.dumps({"queries": []}), encoding="utf-8")
    result = {
        "task_id": "run::acme", "attempt_id": "primary-1", "company": "Acme", "company_id": "acme",
        "official_domain": "acme.com", "jobs": [VALID_JOB, DRIVETRAIN_JOB],
        "rejections": [], "foreign_leads": [], "detail_urls": [], "evidence_quotes": [],
        "source_health": {}, "result_states": {}, "completion_claim": True,
    }
    (live / "result.json").write_text(json.dumps(result), encoding="utf-8")

    output = build_discovery_workbook(evidence_root=evidence, live_root=live, run_id="run", output_root=tmp_path / "out")
    wb = load_workbook(output, read_only=True)
    validated = list(wb["Validated_Jobs"].iter_rows(min_row=2, values_only=True))
    rejected = list(wb["Rejected_Jobs"].iter_rows(min_row=2, values_only=True))
    wb.close()

    assert len(validated) == 1
    assert validated[0][1] == "Java Backend Engineer"
    assert validated[0][-1] == "PYTHON_VALIDATED"
    # The blank-URL Drivetrain proposal must be a rejection, never a validated row.
    assert any(row[5] == "NON_OFFICIAL_URL" for row in rejected)
    assert all(row[1] != "Software Engineer" for row in validated)


def test_v32_workbook_does_not_leak_raw_proposals(tmp_path: Path) -> None:
    """Even a raw (un-committed) result file must be gated: the old leak path."""
    store = StateStore(tmp_path / "state.sqlite")
    with store:
        service = VscodeHuntService(store, tmp_path / "out")
        run_id = service.create_run([
            {"company_id": "acme", "name": "Acme", "official_domain": "acme.example", "careers_url": "https://acme.example/careers"},
        ])
        task = service.next_tasks(run_id, materialize=True)[0]
        attempt = service.record_attempt(run_id, task["task_id"], "attempt-1", Backend.VSCODE_SUBAGENT)

        raw_result = tmp_path / "raw_result.json"
        raw_result.write_text(json.dumps({
            "jobs": [VALID_JOB, DRIVETRAIN_JOB], "rejections": [], "foreign_leads": [],
            "detail_urls": [], "evidence_quotes": [], "source_health": {}, "result_states": [],
        }), encoding="utf-8")
        with store._conn:
            store._conn.execute(
                "UPDATE vscode_hunt_attempts SET result_path=? WHERE attempt_id=?",
                (str(raw_result), attempt.attempt_id),
            )

        output = build_v32_workbook(service.root, run_id, store._conn, kind="Audit")
        wb = load_workbook(output, read_only=True)
        validated = list(wb["Validated_Jobs"].iter_rows(min_row=2, values_only=True))
        rejected = list(wb["Rejected_Jobs"].iter_rows(min_row=2, values_only=True))
        wb.close()

    assert len(validated) == 1
    assert validated[0][1] == "Java Backend Engineer"
    assert validated[0][-1] == "PYTHON_VALIDATED"
    assert validated[0][9] != "MANUAL_REVIEW"  # recommendation column is a real verdict
    assert any(row[5] == "NON_OFFICIAL_URL" for row in rejected)
