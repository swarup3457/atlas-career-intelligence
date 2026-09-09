"""Search Hunt Recovery V2 adversarial + pipeline tests (build spec 21, 24).

Golden negatives (from the real failed workbook), positive controls, the
experience corpus, one-snapshot-six-lanes, hydrate-once, zero-network
requalification, stratified cohort, diversity caps, end-to-end graph + report +
validator, zero-suitable -> zero packs, and validator wrong-stack detection.
"""

from __future__ import annotations

import datetime

import pytest

from atlas.hunt.campaign import build_stratified_cohort, seal_campaign
from atlas.hunt.experience_v2 import ExperienceFitBand, evaluate_experience_fit
from atlas.hunt.matching import HuntCandidate, diversify_shortlist
from atlas.hunt.models import BoardSnapshot, JobDetailRevision, RawJob, RunLineage
from atlas.hunt.pipeline import FixtureProvider, requalify_details, run_campaign
from atlas.hunt.prefilter import hydration_union, prefilter_snapshot
from atlas.hunt.qualification import QualifiableJob, QualificationStatus, qualify_job
from atlas.hunt.role_family import RoleFamily, classify_role_family
from atlas.hunt.role_intent import load_role_intent_policy
from atlas.hunt.signals import signal_present
from atlas.policy.loader import load_policy
from atlas.policy.models import ExperiencePolicy

TODAY = datetime.date(2026, 9, 9)


@pytest.fixture(scope="module")
def intent():
    return load_role_intent_policy()


@pytest.fixture(scope="module")
def policy():
    return load_policy()


@pytest.fixture(scope="module")
def exp_policy():
    return ExperiencePolicy(preferred_ranges=("1-3", "2"), hard_reject_min_years=4, allow_three_year_when_strong=True)


def _q(intent, exp_policy, title, desc="", experience="", cy=None):
    job = QualifiableJob(job_key=title, title=title, description=desc, experience_text=experience)
    return qualify_job(job, intent, experience_policy=exp_policy, candidate_years=cy)


# ---------------------------------------------------------------------------
# 1. Golden negatives (must NOT be admitted into the wrong lane)
# ---------------------------------------------------------------------------
def test_ruby_backend_not_java(intent, exp_policy):
    r = _q(intent, exp_policy, "Senior Backend Engineer (Ruby on Rails)",
           "Build REST APIs in Ruby on Rails with PostgreSQL and microservices.")
    assert r.by_lane["JAVA_BACKEND"].status == QualificationStatus.REJECT_WRONG_STACK.value
    assert r.by_lane["JAVA_BACKEND"].dominant_stack is not None
    # and it does not re-enter general software as a wrong-stack job
    assert r.by_lane["GENERAL_SOFTWARE"].status == QualificationStatus.REJECT_WRONG_STACK.value
    assert r.primary_lane is None


def test_go_backend_not_java(intent, exp_policy):
    r = _q(intent, exp_policy, "Staff Backend Engineer (Go)",
           "Design Go microservices, REST APIs and SQL.", experience="8+ years")
    assert r.by_lane["JAVA_BACKEND"].status == QualificationStatus.REJECT_WRONG_STACK.value
    assert r.primary_lane is None


def test_product_manager_excluded(intent, exp_policy):
    r = _q(intent, exp_policy, "Technical Product Manager for API Management",
           "Own the API product roadmap and REST API strategy.")
    assert classify_role_family("Technical Product Manager for API Management").family == RoleFamily.PRODUCT_MANAGEMENT.value
    assert r.by_lane["JAVA_BACKEND"].status == QualificationStatus.REJECT_ROLE_FAMILY.value
    assert r.primary_lane is None


def test_sdet_excluded_from_development(intent, exp_policy):
    r = _q(intent, exp_policy, "Software Development Engineer in Test II",
           "Selenium test automation, API testing, Java.")
    assert classify_role_family("Software Development Engineer in Test II").family == RoleFamily.QA_AUTOMATION_SDET.value
    assert r.by_lane["GENERAL_SOFTWARE"].status == QualificationStatus.REJECT_ROLE_FAMILY.value
    assert r.primary_lane is None


def test_generic_backend_api_sql_not_java(intent, exp_policy):
    r = _q(intent, exp_policy, "Backend Engineer",
           "Build REST APIs, SQL and microservices for our platform.")
    assert r.by_lane["JAVA_BACKEND"].status == QualificationStatus.REJECT_WRONG_STACK.value


def test_senior_java_high_experience_rejected(intent, exp_policy):
    r = _q(intent, exp_policy, "Senior Java Engineer",
           "Java Spring Boot backend services.", experience="8+ years")
    assert r.by_lane["JAVA_BACKEND"].status == QualificationStatus.REJECT_EXPERIENCE.value
    assert r.primary_lane is None


# ---------------------------------------------------------------------------
# 2. Positive controls
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "title,desc,experience,want",
    [
        ("Java Backend Engineer", "Java, Spring Boot, REST APIs, Hibernate, MySQL.", "2 years", "JAVA_BACKEND"),
        ("Software Engineer", "Spring Boot is mandatory. Build backend services.", "1-3 years", "JAVA_BACKEND"),
        ("Java Full Stack Developer", "Java, Spring Boot backend with React frontend, TypeScript.", "2-3 years", "JAVA_FULLSTACK"),
        ("Full Stack Engineer", "Node.js backend, React frontend, TypeScript, MongoDB.", "2 years", "REACT_FRONTEND"),
        ("React Developer", "React, ReactJS, TypeScript, responsive UI.", "2 years", "REACT_FRONTEND"),
        (".NET Developer", "C#, ASP.NET Core, Entity Framework, SQL Server.", "2 years", "DOTNET"),
        ("Payroll Software Engineer", "Build payroll and HCM integration software using Java Spring Boot.", "2-3 years", "ENTERPRISE_HR_PAYROLL_INTEGRATION"),
        ("Associate Software Engineer", "Software development, application development, SQL.", "1-2 years", "GENERAL_SOFTWARE"),
    ],
)
def test_positive_controls(intent, exp_policy, title, desc, experience, want):
    r = _q(intent, exp_policy, title, desc, experience)
    assert r.primary_lane == want
    assert r.primary.status == QualificationStatus.QUALIFIED.value


def test_node_react_not_java_fullstack(intent, exp_policy):
    r = _q(intent, exp_policy, "Full Stack Engineer", "Node.js backend, React frontend, TypeScript.", "2 years")
    assert r.by_lane["JAVA_FULLSTACK"].status != QualificationStatus.QUALIFIED.value


# ---------------------------------------------------------------------------
# 3. Experience corpus
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "experience,expect_eligible",
    [
        ("0-2 years", True), ("0-3 years", True), ("1-2 years", True), ("1-3 years", True),
        ("2 years", True), ("2+ years", True), ("2-3 years", True), ("3 years", True),
        ("3+ years", True), ("4+ years", False), ("5+ years", False),
        ("2 years required, 5 years preferred", True), ("up to 2 years", True),
    ],
)
def test_experience_corpus(exp_policy, experience, expect_eligible):
    fit = evaluate_experience_fit(title="Software Engineer", experience_text=experience, policy=exp_policy)
    assert fit.eligible is expect_eligible


def test_experience_title_conflict_uses_explicit(exp_policy):
    fit = evaluate_experience_fit(title="Senior Software Engineer", experience_text="2-3 years", policy=exp_policy)
    assert fit.eligible
    assert fit.mandatory_min_years == 2


def test_experience_senior_no_range_is_manual(exp_policy):
    fit = evaluate_experience_fit(title="Staff Engineer", experience_text="Own architecture.", policy=exp_policy)
    assert fit.fit == ExperienceFitBand.MANUAL_VERIFICATION.value
    assert fit.ambiguous_seniority


def test_experience_4plus_satisfied_by_candidate(exp_policy):
    fit = evaluate_experience_fit(title="Engineer", experience_text="5+ years", policy=exp_policy, candidate_years=6.0)
    assert fit.fit == ExperienceFitBand.STRETCH.value


# ---------------------------------------------------------------------------
# 4. react must not match reactive
# ---------------------------------------------------------------------------
def test_react_not_reactive():
    assert signal_present("Build reactive systems", "React") is False
    assert signal_present("React and Redux", "React") is True


def test_dotnet_csharp_signal():
    assert signal_present("Strong C# and ASP.NET Core", "C#")
    assert signal_present("Using .NET 8", ".NET")
    assert signal_present("ASP.NET Core MVC", "ASP.NET")


# ---------------------------------------------------------------------------
# 5. one snapshot feeds six lanes; each job hydrated once
# ---------------------------------------------------------------------------
def _snapshot():
    raw = (
        RawJob("J1", "Java Backend Engineer", "Bengaluru", summary="Java Spring Boot"),
        RawJob("R1", "React Developer", "Bengaluru", summary="React TypeScript"),
    )
    return BoardSnapshot(
        snapshot_id="S1", campaign_id="C1", company_id="Acme", company_name="Acme",
        source_instance_id="acme:careers", source_family="OFFICIAL_CAREERS", route_family="OFFICIAL_CAREERS",
        raw_jobs=raw,
    )


def test_one_snapshot_all_six_lanes(policy, intent):
    # ONE board fetch per company must feed all six local lane evaluations
    # (collect once, classify many) — never six network fetches.
    company = _fixture_campaign(policy, intent, cohort=1).companies[0].name
    snap = _snapshot()
    snap = BoardSnapshot(
        snapshot_id="S1", campaign_id="C1", company_id=company, company_name=company,
        source_instance_id=f"{company}:careers", source_family="OFFICIAL_CAREERS",
        route_family="OFFICIAL_CAREERS", raw_jobs=snap.raw_jobs,
    )
    provider = FixtureProvider(boards={company: snap}, details={})
    campaign = _fixture_campaign(policy, intent, cohort=1)
    cand = HuntCandidate.from_skills(["Java", "React"], target_lanes=intent.lane_keys())
    result = run_campaign(campaign, intent, policy, provider, cand, candidate_years=2.0, today=TODAY)
    # exactly one board snapshot for the one company (no per-lane refetch)
    assert len(result.snapshots) == 1
    # yet coverage evaluates all six lanes from that single snapshot
    assert {c.lane for c in result.coverage} == set(intent.lane_keys())
    # and the prefilter itself iterates every lane against the one snapshot
    lanes_considered = {d.lane for d in prefilter_snapshot(snap, intent)}
    assert lanes_considered  # recall-oriented; at least the matching lanes admitted


def test_hydrate_union_dedups(intent):
    snap = _snapshot()
    decisions = prefilter_snapshot(snap, intent)
    union = hydration_union(decisions)
    assert sorted(union) == ["J1", "R1"]  # each unique job appears once


# ---------------------------------------------------------------------------
# 6. policy-only requalification makes zero network calls
# ---------------------------------------------------------------------------
def _fixture_campaign(policy, intent, cohort=5):
    return seal_campaign(policy, campaign_id="C1", lanes=intent.lane_keys(),
                         role_policy_hash=intent.fingerprint, cohort_size=cohort, max_batch=10)


def _java_detail(company, sid):
    return JobDetailRevision(
        revision_id=f"D-{sid}", snapshot_id=f"S-{company}", source_job_id=sid, company=company,
        title="Java Backend Engineer", description="Java Spring Boot, REST APIs, Hibernate, MySQL.",
        mandatory_requirements=("Java", "Spring Boot"), experience_text="2-3 years", location="Bengaluru",
        posted_date=TODAY, official_url=f"x/{sid}", requisition_id=sid, verification_state="VERIFIED_OFFICIAL",
        has_live_official_page=True, eligibility_text="Bengaluru, India",
    )


def _provider_for(company):
    raw = (RawJob("J1", "Java Backend Engineer", "Bengaluru", summary="Java Spring Boot"),)
    snap = BoardSnapshot(
        snapshot_id=f"S-{company}", campaign_id="C1", company_id=company, company_name=company,
        source_instance_id=f"{company}:careers", source_family="OFFICIAL_CAREERS", route_family="OFFICIAL_CAREERS",
        raw_jobs=raw,
    )
    return FixtureProvider(boards={company: snap}, details={"J1": _java_detail(company, "J1")})


def test_requalify_zero_network(policy, intent):
    campaign = _fixture_campaign(policy, intent)
    first = campaign.companies[0].name
    provider = _provider_for(first)
    cand = HuntCandidate.from_skills(["Java", "Spring Boot"], target_lanes=intent.lane_keys())
    result = run_campaign(campaign, intent, policy, provider, cand, candidate_years=2.0, today=TODAY)
    assert result.qualified_jobs >= 1

    stored_snaps = {s.company_name: s for s in result.snapshots}
    stored_details: dict = {}
    for d in result.details:
        stored_details.setdefault(d.company, []).append(d)

    requal = requalify_details(campaign, intent, policy, cand, stored_snaps, stored_details,
                               candidate_years=2.0, today=TODAY)
    assert requal.network_calls == 0
    assert requal.qualified_jobs == result.qualified_jobs


# ---------------------------------------------------------------------------
# 7. stratified cohort covers every seed group
# ---------------------------------------------------------------------------
def test_stratified_cohort_covers_groups(policy):
    cohort = build_stratified_cohort(policy.company_seed.companies, 30)
    assert len(cohort) == 30
    groups = {c.group for c in cohort}
    all_groups = {c.group for c in policy.company_seed.companies}
    assert groups == all_groups  # every group represented


def test_campaign_obligations(policy, intent):
    campaign = seal_campaign(policy, campaign_id="C1", lanes=intent.lane_keys(),
                             role_policy_hash=intent.fingerprint, cohort_size=30, max_batch=60)
    assert campaign.company_count == 30
    assert campaign.obligation_count == 30 * 6


# ---------------------------------------------------------------------------
# 8. diversity caps (no data loss)
# ---------------------------------------------------------------------------
def test_diversity_caps(intent):
    from atlas.hunt.matching import CandidateMatchDecision
    from atlas.candidate.eligibility import Recommendation

    matches = [
        CandidateMatchDecision(job_key=f"k{i}", lane="JAVA_BACKEND", company="BigCo",
                               title=f"Java Engineer {i}", match_score=90,
                               recommendation=Recommendation.PRIORITY_APPLY.value)
        for i in range(10)
    ]
    shortlist = diversify_shortlist(matches, intent)
    # per-company cap (5) and per-company-per-lane cap (3) both apply
    assert len(shortlist) <= intent.max_display_per_company
    assert sum(1 for m in shortlist if m.company == "BigCo" and m.lane == "JAVA_BACKEND") <= intent.max_display_per_lane_per_company
    # raw input is never mutated
    assert len(matches) == 10


# ---------------------------------------------------------------------------
# 9. end-to-end graph -> unique workbook + validator pass + no wrong stack
# ---------------------------------------------------------------------------
def _demo_runtime(tmp_path, policy, intent, run_id="T_E2E", allow_extension=False):
    from atlas.hunt.graph import HuntRuntime
    from atlas.hunt.live import build_demo_provider

    provider = build_demo_provider()
    campaign = seal_campaign(policy, campaign_id=run_id, lanes=intent.lane_keys(),
                             role_policy_hash=intent.fingerprint, cohort_size=5, max_batch=10)
    cand = HuntCandidate.from_skills(
        ["Java", "Spring Boot", "REST APIs", "React", "TypeScript", "C#", "ASP.NET Core", "SQL", "payroll"],
        target_lanes=intent.lane_keys())
    lineage = RunLineage(run_id=run_id, run_kind="COLLECTION", role_policy_hash=intent.fingerprint,
                         company_plan_hash=campaign.company_plan_hash, created_at="now")
    return HuntRuntime(intent=intent, policy=policy, provider=provider, candidate=cand, campaign=campaign,
                       output_root=tmp_path, run_id=run_id, lineage=lineage, candidate_years=2.0,
                       allow_extension=allow_extension)


def test_end_to_end_graph_valid(tmp_path, policy, intent):
    from atlas.hunt.graph import run_hunt

    rt = _demo_runtime(tmp_path, policy, intent)
    out = run_hunt(rt)
    assert out.outcome == "ENGINEERING_PASS_WITH_MATCHES"
    assert out.validation is not None and out.validation.passed
    assert out.report.workbook_path.exists()
    assert out.report.workbook_path.name.startswith("Atlas_Jobs_")
    assert run_id_in_name(out.report.workbook_path.name, "T_E2E")
    # coverage covers all six lanes for all five companies
    assert out.report.company_coverage_rows == 5 * 6
    # the six lane positives qualify; the four wrong-stack/off-family controls do not
    assert out.report.all_jobs_rows == 6
    # every shortlist row has a development role family (validator would fail otherwise)
    assert out.validation.passed


def run_id_in_name(name: str, run_id: str) -> bool:
    return name.endswith(f"_{run_id}.xlsx")


def test_workbook_not_overwritten(tmp_path, policy, intent):
    from atlas.hunt.graph import run_hunt

    rt = _demo_runtime(tmp_path, policy, intent, run_id="T_DUP")
    run_hunt(rt)
    # a second run into the SAME run dir must not overwrite the workbook
    rt2 = _demo_runtime(tmp_path, policy, intent, run_id="T_DUP")
    with pytest.raises(FileExistsError):
        run_hunt(rt2)


# ---------------------------------------------------------------------------
# 10. zero suitable jobs -> zero packs -> COMPLETE_NO_MATCHES
# ---------------------------------------------------------------------------
def test_zero_matches_zero_packs(tmp_path, policy, intent):
    from atlas.hunt.graph import HuntRuntime, run_hunt

    # a board with ONLY wrong-stack negative controls
    raw = (RawJob("X1", "Senior Backend Engineer (Ruby on Rails)", "Bengaluru", summary="Ruby on Rails"),)
    company = seal_campaign(policy, campaign_id="T_NM", lanes=intent.lane_keys(),
                            role_policy_hash=intent.fingerprint, cohort_size=3, max_batch=3).companies[0].name
    snap = BoardSnapshot(snapshot_id="S-NM", campaign_id="T_NM", company_id=company, company_name=company,
                         source_instance_id=f"{company}:careers", source_family="OFFICIAL_CAREERS",
                         route_family="OFFICIAL_CAREERS", raw_jobs=raw)
    detail = JobDetailRevision(revision_id="D-X1", snapshot_id="S-NM", source_job_id="X1", company=company,
                              title="Senior Backend Engineer (Ruby on Rails)",
                              description="Ruby on Rails, PostgreSQL, REST APIs.", experience_text="6+ years",
                              location="Bengaluru", posted_date=TODAY, verification_state="VERIFIED_OFFICIAL",
                              has_live_official_page=True, eligibility_text="Bengaluru, India")
    provider = FixtureProvider(boards={company: snap}, details={"X1": detail})
    campaign = seal_campaign(policy, campaign_id="T_NM", lanes=intent.lane_keys(),
                             role_policy_hash=intent.fingerprint, cohort_size=3, max_batch=3)
    cand = HuntCandidate.from_skills(["Java", "React"], target_lanes=intent.lane_keys())
    lineage = RunLineage(run_id="T_NM", run_kind="COLLECTION", role_policy_hash=intent.fingerprint,
                         company_plan_hash=campaign.company_plan_hash, created_at="now")
    rt = HuntRuntime(intent=intent, policy=policy, provider=provider, candidate=cand, campaign=campaign,
                     output_root=tmp_path, run_id="T_NM", lineage=lineage, candidate_years=2.0, allow_extension=False)
    out = run_hunt(rt)
    assert out.outcome == "ENGINEERING_PASS_COMPLETE_NO_MATCHES"
    assert out.report.all_jobs_rows == 0
    assert out.report.application_packs == 0
    assert out.validation.passed  # empty shortlist is still a valid (truthful) report


# ---------------------------------------------------------------------------
# 11. validator catches an injected wrong-stack / non-terminal row
# ---------------------------------------------------------------------------
def test_validator_flags_wrong_stack_row(tmp_path, intent):
    from openpyxl import Workbook
    from atlas.hunt.report import (
        ALL_JOBS_COLUMNS, COMPANY_COVERAGE_COLUMNS, SOURCE_COVERAGE_COLUMNS,
        CLOSED_REJECTED_COLUMNS, RESUME_TAILORING_COLUMNS,
    )
    from atlas.hunt.validator import validate_run_dir

    run_dir = tmp_path / "runs" / "BAD"
    run_dir.mkdir(parents=True)
    wb = Workbook()
    wb.remove(wb.active)
    aj = wb.create_sheet("All_Jobs")
    aj.append(list(ALL_JOBS_COLUMNS))
    # a JAVA_BACKEND row whose Stack_Anchors has NO java anchor (wrong-stack leak)
    row = {c: "" for c in ALL_JOBS_COLUMNS}
    row.update({"Company": "BadCo", "Role_Title": "Backend Engineer", "Lane": "JAVA_BACKEND",
                "Role_Family": "BACKEND_DEVELOPMENT", "Stack_Anchors": "Ruby, PostgreSQL",
                "Qualification_Status": "QUALIFIED", "Recommendation": "PRIORITY_APPLY"})
    aj.append([row[c] for c in ALL_JOBS_COLUMNS])
    cc = wb.create_sheet("Company_Coverage")
    cc.append(list(COMPANY_COVERAGE_COLUMNS))
    for lane in intent.lane_keys():
        r = {c: "" for c in COMPANY_COVERAGE_COLUMNS}
        r.update({"Company": "BadCo", "Lane": lane, "Terminal_Status": "CHECKED"})
        cc.append([r[c] for c in COMPANY_COVERAGE_COLUMNS])
    sc = wb.create_sheet("Source_Coverage")
    sc.append(list(SOURCE_COVERAGE_COLUMNS))
    sc.append(["BadCo:careers"] + [""] * (len(SOURCE_COVERAGE_COLUMNS) - 1))
    wb.create_sheet("Closed_or_Rejected").append(list(CLOSED_REJECTED_COLUMNS))
    rs = wb.create_sheet("Run_Summary")
    rs.append(["Metric", "Value"])
    rs.append(["Relevant (shortlist) jobs", 1])
    wb.create_sheet("Resume_Tailoring").append(list(RESUME_TAILORING_COLUMNS))
    wb.save(run_dir / "Atlas_Jobs_20260909-000000_BAD.xlsx")

    report = validate_run_dir(run_dir, intent=intent)
    assert not report.passed
    assert any(i.code == "MISSING_LANE_ANCHOR" for i in report.issues)
