from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest

from atlas.persistence.sqlite import StateStore
from atlas.vscode_hunt.models import Backend
from atlas.vscode_hunt.service import VscodeHuntService
from atlas.vscode_hunt.verification_batch import (
    LeadClassification,
    evaluate_verification_batch_completion,
    manifest_hash,
    validate_verification_batch_result,
    normalize_selection_manifest,
)
from atlas.vscode_hunt.mcp_server import build_runtime_server
import asyncio
from atlas.vscode_hunt.report import build_v32_workbook
from atlas.vscode_hunt.outcomes import outcome_rows
from atlas.vscode_hunt.history import build_master_workbook


def _manifest() -> dict:
    return {"batches": [{"batch_id": f"batch-{index}", "leads": [{"lead_id": f"lead-{index}", "company": f"Company {index}", "source_url": "https://example.com/lead"}]} for index in range(1, 5)]}


def _result(run_id: str, task: dict, attempt_id: str, classification: str = LeadClassification.SOURCE_UNAVAILABLE.value) -> dict:
    return {
        "schema_version": 1, "run_id": run_id, "task_id": task["task_id"], "attempt_id": attempt_id,
        "batch_id": task["batch_id"], "manifest_hash": task["manifest_hash"],
        "assigned_lead_ids": task["assigned_lead_ids"],
        "outcomes": [{"lead_id": lead_id, "classification": classification, "company": next(lead["company"] for lead in task["leads"] if lead["lead_id"] == lead_id), "attempted_url": "https://example.com/lead", "error_details": "sanitized test outcome", "source_health": {"status": "unavailable"}} for lead_id in task["assigned_lead_ids"]],
        "queries": [], "navigation_summary": [], "browser_used": False,
        "source_health": {"status": "unavailable"}, "errors": [], "completion_claim": True,
    }


def test_four_batch_creation_is_idempotent_and_persistent(tmp_path: Path) -> None:
    with StateStore(tmp_path / "state.sqlite") as store:
        service = VscodeHuntService(store, tmp_path / "out")
        manifest = _manifest()
        run_id = service.create_verification_run(manifest, "run-four")
        assert service.create_verification_run(manifest, "run-four") == run_id
        assert len(service.get_ready_tasks(run_id)) == 4
        assert store.schema_version() >= 16
        assert store._conn.execute("SELECT COUNT(*) AS n FROM vscode_verification_items").fetchone()["n"] == 4
        with pytest.raises(ValueError, match="conflict"):
            different = {"batches": [{"batch_id": f"other-{index}", "leads": [{"lead_id": f"other-lead-{index}"}]} for index in range(1, 5)]}
            service.create_verification_run(different, run_id)


def test_batch_commit_tracks_each_terminal_lead_and_finalize(tmp_path: Path) -> None:
    with StateStore(tmp_path / "state.sqlite") as store:
        service = VscodeHuntService(store, tmp_path / "out")
        run_id = service.create_verification_run(_manifest(), "run-commit")
        task = service.get_ready_tasks(run_id)[0]
        attempt = service.create_or_start_attempt(run_id, task["task_id"], worker_invocation_id="worker-1")
        for lead_id in task["assigned_lead_ids"]:
            service.record_lead_checkpoint(run_id, task["task_id"], attempt.attempt_id, lead_id, {"classification": "SOURCE_UNAVAILABLE"})
        ack = service.commit_result_payload(run_id, task["task_id"], attempt.attempt_id, _result(run_id, task, attempt.attempt_id))
        assert ack["task_status"] == "COMPLETE"
        assert store._conn.execute("SELECT COUNT(*) AS n FROM vscode_verification_items WHERE status='TERMINAL'").fetchone()["n"] == 1
        assert service.finalize_run(run_id)["can_finish"] is False


def test_missing_lead_is_follow_up_and_unknown_identity_is_invalid() -> None:
    task = {"run_id": "run", "task_id": "task", "attempt_id": "attempt", "manifest_hash": "hash", "assigned_lead_ids": ["a", "b"]}
    result = {"schema_version": 1, "run_id": "run", "task_id": "task", "attempt_id": "attempt", "batch_id": "b", "manifest_hash": "hash", "assigned_lead_ids": ["a", "b"], "outcomes": [{"lead_id": "a", "classification": "SOURCE_UNAVAILABLE"}], "completion_claim": True}
    decision = evaluate_verification_batch_completion(result, task)
    assert decision["action"] == "FOLLOW_UP_REQUIRED"
    assert decision["missing_lead_ids"] == ["b"]
    result["outcomes"].append({"lead_id": "c", "classification": "SOURCE_UNAVAILABLE"})
    assert any("unknown lead" in error for error in validate_verification_batch_result(result, task))


def test_agent_contracts_are_fail_closed() -> None:
    root = Path(__file__).resolve().parents[1]
    verifier = (root / ".github/agents/atlas-job-lead-verifier-vscode.agent.md").read_text(encoding="utf-8")
    coordinator = (root / ".github/agents/atlas-r1-verification-coordinator.agent.md").read_text(encoding="utf-8")
    assert "name: atlas-job-lead-verifier-vscode" in verifier
    assert "name: atlas-r1-verification-coordinator" in coordinator
    assert "runSubagent" in coordinator and "atlas-runtime" in coordinator
    assert "openBrowserPage" not in coordinator and "navigatePage" not in coordinator
    assert "agents: [atlas-job-lead-verifier-vscode]" in coordinator


def test_mcp_exposes_verification_control_plane(tmp_path: Path) -> None:
    with StateStore(tmp_path / "state.sqlite") as store:
        service = VscodeHuntService(store, tmp_path / "out")
        names = {tool.name for tool in asyncio.run(build_runtime_server(service).list_tools())}
    assert {"create_verification_run", "create_or_start_attempt", "record_lead_checkpoint", "get_ready_tasks", "advance_verification_run"} <= names


def test_real_selection_manifest_normalizes_read_only() -> None:
    path = Path(r"C:\Atlas-Copilot-Runs\ATLAS_PRODUCTION_R1_20260911_171517\verification_selection_manifest.json")
    if not path.exists():
        pytest.skip("R1 parent manifest is not present")
    raw = json.loads(path.read_text(encoding="utf-8"))
    normalized = normalize_selection_manifest(raw, parent_run_id=raw["run_id"], source_reference=str(path))
    assert normalized["parent_run_id"] == "r1-prod-20260911_171517"
    assert sum(len(batch["leads"]) for batch in normalized["batches"]) == 28
    assert len({lead["lead_id"] for batch in normalized["batches"] for lead in batch["leads"]}) == 28
    assert normalized["normalized_manifest_hash"]


def test_dispatch_wave_reserves_four_distinct_attempts(tmp_path: Path) -> None:
    with StateStore(tmp_path / "state.sqlite") as store:
        service = VscodeHuntService(store, tmp_path / "out")
        run_id = service.create_verification_run(_manifest(), "run-wave")
        wave = service.advance_verification_run(run_id)
        assert wave["action"] == "DISPATCH_PRIMARY"
        assert len(wave["tasks"]) == 4
        assert len({item["task_id"] for item in wave["tasks"]}) == 4
        assert len({item["attempt_id"] for item in wave["tasks"]}) == 4
        assert service.advance_verification_run(run_id)["action"] == "WAIT_FOR_ACTIVE_WORKERS"
        assert all("resume_text" not in context and "private_profile" not in context and "candidate_name" not in context for context in (task["context"] for task in wave["tasks"]))


def test_dry_run_reconciles_committed_outcomes_and_history(tmp_path: Path) -> None:
    with StateStore(tmp_path / "state.sqlite") as store:
        service = VscodeHuntService(store, tmp_path / "out")
        run_id = service.create_verification_run(_manifest(), "run-dry")
        wave = service.advance_verification_run(run_id)
        for index, item in enumerate(wave["tasks"]):
            result = _result(run_id, item["context"], item["attempt_id"])
            result["batch_id"] = item["context"]["batch_id"]
            result["task_id"] = item["task_id"]
            result["attempt_id"] = item["attempt_id"]
            service.commit_result_payload(run_id, item["task_id"], item["attempt_id"], result)
        rows = outcome_rows(store._conn, run_id)
        assert len(rows) == 4
        workbook = build_v32_workbook(service.root, run_id, store._conn, kind="Final")
        master = build_master_workbook(store, service.root)
        from openpyxl import load_workbook
        current = load_workbook(workbook, read_only=True)
        cumulative = load_workbook(master, read_only=True)
        assert "Lead_Outcome_Audit" in current.sheetnames
        assert "Run_History" in cumulative.sheetnames
        current.close(); cumulative.close()