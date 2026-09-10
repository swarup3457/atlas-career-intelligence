from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from atlas.persistence.sqlite import StateStore
from atlas.vscode_hunt.models import Action, Backend
from atlas.vscode_hunt.report import SHEETS, build_workbook
from atlas.vscode_hunt.service import VscodeHuntService


def _result(run_id: str, task_id: str, attempt_id: str, lanes: list[str], *, incomplete: bool = False) -> dict:
    return {
        "run_id": run_id, "task_id": task_id, "attempt_id": attempt_id,
        "company_id": "fixture-a", "worker_invocation_id": "worker-1",
        "official_domain": "fixture.example", "career_url": "https://fixture.example/careers",
        "queries": ["java India"], "lanes_attempted": [] if incomplete else lanes,
        "result_states": ["NO_MATCH"], "detail_urls": [], "jobs": [],
        "rejections": [], "foreign_leads": [], "browser_errors": [],
        "evidence_quotes": [], "completion_claim": not incomplete,
    }


def test_stateless_follow_up_and_workbook(tmp_path: Path) -> None:
    with StateStore(tmp_path / "state.sqlite") as store:
        service = VscodeHuntService(store, tmp_path / "evidence")
        run_id = service.create_run([{"company_id": "fixture-a", "name": "Fixture A", "official_domain": "fixture.example", "lanes": ["JAVA_BACKEND"]}])
        task = service.next_tasks(run_id, materialize=True)[0]
        first = service.record_attempt(run_id, task["task_id"], "attempt-1", Backend.VSCODE_SUBAGENT)
        result_path = Path(task["result_path"])
        result_path.write_text(json.dumps(_result(run_id, task["task_id"], first.attempt_id, ["JAVA_BACKEND"], incomplete=True)), encoding="utf-8")
        follow_up = service.ingest(run_id, task["task_id"], first.attempt_id, result_path)
        assert follow_up["action"] == Action.FOLLOW_UP_REQUIRED.value
        assert follow_up["new_attempt_id"] != first.attempt_id
        second = service.record_attempt(run_id, task["task_id"], follow_up["new_attempt_id"], Backend.VSCODE_SUBAGENT, first.attempt_id)
        result_path.write_text(json.dumps(_result(run_id, task["task_id"], second.attempt_id, ["JAVA_BACKEND"])), encoding="utf-8")
        assert service.ingest(run_id, task["task_id"], second.attempt_id, result_path)["action"] == Action.COMPLETE.value
        workbook_path = build_workbook(service.root, run_id, store._conn)
        workbook = load_workbook(workbook_path, read_only=True)
        assert workbook.sheetnames == list(SHEETS)
        workbook.close()


def test_identity_and_schema_migration_are_idempotent(tmp_path: Path) -> None:
    db = tmp_path / "state.sqlite"
    with StateStore(db) as store:
        assert store.schema_version() == 13
    with StateStore(db) as store:
        assert store.schema_version() == 13
        service = VscodeHuntService(store, tmp_path / "evidence")
        run_id = service.create_run([{"company_id": "fixture-a", "name": "Fixture A", "official_domain": "fixture.example"}])
        task = service.next_tasks(run_id)[0]
        service.record_attempt(run_id, task["task_id"], "attempt-real", Backend.VSCODE_SUBAGENT)
        path = Path(task["result_path"])
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("{}", encoding="utf-8")
        assert service.ingest(run_id, task["task_id"], "attempt-fake", path)["action"] == Action.REJECTED_INVALID_RESULT.value
